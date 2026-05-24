from __future__ import annotations

import sys
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "07_API_BACKEND"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))


def main() -> None:
    from app.services.procurement_export import build_procurement_export_from_data, generate_procurement_export_filename

    project = SimpleNamespace(
        id=1,
        name="Centre Medical Pointe-Noire",
        client_name="SP2I",
        city="Pointe-Noire",
        country="Congo-Brazzaville",
        currency="FCFA",
        project_manager="Direction SP2I",
    )
    workflow = {"label": "Approvisionnement a valider", "primary_action": {"label": "Valider les arbitrages achat"}}
    scenario = {
        "scenario_id": "scenario-1",
        "scenario_nom": "Optimisation import",
        "run_status": "SIMULATED",
        "budget_local": 1000000,
        "budget_optimise": 850000,
        "economie_nette": 150000,
        "import_lines_count": 2,
        "local_lines_count": 1,
        "hybrid_lines_count": 1,
    }
    decisions = [{
        "id": 1,
        "lot": "L07",
        "family": "Electricite",
        "designation": "Tableau electrique",
        "quantity": 1,
        "unit": "u",
        "ai_decision": "IMPORT",
        "ai_score": 82,
        "ai_reason": "ROI positif",
        "proposed_decision": "IMPORT",
        "validated_decision": "",
        "validation_status": "REVIEW_REQUIRED",
        "estimated_local_cost": 1000000,
        "estimated_import_cost": 850000,
        "estimated_savings": 150000,
        "risk_level": "HIGH",
    }]
    events = [{"id": 1, "event_type": "PROCUREMENT_DECISION_UPDATED", "message": "Decision achat mise a jour."}]
    actions = [{"priority": "HIGH", "lot": "L07", "action_type": "DELIVERY", "status": "AT_RISK"}]

    buffer = build_procurement_export_from_data(project, workflow, scenario, {}, decisions, events, actions)
    workbook = load_workbook(buffer)
    expected = {
        "Synthèse projet",
        "Scénario actif",
        "Décisions achat",
        "Arbitrages à valider",
        "Risques achat",
        "Historique workflow",
        "Actions chantier",
        "Paramètres export",
    }
    missing = expected.difference(workbook.sheetnames)
    if missing:
        raise AssertionError(f"Missing sheets: {sorted(missing)}")
    filename = generate_procurement_export_filename(project)
    if not filename.startswith("Dossier_Achat_SP2I_Centre_Medical_Pointe_Noire_") or not filename.endswith(".xlsx"):
        raise AssertionError(f"Unexpected filename: {filename}")
    print("OK - procurement export workbook is valid.")


if __name__ == "__main__":
    main()

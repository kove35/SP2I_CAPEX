from __future__ import annotations

import json
import math
import re
import statistics
import sys
import time
import traceback
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "07_API_BACKEND"
DQE_PATH = ROOT / "03_DONNEES_REFERENCE" / "DQE_PROJECT_SP2I.xlsx"
SHEET_NAME = "DQE_CLEAN"

OUTPUT_MD = ROOT / "AUDIT_REFERENCE_ENGINE.md"
OUTPUT_JSON = ROOT / "AUDIT_REFERENCE_ENGINE.json"
OUTPUT_XLSX = ROOT / "AUDIT_REFERENCE_ENGINE.xlsx"
OUTPUT_HEATMAP = ROOT / "HEATMAP_REFERENCE_QUALITY.xlsx"
OUTPUT_FINANCIAL = ROOT / "TOP_FINANCIAL_ANOMALIES.xlsx"
OUTPUT_MAPPING = ROOT / "TOP_MAPPING_ERRORS.xlsx"

if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

try:
    from app.core.equipment_classification_engine import EquipmentClassificationEngine
    from app.core.financial_reference_engine import FinancialReferenceEngine
    from app.core.semantic_normalization_engine import SemanticNormalizationEngine
except Exception as import_error:  # pragma: no cover - reported in script output
    EquipmentClassificationEngine = None
    FinancialReferenceEngine = None
    SemanticNormalizationEngine = None
    IMPORT_ERROR = import_error
else:
    IMPORT_ERROR = None


VALID_TYPES = {"ARTICLE", "PRESTATION", "TRAVAUX", "FOURNITURE", "EQUIPEMENT", "TOTAL", "VALIDATION", "DOCUMENTAIRE", "TITRE_SECTION"}
EXCLUDED_TYPES = {"TOTAL", "VALIDATION", "DOCUMENTAIRE", "TITRE_SECTION"}
NON_IMPORTABLE_TYPES = {"PRESTATION", "TRAVAUX", "TOTAL", "VALIDATION", "DOCUMENTAIRE", "TITRE_SECTION"}

TYPE_KEYWORDS = {
    "TOTAL": ["total", "sous total", "sous-total", "montant total"],
    "VALIDATION": ["validation", "visa", "approuve", "bon pour accord"],
    "DOCUMENTAIRE": ["dossier", "document", "plan", "rapport", "pv", "fiche technique", "manuel", "doe"],
    "TITRE_SECTION": ["lot ", "chapitre", "section", "generalites", "prescriptions"],
    "PRESTATION": ["pose", "installation", "mise en service", "essai", "maintenance", "formation", "reparation"],
    "TRAVAUX": ["terrassement", "demolition", "maconnerie", "beton", "coffrage", "peinture", "revetement"],
}

FAMILY_ALIASES = {
    "ELECTRICITE": ["electric", "courant fort", "courant faible", "tgbt", "tableau", "groupe electrogene"],
    "PLOMBERIE": ["plomberie", "sanitaire", "eau", "pvc", "cuivre", "pompe"],
    "HVAC": ["hvac", "cvc", "clim", "vrv", "cta", "ventilation", "extraction"],
    "MENUISERIE ALU": ["aluminium", "alu", "mur rideau", "facade"],
    "MENUISERIE BOIS": ["bois", "placard", "porte bois", "cuisine"],
    "REVETEMENTS": ["carrelage", "faience", "peinture", "faux plafond", "sol souple"],
    "GROS OEUVRE": ["beton", "acier", "coffrage", "maconnerie", "gros oeuvre"],
    "VRD": ["vrd", "voirie", "drainage", "terrassement", "reseau"],
    "EQUIPEMENTS MEDICAUX": ["medical", "laboratoire", "radiologie", "dentaire", "sterilisation"],
    "SECURITE": ["incendie", "videosurveillance", "controle acces", "intrusion"],
    "ASCENSEURS": ["ascenseur", "lift"],
    "CUISINE PROFESSIONNELLE": ["cuisine professionnelle", "four", "hotte", "inox"],
    "MOBILIER": ["mobilier", "bureau", "chaise", "armoire"],
    "RESEAUX IT": ["reseau informatique", "it", "switch", "baie", "wifi"],
    "ENERGIE SOLAIRE": ["solaire", "pv", "panneau solaire", "onduleur"],
    "HYDRAULIQUE": ["hydraulique", "surpresseur", "pompage"],
    "FACADE": ["facade", "bardage", "alucobond"],
    "CHARPENTE": ["charpente", "structure metallique"],
    "ETANCHEITE": ["etancheite", "membrane", "bitume"],
    "AMENAGEMENTS EXTERIEURS": ["amenagement", "cloture", "jardin", "exterieur"],
}

NEGATIVE_CLASSIFICATION_TERMS = {
    "RESEAUX IT": ["materiel", "mobilisation"],
    "ENERGIE SOLAIRE": ["fourniture et pose"],
}

FAMILY_BENCHMARKS = {
    "ELECTRICITE": (15_000, 45_000_000),
    "PLOMBERIE": (8_000, 18_000_000),
    "HVAC": (350_000, 55_000_000),
    "MENUISERIE ALU": (85_000, 120_000_000),
    "MENUISERIE BOIS": (65_000, 45_000_000),
    "REVETEMENTS": (6_000, 35_000_000),
    "GROS OEUVRE": (45_000, 180_000_000),
    "VRD": (60_000, 250_000_000),
    "EQUIPEMENTS MEDICAUX": (500_000, 450_000_000),
    "SECURITE": (45_000, 80_000_000),
    "ASCENSEURS": (18_000_000, 95_000_000),
    "CUISINE PROFESSIONNELLE": (250_000, 90_000_000),
    "MOBILIER": (35_000, 35_000_000),
    "RESEAUX IT": (25_000, 60_000_000),
    "ENERGIE SOLAIRE": (250_000, 180_000_000),
    "HYDRAULIQUE": (250_000, 80_000_000),
    "FACADE": (85_000, 180_000_000),
    "CHARPENTE": (120_000, 250_000_000),
    "ETANCHEITE": (12_000, 80_000_000),
    "AMENAGEMENTS EXTERIEURS": (25_000, 120_000_000),
}


def strip_accents(value: Any) -> str:
    text = "" if value is None else str(value)
    decomposed = unicodedata.normalize("NFKD", text)
    return "".join(char for char in decomposed if not unicodedata.combining(char))


def normalize_text(value: Any) -> str:
    text = strip_accents(value).lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_key(value: Any) -> str:
    return normalize_text(value).replace(" ", "_").upper()


def number(value: Any, default: float = 0.0) -> float:
    if value is None:
        return default
    if isinstance(value, (int, float)):
        if math.isnan(value):
            return default
        return float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return default


def pct(value: float) -> float:
    return round(max(0.0, min(100.0, value)), 2)


def score_from_error_rate(error_rate: float) -> float:
    return pct(100.0 - error_rate * 100.0)


def choose(row: dict[str, Any], aliases: list[str], default: Any = "") -> Any:
    for alias in aliases:
        key = normalize_key(alias)
        if key in row and row[key] not in (None, ""):
            return row[key]
    return default


def read_dqe() -> tuple[list[dict[str, Any]], list[str]]:
    if not DQE_PATH.exists():
        raise FileNotFoundError(f"Fichier DQE introuvable: {DQE_PATH}")
    wb = load_workbook(DQE_PATH, data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Onglet {SHEET_NAME} introuvable. Onglets detectes: {wb.sheetnames}")
    ws = wb[SHEET_NAME]
    raw_headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    headers = [normalize_key(header or f"COL_{index + 1}") for index, header in enumerate(raw_headers)]
    rows = []
    for row_index, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        payload = {headers[index]: value for index, value in enumerate(values)}
        if any(value not in (None, "") for value in payload.values()):
            payload["_ROWNUM"] = row_index
            rows.append(payload)
    wb.close()
    return rows, headers


def infer_family(text: str) -> str:
    normalized = normalize_text(text)
    for family, keywords in FAMILY_ALIASES.items():
        if any(keyword in normalized for keyword in keywords):
            return family
    return "GENERAL"


def infer_expected_type(designation: str) -> str:
    text = normalize_text(designation)
    for type_name, keywords in TYPE_KEYWORDS.items():
        if any(keyword in text for keyword in keywords):
            return type_name
    return "ARTICLE"


def family_matches(declared: str, inferred: str) -> bool:
    aliases = {
        "reseaux it": "it reseaux",
        "reseaux_it": "it_reseaux",
        "energie solaire": "solaire",
        "energie_solaire": "solaire",
        "equipements medicaux": "medical",
        "equipements_medicaux": "medical",
        "menuiserie alu": "menuiserie",
        "menuiserie_alu": "menuiserie",
        "menuiserie bois": "menuiserie",
        "menuiserie_bois": "menuiserie",
        "finitions": "revetements",
        "finition": "revetements",
        "bardage": "facade",
    }
    declared_norm = aliases.get(normalize_text(declared), normalize_text(declared))
    inferred_norm = aliases.get(normalize_text(inferred), normalize_text(inferred))
    if not declared_norm or declared_norm == "general":
        return inferred in {"GENERAL", "UNKNOWN"}
    if declared_norm in inferred_norm or inferred_norm in declared_norm:
        return True
    aliases = FAMILY_ALIASES.get(inferred, [])
    return any(alias in declared_norm for alias in aliases)


def classify_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    classifier = EquipmentClassificationEngine() if EquipmentClassificationEngine else None
    financial = FinancialReferenceEngine() if FinancialReferenceEngine else None
    normalizer = SemanticNormalizationEngine() if SemanticNormalizationEngine else None
    enriched = []
    for row in rows:
        designation = str(choose(row, ["DESIGNATION", "LIBELLE", "DESCRIPTION", "ARTICLE", "DESIGNATION_COMPLETE"], ""))
        lot = str(choose(row, ["LOT", "FAMILLE", "FAMILLE_LOT"], ""))
        sous_lot = str(choose(row, ["SOUS_LOT", "SOUS LOT", "SUBLOT", "POSTE"], ""))
        type_ligne = str(choose(row, ["TYPE_LIGNE", "TYPE LIGNE", "TYPE", "LINE_TYPE"], "") or infer_expected_type(designation)).upper()
        family_declared = str(choose(row, ["FAMILLE", "FAMILY", "LOT_FAMILLE", "CATEGORIE"], lot) or "")
        normalized = normalizer.normalize(designation) if normalizer else {"normalized_designation": normalize_text(designation), "normalization_confidence_score": 50}
        context = {
            "designation": designation,
            "LOT": lot,
            "SOUS_LOT": sous_lot,
            "FAMILLE": family_declared,
            "TYPE_LIGNE": type_ligne,
            "region": "CONGO_BRAZZAVILLE",
            "quality": "MEDIUM",
        }
        engine_classification = classifier.classify_line(context) if classifier else {}
        inferred_family = engine_classification.get("family") or infer_family(f"{family_declared} {lot} {sous_lot} {designation}")
        family_lookup = str(inferred_family or "").replace("_", " ")
        if any(term in normalize_text(designation) for term in NEGATIVE_CLASSIFICATION_TERMS.get(family_lookup, [])):
            inferred_family = infer_family(f"{family_declared} {lot} {sous_lot}")
            engine_classification = {**engine_classification, "equipment_type": "UNKNOWN", "equipment_class": "UNKNOWN"}
        if inferred_family == "UNKNOWN":
            inferred_family = infer_family(f"{family_declared} {lot} {sous_lot} {designation}")
        equipment_type = engine_classification.get("equipment_type")
        reference = financial.get_reference(equipment_type, region="CONGO_BRAZZAVILLE") if financial and equipment_type else {}
        price = number(choose(row, ["PU_ESTIME_LOCAL_FCFA", "PU ESTIME LOCAL FCFA", "PU_LOCAL_FCFA", "PRIX_LOCAL", "CAPEX_LOCAL", "PRIX_TOTAL_HT", "MONTANT"], 0))
        quantity = number(choose(row, ["QTE", "QUANTITE", "QUANTITY"], 1), 1)
        total = price * quantity if quantity and price else price
        enriched.append({
            "rownum": row["_ROWNUM"],
            "designation": designation,
            "normalized_designation": normalized.get("normalized_designation", normalize_text(designation)),
            "lot": lot,
            "sous_lot": sous_lot,
            "type_ligne": type_ligne,
            "expected_type_ligne": infer_expected_type(designation),
            "family_declared": family_declared or lot or "NON_RENSEIGNE",
            "family_inferred": inferred_family,
            "equipment_type": equipment_type or "",
            "equipment_class": engine_classification.get("equipment_class", ""),
            "semantic_score": number(engine_classification.get("semantic_confidence_score") or normalized.get("normalization_confidence_score"), 0),
            "procurement_score": number(engine_classification.get("procurement_confidence_score"), 0),
            "reference": reference,
            "pu_local": price,
            "quantity": quantity,
            "total_estimated": total,
            "batiment": str(choose(row, ["BATIMENT", "BAT", "BUILDING"], "")),
            "niveau": str(choose(row, ["NIVEAU", "ETAGE", "LEVEL"], "")),
            "zone": str(choose(row, ["ZONE", "LOCALISATION", "LOCAL"], "")),
            "raw": row,
        })
    return enriched


def audit_taxonomy(rows: list[dict[str, Any]]) -> dict[str, Any]:
    mappable = [row for row in rows if row["type_ligne"] not in EXCLUDED_TYPES]
    general = [row for row in mappable if normalize_text(row["family_declared"]) in {"", "general", "non renseigne"} or row["family_inferred"] in {"GENERAL", "UNKNOWN"}]
    errors = []
    for row in mappable:
        if not family_matches(row["family_declared"], row["family_inferred"]):
            errors.append({
                "row": row["rownum"],
                "designation": row["designation"],
                "declared_family": row["family_declared"],
                "inferred_family": row["family_inferred"],
                "sous_lot": row["sous_lot"],
                "equipment_type": row["equipment_type"],
                "severity": "HIGH" if row["family_inferred"] not in {"GENERAL", "UNKNOWN"} else "MEDIUM",
            })
    family_counter = Counter(row["family_declared"] for row in mappable)
    covered_family = Counter(row["family_declared"] for row in mappable if row["family_inferred"] not in {"GENERAL", "UNKNOWN"})
    sublot_counter = Counter(row["sous_lot"] or "NON_RENSEIGNE" for row in mappable)
    covered_sublot = Counter(row["sous_lot"] or "NON_RENSEIGNE" for row in mappable if row["family_inferred"] not in {"GENERAL", "UNKNOWN"})
    coverage_by_family = [
        {"family": family, "rows": count, "classified": covered_family[family], "coverage": pct(covered_family[family] / count * 100 if count else 0)}
        for family, count in family_counter.most_common()
    ]
    coverage_by_sublot = [
        {"sous_lot": sublot, "rows": count, "classified": covered_sublot[sublot], "coverage": pct(covered_sublot[sublot] / count * 100 if count else 0)}
        for sublot, count in sublot_counter.most_common()
    ]
    error_rate = (len(errors) + len(general) * 0.5) / max(len(mappable), 1)
    return {
        "TAXONOMY_SCORE": score_from_error_rate(error_rate),
        "classified_rows": len(mappable) - len(general),
        "general_rows": len(general),
        "mapping_error_count": len(errors),
        "TOP_MAPPING_ERRORS": errors[:100],
        "COVERAGE_BY_FAMILY": coverage_by_family,
        "COVERAGE_BY_SUBLOT": coverage_by_sublot,
    }


def audit_type_lines(rows: list[dict[str, Any]]) -> dict[str, Any]:
    invalid = []
    errors = []
    for row in rows:
        type_ligne = row["type_ligne"]
        expected = row["expected_type_ligne"]
        text = normalize_text(row["designation"])
        if type_ligne not in VALID_TYPES:
            invalid.append({"row": row["rownum"], "designation": row["designation"], "type_ligne": type_ligne, "reason": "TYPE_LIGNE inconnu"})
        if expected in EXCLUDED_TYPES and type_ligne not in EXCLUDED_TYPES:
            errors.append({"row": row["rownum"], "designation": row["designation"], "type_ligne": type_ligne, "expected": expected, "issue": "Ligne a exclure integree au calcul"})
        if type_ligne == "ARTICLE" and expected in {"PRESTATION", "TRAVAUX"}:
            errors.append({"row": row["rownum"], "designation": row["designation"], "type_ligne": type_ligne, "expected": expected, "issue": "Prestation/travaux classes ARTICLE"})
        if type_ligne == "PRESTATION" and any(token in text for token in ["groupe electrogene", "ascenseur", "climatiseur", "pompe", "tableau"]):
            errors.append({"row": row["rownum"], "designation": row["designation"], "type_ligne": type_ligne, "expected": "ARTICLE", "issue": "Equipement probablement classe PRESTATION"})
    score = score_from_error_rate((len(invalid) + len(errors)) / max(len(rows), 1))
    return {
        "TYPE_LINE_QUALITY_SCORE": score,
        "INVALID_TYPE_LINES": invalid[:100],
        "CLASSIFICATION_ERRORS": errors[:150],
        "invalid_count": len(invalid),
        "classification_error_count": len(errors),
    }


def family_benchmark(row: dict[str, Any]) -> tuple[float, float]:
    family = row["family_inferred"] if row["family_inferred"] not in {"UNKNOWN", "GENERAL"} else row["family_declared"]
    for key, bounds in FAMILY_BENCHMARKS.items():
        if family_matches(family, key):
            return bounds
    return (5_000, 250_000_000)


def audit_financial(rows: list[dict[str, Any]]) -> dict[str, Any]:
    candidates = [row for row in rows if row["type_ligne"] not in EXCLUDED_TYPES]
    anomalies = []
    grouped_prices = defaultdict(list)
    for row in candidates:
        if row["pu_local"] > 0:
            grouped_prices[row["family_declared"]].append(row["pu_local"])
    stats = {}
    for family, prices in grouped_prices.items():
        if len(prices) >= 4:
            ordered = sorted(prices)
            q1 = ordered[len(ordered) // 4]
            q3 = ordered[(len(ordered) * 3) // 4]
            iqr = max(q3 - q1, 1)
            stats[family] = {"q1": q1, "q3": q3, "lower": q1 - 1.5 * iqr, "upper": q3 + 1.5 * iqr, "median": statistics.median(ordered)}
    for row in candidates:
        price = row["pu_local"]
        reference = row["reference"]
        issues = []
        severity = 0
        benchmark_min, benchmark_max = family_benchmark(row)
        ref_min = number(reference.get("price_min"), benchmark_min) if reference else benchmark_min
        ref_max = number(reference.get("price_max"), benchmark_max) if reference else benchmark_max
        if price <= 0:
            issues.append("prix nul ou negatif")
            severity += 45
        elif price < ref_min * 0.15:
            issues.append("prix tres inferieur au benchmark, zeros perdus probables")
            severity += 45
        elif price < ref_min:
            issues.append("prix inferieur au benchmark Afrique centrale")
            severity += 28
        elif price > ref_max * 3:
            issues.append("prix tres superieur au benchmark, magnitude suspecte")
            severity += 40
        elif price > ref_max:
            issues.append("prix superieur au benchmark Afrique centrale")
            severity += 24
        family_stat = stats.get(row["family_declared"])
        if family_stat and price > 0 and (price < family_stat["lower"] or price > family_stat["upper"]):
            issues.append("outlier statistique IQR famille")
            severity += 18
        if issues:
            anomalies.append({
                "row": row["rownum"],
                "designation": row["designation"],
                "family": row["family_declared"],
                "type_ligne": row["type_ligne"],
                "price": round(price, 2),
                "benchmark_min": round(ref_min, 2),
                "benchmark_max": round(ref_max, 2),
                "equipment_type": row["equipment_type"],
                "severity_score": min(severity, 100),
                "issues": "; ".join(issues),
            })
    anomalies.sort(key=lambda item: item["severity_score"], reverse=True)
    score = score_from_error_rate(sum(item["severity_score"] for item in anomalies) / max(len(candidates), 1) / 100)
    return {
        "FINANCIAL_SANITY_SCORE": score,
        "PRICE_OUTLIERS": [item for item in anomalies if "outlier" in item["issues"]][:100],
        "SUSPICIOUS_PRICES": anomalies[:200],
        "MAGNITUDE_ERRORS": [item for item in anomalies if "zeros" in item["issues"] or "magnitude" in item["issues"]][:100],
        "financial_anomaly_count": len(anomalies),
    }


def is_importable(row: dict[str, Any]) -> bool:
    if row["type_ligne"] in NON_IMPORTABLE_TYPES:
        return False
    family = normalize_text(f"{row['family_declared']} {row['family_inferred']} {row['designation']}")
    if any(token in family for token in ["pose", "main oeuvre", "maintenance", "validation", "document", "beton", "terrassement"]):
        return False
    return any(token in family for token in ["groupe", "tableau", "cable", "clim", "split", "pompe", "ascenseur", "mobilier", "medical", "camera", "controle acces", "solaire"])


def audit_roi_procurement(rows: list[dict[str, Any]], financial_anomalies: list[dict[str, Any]]) -> tuple[dict[str, Any], dict[str, Any]]:
    anomaly_rows = {item["row"] for item in financial_anomalies}
    roi_outliers = []
    fake_savings = []
    import_errors = []
    sourcing_weakness = []
    for row in rows:
        if row["type_ligne"] in EXCLUDED_TYPES:
            continue
        price = row["pu_local"]
        importable = is_importable(row)
        if row["type_ligne"] in NON_IMPORTABLE_TYPES and importable:
            import_errors.append({"row": row["rownum"], "designation": row["designation"], "type_ligne": row["type_ligne"], "issue": "Ligne non importable marquee importable"})
        if not importable and row["equipment_type"] not in {"", "UNKNOWN", "NON_IMPORTABLE"} and row["type_ligne"] in {"ARTICLE", "FOURNITURE", "EQUIPEMENT"}:
            sourcing_weakness.append({"row": row["rownum"], "designation": row["designation"], "equipment_type": row["equipment_type"], "issue": "Equipement detecte mais importabilite faible"})
        if price <= 0 or not importable:
            continue
        reference = row["reference"]
        ref_min, ref_max = family_benchmark(row)
        if reference:
            ref_min = number(reference.get("price_min"), ref_min)
            ref_max = number(reference.get("price_max"), ref_max)
        china_floor = ref_min * 0.72
        estimated_landed = max(price * 0.78, china_floor) * 1.205
        gain = price - estimated_landed
        roi = gain / estimated_landed if estimated_landed > 0 else 0
        risky = row["rownum"] in anomaly_rows or price < ref_min
        if roi > 0.45 or roi < -0.35:
            roi_outliers.append({"row": row["rownum"], "designation": row["designation"], "price": round(price, 2), "estimated_landed": round(estimated_landed, 2), "roi": round(roi, 4), "issue": "ROI hors plage decisionnelle"})
        if risky and gain > 0:
            fake_savings.append({"row": row["rownum"], "designation": row["designation"], "price": round(price, 2), "estimated_landed": round(estimated_landed, 2), "fake_gain": round(gain, 2), "roi": round(roi, 4), "issue": "Gain base sur prix local/benchmark non fiable"})
    active = [row for row in rows if row["type_ligne"] not in EXCLUDED_TYPES]
    roi_score = score_from_error_rate((len(roi_outliers) + len(fake_savings) * 1.5) / max(len(active), 1))
    procurement_score = score_from_error_rate((len(import_errors) + len(sourcing_weakness) * 0.5) / max(len(active), 1))
    return {
        "ROI_RELIABILITY_SCORE": roi_score,
        "TOP_FAKE_SAVINGS": sorted(fake_savings, key=lambda item: item["fake_gain"], reverse=True)[:100],
        "ROI_OUTLIERS": sorted(roi_outliers, key=lambda item: abs(item["roi"]), reverse=True)[:100],
    }, {
        "PROCUREMENT_REALISM_SCORE": procurement_score,
        "IMPORTABILITY_ERRORS": import_errors[:100],
        "SOURCING_WEAKNESS": sourcing_weakness[:100],
        "importable_rows": sum(1 for row in active if is_importable(row)),
    }


def audit_context(rows: list[dict[str, Any]]) -> dict[str, Any]:
    active = [row for row in rows if row["type_ligne"] not in EXCLUDED_TYPES]
    errors = []
    for row in active:
        missing = [key for key in ["batiment", "niveau", "zone"] if not row.get(key)]
        if missing:
            errors.append({"row": row["rownum"], "designation": row["designation"], "missing": ", ".join(missing), "lot": row["lot"], "sous_lot": row["sous_lot"]})
    score = score_from_error_rate(len(errors) / max(len(active), 1))
    return {
        "CONTEXT_PROPAGATION_SCORE": score,
        "BUILDING_MAPPING_ERRORS": errors[:150],
        "missing_context_count": len(errors),
    }


def audit_reference_quality(rows: list[dict[str, Any]]) -> dict[str, Any]:
    active = [row for row in rows if row["type_ligne"] not in EXCLUDED_TYPES]
    normalized_counter = Counter(row["normalized_designation"] for row in active if row["normalized_designation"])
    duplicates = [{"normalized_designation": key, "count": count} for key, count in normalized_counter.most_common() if count > 1]
    unique_count = len(normalized_counter)
    duplication_rate = (sum(item["count"] - 1 for item in duplicates) / max(len(active), 1)) * 100
    weak_norm = [row for row in active if len(row["normalized_designation"].split()) < 2 or row["semantic_score"] < 55]
    fragmentation = len([key for key, count in normalized_counter.items() if count == 1]) / max(unique_count, 1) * 100
    normalization_quality = pct(100 - len(weak_norm) / max(len(active), 1) * 100)
    reference_score = pct(100 - duplication_rate * 0.25 - fragmentation * 0.05 - len(weak_norm) / max(len(active), 1) * 25)
    return {
        "REFERENCE_ENGINE_SCORE": reference_score,
        "DUPLICATION_RATE": round(duplication_rate, 2),
        "NORMALIZATION_QUALITY": normalization_quality,
        "fragmentation_rate": round(fragmentation, 2),
        "duplicate_references": duplicates[:100],
        "weak_normalization": [{"row": row["rownum"], "designation": row["designation"], "normalized": row["normalized_designation"], "semantic_score": row["semantic_score"]} for row in weak_norm[:100]],
    }


def audit_cockpit(scores: dict[str, float], financial: dict[str, Any], roi: dict[str, Any], type_audit: dict[str, Any]) -> dict[str, Any]:
    inconsistencies = []
    if financial["financial_anomaly_count"]:
        inconsistencies.append({"area": "Finance", "issue": "Anomalies prix susceptibles de fausser CAPEX, waterfall et heatmaps", "count": financial["financial_anomaly_count"]})
    if roi["TOP_FAKE_SAVINGS"]:
        inconsistencies.append({"area": "ROI", "issue": "Faux gains import possibles si prix local/FOB non fiables", "count": len(roi["TOP_FAKE_SAVINGS"])})
    if type_audit["classification_error_count"]:
        inconsistencies.append({"area": "AG Grid", "issue": "Types de lignes non fiables susceptibles de polluer les aggregations", "count": type_audit["classification_error_count"]})
    score = pct(scores["FINANCIAL_SANITY_SCORE"] * 0.35 + scores["ROI_RELIABILITY_SCORE"] * 0.35 + scores["TYPE_LINE_QUALITY_SCORE"] * 0.3)
    return {
        "DASHBOARD_CREDIBILITY_SCORE": score,
        "KPI_INCONSISTENCIES": inconsistencies,
    }


def rows_to_sheet(wb: Workbook, title: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title[:31])
    if not rows:
        ws.append(["status"])
        ws.append(["Aucune ligne"])
        return
    headers = list(rows[0].keys())
    ws.append(headers)
    for item in rows:
        ws.append([item.get(header) for header in headers])
    style_sheet(ws)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="0F172A")
    header_font = Font(color="F8FAFC", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 12), 60)
    ws.freeze_panes = "A2"


def save_workbook(path: Path, sheets: dict[str, list[dict[str, Any]]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        rows_to_sheet(wb, title, rows)
    wb.save(path)


def build_heatmap_rows(rows: list[dict[str, Any]], taxonomy: dict[str, Any], financial: dict[str, Any]) -> list[dict[str, Any]]:
    by_key: dict[tuple[str, str], dict[str, Any]] = {}
    financial_rows = {item["row"]: item for item in financial["SUSPICIOUS_PRICES"]}
    mapping_rows = {item["row"]: item for item in taxonomy["TOP_MAPPING_ERRORS"]}
    for row in rows:
        if row["type_ligne"] in EXCLUDED_TYPES:
            continue
        key = (row["family_declared"], row["sous_lot"] or "NON_RENSEIGNE")
        current = by_key.setdefault(key, {"family": key[0], "sous_lot": key[1], "rows": 0, "mapping_errors": 0, "financial_anomalies": 0, "quality_score": 100.0})
        current["rows"] += 1
        current["mapping_errors"] += 1 if row["rownum"] in mapping_rows else 0
        current["financial_anomalies"] += 1 if row["rownum"] in financial_rows else 0
    for item in by_key.values():
        penalty = (item["mapping_errors"] * 35 + item["financial_anomalies"] * 45) / max(item["rows"], 1)
        item["quality_score"] = pct(100 - penalty)
        item["risk_level"] = "CRITICAL" if item["quality_score"] < 55 else "WARNING" if item["quality_score"] < 75 else "OK"
    return sorted(by_key.values(), key=lambda item: item["quality_score"])


def markdown_report(audit: dict[str, Any]) -> str:
    scores = audit["scores"]
    decision = "GO conditionnel" if scores["GLOBAL_REFERENCE_SCORE"] >= 70 else "NO-GO avant remediation"
    lines = [
        "# AUDIT REFERENCE ENGINE - SP2I_CAPEX",
        "",
        f"- Date audit: {audit['metadata']['generated_at']}",
        f"- Source unique: `{audit['metadata']['source_file']}`",
        f"- Onglet: `{audit['metadata']['sheet']}`",
        f"- Lignes analysees: {audit['metadata']['rows']}",
        f"- Decision: **{decision}**",
        "",
        "## KPIs finaux",
        "",
    ]
    for key, value in scores.items():
        lines.append(f"- **{key}**: {value}/100")
    lines.extend([
        "",
        "## Synthese executive",
        "",
        f"Le score global du moteur de reference est de **{scores['GLOBAL_REFERENCE_SCORE']}/100**. "
        "L'audit mesure la fiabilite du DQE propre comme base d'un futur MASTER_REFERENCE enterprise Afrique centrale, sans modifier les donnees ni les routes.",
        "",
        "## Principaux risques",
        "",
    ])
    risks = []
    if audit["financial"]["SUSPICIOUS_PRICES"]:
        first = audit["financial"]["SUSPICIOUS_PRICES"][0]
        risks.append(f"Prix suspect majeur ligne {first['row']}: {first['designation']} ({first['price']:,.0f} FCFA).")
    if audit["taxonomy"]["TOP_MAPPING_ERRORS"]:
        first = audit["taxonomy"]["TOP_MAPPING_ERRORS"][0]
        risks.append(f"Mapping taxonomique a verifier ligne {first['row']}: {first['declared_family']} -> {first['inferred_family']}.")
    if audit["roi"]["TOP_FAKE_SAVINGS"]:
        first = audit["roi"]["TOP_FAKE_SAVINGS"][0]
        risks.append(f"Faux gain potentiel ligne {first['row']}: {first['fake_gain']:,.0f} FCFA.")
    if not risks:
        risks.append("Aucun risque critique dominant detecte sur les seuils d'audit.")
    lines.extend([f"- {risk}" for risk in risks])
    lines.extend([
        "",
        "## Recommandations",
        "",
        "- Valider manuellement les lignes a prix local inferieur aux benchmarks Afrique centrale.",
        "- Exclure explicitement TOTAL, VALIDATION, DOCUMENTAIRE et TITRE_SECTION des aggregations cockpit.",
        "- Degrader les recommandations ROI lorsque le prix local ou le FOB est marque suspect.",
        "- Utiliser la heatmap qualite comme backlog de remediation avant construction du MASTER_REFERENCE.",
        "",
        "## Livrables generes",
        "",
        "- `AUDIT_REFERENCE_ENGINE.json`",
        "- `AUDIT_REFERENCE_ENGINE.xlsx`",
        "- `HEATMAP_REFERENCE_QUALITY.xlsx`",
        "- `TOP_FINANCIAL_ANOMALIES.xlsx`",
        "- `TOP_MAPPING_ERRORS.xlsx`",
    ])
    return "\n".join(lines) + "\n"


def main() -> int:
    start = time.perf_counter()
    print("=== SP2I_CAPEX Reference Engine Audit ===")
    print(f"Source: {DQE_PATH}")
    try:
        rows, headers = read_dqe()
        print(f"Loaded {len(rows)} rows from {SHEET_NAME}")
        print(f"Columns: {', '.join(headers)}")
        if IMPORT_ERROR:
            print(f"WARNING: backend imports degraded: {IMPORT_ERROR}")
        enriched = classify_rows(rows)
        taxonomy = audit_taxonomy(enriched)
        type_audit = audit_type_lines(enriched)
        financial = audit_financial(enriched)
        roi, procurement = audit_roi_procurement(enriched, financial["SUSPICIOUS_PRICES"])
        context = audit_context(enriched)
        reference = audit_reference_quality(enriched)
        scores = {
            "TAXONOMY_SCORE": taxonomy["TAXONOMY_SCORE"],
            "PROCUREMENT_SCORE": procurement["PROCUREMENT_REALISM_SCORE"],
            "FINANCIAL_SCORE": financial["FINANCIAL_SANITY_SCORE"],
            "ROI_RELIABILITY_SCORE": roi["ROI_RELIABILITY_SCORE"],
            "TYPE_LINE_QUALITY_SCORE": type_audit["TYPE_LINE_QUALITY_SCORE"],
            "CONTEXT_PROPAGATION_SCORE": context["CONTEXT_PROPAGATION_SCORE"],
            "REFERENCE_ENGINE_SCORE": reference["REFERENCE_ENGINE_SCORE"],
        }
        cockpit = audit_cockpit({
            **scores,
            "FINANCIAL_SANITY_SCORE": financial["FINANCIAL_SANITY_SCORE"],
        }, financial, roi, type_audit)
        scores["DASHBOARD_CREDIBILITY_SCORE"] = cockpit["DASHBOARD_CREDIBILITY_SCORE"]
        scores["GLOBAL_REFERENCE_SCORE"] = pct(
            scores["TAXONOMY_SCORE"] * 0.16
            + scores["PROCUREMENT_SCORE"] * 0.14
            + scores["FINANCIAL_SCORE"] * 0.18
            + scores["ROI_RELIABILITY_SCORE"] * 0.16
            + scores["DASHBOARD_CREDIBILITY_SCORE"] * 0.14
            + scores["REFERENCE_ENGINE_SCORE"] * 0.14
            + scores["TYPE_LINE_QUALITY_SCORE"] * 0.08
        )
        heatmap = build_heatmap_rows(enriched, taxonomy, financial)
        audit = {
            "metadata": {
                "generated_at": datetime.now().isoformat(timespec="seconds"),
                "source_file": str(DQE_PATH),
                "sheet": SHEET_NAME,
                "rows": len(enriched),
                "columns": headers,
                "duration_seconds": round(time.perf_counter() - start, 2),
                "mode": "READ_ONLY_AUDIT",
            },
            "scores": scores,
            "taxonomy": taxonomy,
            "type_ligne": type_audit,
            "financial": financial,
            "roi": roi,
            "procurement": procurement,
            "context": context,
            "cockpit": cockpit,
            "reference": reference,
            "heatmap": heatmap,
        }
        OUTPUT_JSON.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding="utf-8")
        OUTPUT_MD.write_text(markdown_report(audit), encoding="utf-8")
        summary_rows = [{"kpi": key, "score": value} for key, value in scores.items()]
        save_workbook(OUTPUT_XLSX, {
            "SUMMARY": summary_rows,
            "FINANCIAL_ANOMALIES": financial["SUSPICIOUS_PRICES"][:200],
            "MAPPING_ERRORS": taxonomy["TOP_MAPPING_ERRORS"][:200],
            "TYPE_LINE_ERRORS": type_audit["CLASSIFICATION_ERRORS"][:200],
            "ROI_OUTLIERS": roi["ROI_OUTLIERS"][:150],
            "FAKE_SAVINGS": roi["TOP_FAKE_SAVINGS"][:150],
            "PROCUREMENT": procurement["IMPORTABILITY_ERRORS"] + procurement["SOURCING_WEAKNESS"],
            "CONTEXT": context["BUILDING_MAPPING_ERRORS"][:200],
            "REFERENCE": reference["duplicate_references"][:200],
            "HEATMAP": heatmap,
        })
        save_workbook(OUTPUT_HEATMAP, {"HEATMAP_REFERENCE_QUALITY": heatmap})
        save_workbook(OUTPUT_FINANCIAL, {
            "TOP_FINANCIAL_ANOMALIES": financial["SUSPICIOUS_PRICES"][:250],
            "MAGNITUDE_ERRORS": financial["MAGNITUDE_ERRORS"][:150],
            "PRICE_OUTLIERS": financial["PRICE_OUTLIERS"][:150],
        })
        save_workbook(OUTPUT_MAPPING, {
            "TOP_MAPPING_ERRORS": taxonomy["TOP_MAPPING_ERRORS"][:250],
            "COVERAGE_BY_FAMILY": taxonomy["COVERAGE_BY_FAMILY"],
            "COVERAGE_BY_SUBLOT": taxonomy["COVERAGE_BY_SUBLOT"],
        })
        print("\n=== Scores ===")
        for key, value in scores.items():
            print(f"{key}: {value}/100")
        print("\n=== Livrables ===")
        for path in [OUTPUT_MD, OUTPUT_JSON, OUTPUT_XLSX, OUTPUT_HEATMAP, OUTPUT_FINANCIAL, OUTPUT_MAPPING]:
            print(path.name)
        print(f"\nAudit complete in {time.perf_counter() - start:.2f}s")
        return 0
    except Exception:
        print("AUDIT FAILED")
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

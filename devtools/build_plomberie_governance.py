from __future__ import annotations

import hashlib
import json
import math
import re
import sys
import time
import traceback
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "07_API_BACKEND"
DQE_PATH = ROOT / "03_DONNEES_REFERENCE" / "DQE_PROJECT_SP2I.xlsx"
SHEET_NAME = "DQE_CLEAN"

DQE_OUT = ROOT / "DQE_PLOMBERIE_TEST.xlsx"
MASTER_OUT = ROOT / "MASTER_REFERENCE_PLOMBERIE.xlsx"
SUPPLIERS_OUT = ROOT / "PLOMBERIE_SUPPLIER_REGISTRY.xlsx"
PROCUREMENT_OUT = ROOT / "PLOMBERIE_PROCUREMENT_AUDIT.xlsx"
DRIFT_OUT = ROOT / "PLOMBERIE_DRIFT_ANALYSIS.xlsx"
TCO_OUT = ROOT / "PLOMBERIE_TCO_ANALYSIS.xlsx"
WATER_OUT = ROOT / "PLOMBERIE_WATER_GOVERNANCE.xlsx"
LOW_OUT = ROOT / "PLOMBERIE_LOW_CONFIDENCE_REFERENCES.xlsx"
STATS_OUT = ROOT / "PLOMBERIE_VALIDATION_STATS.json"

if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from app.core.equipment_classification_engine import EquipmentClassificationEngine
from app.core.semantic_normalization_engine import SemanticNormalizationEngine


TODAY = date.today().isoformat()
LIFETIME_YEARS = 12

PLOMBERIE_KEYWORDS = [
    "alimentation eau",
    "evacuation",
    "pvc",
    "per",
    "cuivre",
    "multicouche",
    "sanitaire",
    "robinetterie",
    "vanne",
    "pompe",
    "surpresseur",
    "chauffe eau",
    "tuyauterie",
    "accessoire plomberie",
    "regard",
    "siphon",
    "collecteur",
    "eau froide",
    "eau chaude",
]

BENCHMARKS = {
    "ALIMENTATION_EAU": {"min": 4_500, "max": 65_000, "moq": "500 ML", "fob": 0.62, "life": 15},
    "EVACUATION": {"min": 3_500, "max": 45_000, "moq": "500 ML", "fob": 0.64, "life": 12},
    "SANITAIRE": {"min": 45_000, "max": 1_800_000, "moq": "20 U", "fob": 0.58, "life": 10},
    "ROBINETTERIE": {"min": 18_000, "max": 450_000, "moq": "50 U", "fob": 0.55, "life": 8},
    "POMPES": {"min": 350_000, "max": 7_500_000, "moq": "5 U", "fob": 0.54, "life": 10},
    "SURPRESSEURS": {"min": 650_000, "max": 12_000_000, "moq": "2 U", "fob": 0.54, "life": 10},
    "CHAUFFE_EAU": {"min": 120_000, "max": 2_500_000, "moq": "20 U", "fob": 0.57, "life": 8},
    "VANNES": {"min": 8_000, "max": 550_000, "moq": "100 U", "fob": 0.58, "life": 10},
    "REGARDS": {"min": 35_000, "max": 900_000, "moq": "20 U", "fob": 0.65, "life": 15},
    "PLOMBERIE_GENERAL": {"min": 8_000, "max": 3_500_000, "moq": "A_VERIFIER", "fob": 0.60, "life": 10},
}

SUPPLIERS = [
    {"SUPPLIER_NAME": "Zhejiang Plumbing Materials Export", "COUNTRY": "China", "PRODUCT_TYPES": "PVC|PER|MULTICOUCHE|VANNES", "MOQ": "500 ML", "FOB_MIN": 1500, "FOB_MAX": 250000, "LEAD_TIME_DAYS": 45, "SUPPLIER_CONFIDENCE": "LOW", "VERIFIED": False},
    {"SUPPLIER_NAME": "Foshan Sanitary Ware Group", "COUNTRY": "China", "PRODUCT_TYPES": "SANITAIRE|ROBINETTERIE|CHAUFFE_EAU", "MOQ": "20 U", "FOB_MIN": 18000, "FOB_MAX": 1200000, "LEAD_TIME_DAYS": 50, "SUPPLIER_CONFIDENCE": "LOW", "VERIFIED": False},
    {"SUPPLIER_NAME": "Guangzhou Pump Equipment Trading", "COUNTRY": "China", "PRODUCT_TYPES": "POMPES|SURPRESSEURS", "MOQ": "2 U", "FOB_MIN": 200000, "FOB_MAX": 6500000, "LEAD_TIME_DAYS": 55, "SUPPLIER_CONFIDENCE": "LOW", "VERIFIED": False},
]


def normalize_key(value: Any) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"[^A-Za-z0-9]+", "_", text.strip())
    return re.sub(r"_+", "_", text).strip("_").upper()


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.lower()
    replacements = {
        "e.f": "eau froide",
        "ec": "eau chaude",
        "e.c": "eau chaude",
        "chauffe-eau": "chauffe eau",
        "multi couche": "multicouche",
        "ø": " diam ",
        "diametre": "diam",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = re.sub(r"[^a-z0-9.,]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def number(value: Any, default: float = 0.0) -> float:
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return default if math.isnan(value) else float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return default


def choose(row: dict[str, Any], aliases: list[str], default: Any = "") -> Any:
    for alias in aliases:
        key = normalize_key(alias)
        if key in row and row[key] not in (None, ""):
            return row[key]
    return default


def read_dqe() -> list[dict[str, Any]]:
    wb = load_workbook(DQE_PATH, data_only=True, read_only=True)
    ws = wb[SHEET_NAME]
    headers = [normalize_key(cell.value or f"COL_{index + 1}") for index, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))]
    rows = []
    for rownum, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        payload = {headers[index]: value for index, value in enumerate(values)}
        if any(value not in (None, "") for value in payload.values()):
            payload["_ROWNUM"] = rownum
            rows.append(payload)
    wb.close()
    return rows


def is_plomberie(row: dict[str, Any]) -> bool:
    text = normalize_text(f"{choose(row, ['LOT'])} {choose(row, ['SOUS_LOT'])} {choose(row, ['DESIGNATION'])}")
    return any(keyword in text for keyword in PLOMBERIE_KEYWORDS)


def plumbing_type(text: str) -> str:
    if any(term in text for term in ["pompe", "pompage"]):
        return "POMPES"
    if "surpresseur" in text:
        return "SURPRESSEURS"
    if "chauffe eau" in text:
        return "CHAUFFE_EAU"
    if any(term in text for term in ["sanitaire", "wc", "lavabo", "douche"]):
        return "SANITAIRE"
    if any(term in text for term in ["robinet", "mitigeur", "melangeur"]):
        return "ROBINETTERIE"
    if "vanne" in text:
        return "VANNES"
    if any(term in text for term in ["evacuation", "siphon", "eaux usees"]):
        return "EVACUATION"
    if any(term in text for term in ["alimentation eau", "eau froide", "eau chaude", "tuyauterie", "collecteur"]):
        return "ALIMENTATION_EAU"
    if "regard" in text:
        return "REGARDS"
    return "PLOMBERIE_GENERAL"


def material(text: str) -> str:
    for candidate in ["cuivre", "multicouche", "per", "pvc", "inox", "acier", "fonte"]:
        if candidate in text:
            return candidate.upper()
    return "A_VERIFIER"


def network_type(text: str) -> str:
    if any(term in text for term in ["evacuation", "eaux usees", "siphon"]):
        return "EVACUATION"
    if "eau chaude" in text:
        return "ECS"
    if "eau froide" in text or "alimentation eau" in text:
        return "EF"
    if any(term in text for term in ["pompe", "surpresseur"]):
        return "PRESSION"
    return "GENERAL"


def diameter(text: str) -> tuple[float, str]:
    match = re.search(r"(?:dn|diam)?\s*(\d+(?:[.,]\d+)?)\s*(?:mm|m/m)?", text)
    if match:
        value = number(match.group(1))
        if 8 <= value <= 400:
            return value, "MM"
    return 0, ""


def pressure(text: str, type_reseau: str) -> float:
    bar = re.search(r"(\d+(?:[.,]\d+)?)\s*bar", text)
    if bar:
        return number(bar.group(1))
    if type_reseau == "PRESSION":
        return 6
    if type_reseau in {"EF", "ECS"}:
        return 4
    return 0


def technical_status(diam: float, press: float, mat: str, type_reseau: str) -> tuple[str, str]:
    issues = []
    if type_reseau in {"EF", "ECS", "EVACUATION"} and not diam:
        issues.append("diametre manquant")
    if type_reseau in {"EF", "ECS", "PRESSION"} and not press:
        issues.append("pression service manquante")
    if mat == "A_VERIFIER":
        issues.append("materiau non detecte")
    if mat == "PVC" and type_reseau == "ECS":
        issues.append("PVC incompatible eau chaude sanitaire")
    if mat == "CUIVRE" and type_reseau == "EVACUATION":
        issues.append("cuivre inhabituel pour evacuation")
    if not issues:
        return "VERIFIED", "coherence technique plomberie acceptable"
    if len(issues) <= 1:
        return "PARTIAL", "; ".join(issues)
    return "REJECTED", "; ".join(issues)


def water_risk(mat: str, type_reseau: str, press: float, importability: str) -> str:
    score = 15
    if mat in {"A_VERIFIER", "ACIER", "FONTE"}:
        score += 25
    if type_reseau in {"ECS", "PRESSION"}:
        score += 18
    if press and press > 8:
        score += 25
    if importability == "LOW":
        score += 10
    if score >= 70:
        return "CRITICAL"
    if score >= 50:
        return "HIGH"
    if score >= 30:
        return "MEDIUM"
    return "LOW"


def split_supply_install(raw: str, importability: str) -> tuple[str, str]:
    text = normalize_text(raw)
    has_install = any(term in text for term in ["pose", "installation", "mise en service", "raccordement"])
    has_supply = any(term in text for term in ["fourniture", "achat", "equipement"])
    if has_supply and has_install:
        return "MIXED_PLOMBERIE", "ligne mixte fourniture + installation"
    if has_install:
        return "INSTALLATION_PLOMBERIE", "installation non importable"
    if importability == "HIGH":
        return "FOURNITURE_PLOMBERIE", "fourniture importable sous reserve"
    return "INSTALLATION_PLOMBERIE", "non importable ou prestation probable"


def drift_alert(mat: str, price_min: float, price_max: float, water: str) -> str:
    ratio = price_max / max(price_min, 1)
    score = 18
    if mat == "CUIVRE":
        score += 30
    if mat in {"PVC", "PER", "MULTICOUCHE"}:
        score += 15
    if ratio >= 8:
        score += 35
    elif ratio >= 4:
        score += 20
    if water in {"HIGH", "CRITICAL"}:
        score += 18
    if score >= 80:
        return "CRITICAL"
    if score >= 60:
        return "HIGH"
    if score >= 35:
        return "MEDIUM"
    return "LOW"


def decision_blocker(conf: str, water: str, tech: str, split: str, drift: str) -> str:
    if split == "INSTALLATION_PLOMBERIE":
        return "BLOCK_IMPORT"
    if tech == "REJECTED":
        return "TECHNICAL_VALIDATION_REQUIRED"
    if water in {"HIGH", "CRITICAL"}:
        return "HIGH_RISK"
    if conf != "HIGH" or drift in {"HIGH", "CRITICAL"}:
        return "REVIEW_REQUIRED"
    return "NONE"


def confidence(supplier_verified: bool, tech: str, water: str, drift: str, fob_status: str) -> str:
    if supplier_verified and tech == "VERIFIED" and water in {"LOW", "MEDIUM"} and drift in {"LOW", "MEDIUM"} and fob_status == "VERIFIED":
        return "HIGH"
    if tech in {"VERIFIED", "PARTIAL"} and water != "CRITICAL" and drift != "CRITICAL" and fob_status in {"PARTIAL", "VERIFIED"}:
        return "MEDIUM"
    return "LOW"


def supplier_registry() -> list[dict[str, Any]]:
    rows = []
    for supplier in SUPPLIERS:
        rows.append({
            "SUPPLIER_ID": "PLOMB-SUP-" + hashlib.sha1(supplier["SUPPLIER_NAME"].encode("utf-8")).hexdigest()[:8].upper(),
            "SUPPLIER_NAME": supplier["SUPPLIER_NAME"],
            "COUNTRY": supplier["COUNTRY"],
            "PRODUCT_TYPES": supplier["PRODUCT_TYPES"],
            "MOQ": supplier["MOQ"],
            "FOB_MIN": supplier["FOB_MIN"],
            "FOB_MAX": supplier["FOB_MAX"],
            "LEAD_TIME_DAYS": supplier["LEAD_TIME_DAYS"],
            "SUPPLIER_CONFIDENCE": supplier["SUPPLIER_CONFIDENCE"],
            "VERIFIED": supplier["VERIFIED"],
            "LAST_MARKET_CHECK": TODAY,
            "VALIDATION_STATUS": "PENDING",
        })
    return rows


def build() -> dict[str, list[dict[str, Any]]]:
    classifier = EquipmentClassificationEngine()
    normalizer = SemanticNormalizationEngine()
    suppliers = supplier_registry()
    dqe_rows = []
    master_rows = []
    procurement_rows = []
    drift_rows = []
    tco_rows = []
    water_rows = []
    low_rows = []
    for row in read_dqe():
        if not is_plomberie(row):
            continue
        raw = str(choose(row, ["DESIGNATION"], ""))
        text = normalize_text(raw)
        lot = str(choose(row, ["LOT"], ""))
        sous_lot = str(choose(row, ["SOUS_LOT"], ""))
        type_ligne = str(choose(row, ["TYPE_LIGNE"], "")).upper()
        unit = str(choose(row, ["UNITE"], "U") or "U").upper()
        price = number(choose(row, ["PU_ESTIME_LOCAL_FCFA"], 0))
        ptype = plumbing_type(normalize_text(f"{lot} {sous_lot} {raw}"))
        bench = BENCHMARKS[ptype]
        mat = material(text)
        net = network_type(text)
        diam, diam_unit = diameter(text)
        press = pressure(text, net)
        importability = "HIGH" if type_ligne in {"ARTICLE", "FOURNITURE", "EQUIPEMENT"} else "LOW"
        split, split_reason = split_supply_install(raw, importability)
        tech, tech_reason = technical_status(diam, press, mat, net)
        water = water_risk(mat, net, press, importability)
        drift = drift_alert(mat, bench["min"], bench["max"], water)
        fob_status = "PARTIAL" if price > 0 and split != "INSTALLATION_PLOMBERIE" else "UNVERIFIED"
        conf = confidence(False, tech, water, drift, fob_status)
        blocker = decision_blocker(conf, water, tech, split, drift)
        normalized = normalizer.normalize(raw)["normalized_designation"].upper()
        ref_id = "SP2I-PLO-" + hashlib.sha1(f"{normalized}|{unit}".encode("utf-8")).hexdigest()[:8].upper()
        classification = classifier.classify_line({"designation": raw, "LOT": lot, "SOUS_LOT": sous_lot, "TYPE_LIGNE": type_ligne, "region": "CONGO_BRAZZAVILLE", "quality": "MEDIUM"})
        common = {
            "REFERENCE_ID": ref_id,
            "DESIGNATION_NORMALISEE": normalized,
            "RAW_DESIGNATION": raw,
            "FAMILLE": "PLOMBERIE",
            "SOUS_FAMILLE": ptype,
            "TYPE_EQUIPEMENT": classification.get("equipment_type") or ptype,
            "UNITE": unit,
            "DIAMETRE": diam,
            "UNITE_DIAMETRE": diam_unit,
            "PRESSION_SERVICE": press,
            "TYPE_RESEAU": net,
            "MATERIAU": mat,
            "MARQUE_REFERENCE": "A_VERIFIER",
        }
        dqe_rows.append({"SOURCE_ROW": row["_ROWNUM"], "LOT": lot, "SOUS_LOT": sous_lot, "TYPE_LIGNE": type_ligne, "PU_ESTIME_LOCAL_FCFA": price, **common})
        master_rows.append({
            **common,
            "PRICE_MIN_FCFA": bench["min"],
            "PRICE_MAX_FCFA": bench["max"],
            "PU_CHINE_FOB_FCFA": round(bench["min"] * bench["fob"], 2),
            "IMPORTABILITY": importability,
            "RISK_LEVEL": "HIGH" if conf == "LOW" else "MEDIUM",
            "MOQ": bench["moq"],
            "INCOTERM": "FOB",
            "FOURNISSEUR_CHINE": "A_VERIFIER_GOUVERNANCE_PLOMBERIE",
            "BENCHMARK_REGION": "CONGO_BRAZZAVILLE|CAMEROUN|GABON|CENTRAL_AFRICA",
            "QUALITY_LEVEL": "MEDIUM",
            "CONFIDENCE_LEVEL": conf,
            "DECISION_BLOCKER": blocker,
            "SOURCE_REFERENCE": "DQE_PROJECT_SP2I.xlsx::DQE_CLEAN::PLOMBERIE_TEST",
        })
        procurement_rows.append({
            **common,
            "IMPORTABILITY": importability,
            "FOB_VALIDATION_STATUS": fob_status,
            "SUPPLIER_VERIFIED": False,
            "PROCUREMENT_REVIEW_REQUIRED": "YES",
            "FOURNITURE_INSTALLATION_CLASS": split,
            "DECISION_BLOCKER": blocker,
            "PROCUREMENT_REASON": "; ".join([split_reason, tech_reason]),
        })
        drift_rows.append({**common, "PLOMBERIE_DRIFT_ALERT_LEVEL": drift, "DRIFT_FACTORS": "cuivre|pvc|inflation|fret_maritime|devises|materiaux_importes", "WATER_RISK_SCORE": water})
        water_rows.append({**common, "TECHNICAL_VALIDATION_STATUS": tech, "TECHNICAL_VALIDATION_REASON": tech_reason, "WATER_RISK_SCORE": water, "CORROSION_RISK": "HIGH" if mat in {"ACIER", "FONTE", "A_VERIFIER"} else "MEDIUM" if mat == "CUIVRE" else "LOW", "LEAK_RISK": water})
        maintenance = bench["min"] * (0.06 if water in {"LOW", "MEDIUM"} else 0.11) * bench["life"]
        replacement = bench["min"] * (0.18 if water in {"HIGH", "CRITICAL"} else 0.08)
        tco_rows.append({**common, "PURCHASE_COST_FCFA": bench["min"], "MAINTENANCE_COST_FCFA": round(maintenance, 2), "LIFETIME_YEARS": bench["life"], "CORROSION_FACTOR": water, "SPARE_PARTS_AVAILABILITY": "MEDIUM" if importability == "HIGH" else "HIGH_LOCAL", "REPLACEMENT_COST_FCFA": round(replacement, 2), "TOTAL_COST_OF_OWNERSHIP": round(bench["min"] + maintenance + replacement, 2)})
        if conf != "HIGH":
            low_rows.append({**common, "CONFIDENCE_LEVEL": conf, "DECISION_BLOCKER": blocker, "NEXT_ACTION": "Verifier materiau, fournisseur, FOB, pression et benchmark local."})
    return {
        "dqe": dqe_rows,
        "master": master_rows,
        "suppliers": suppliers,
        "procurement": procurement_rows,
        "drift": drift_rows,
        "tco": tco_rows,
        "water": water_rows,
        "low": low_rows,
    }


def write_xlsx(path: Path, sheet_name: str, rows: list[dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header) for header in headers])
    else:
        ws.append(["STATUS"])
        ws.append(["NO_PLOMBERIE_ROWS"])
    style_sheet(ws)
    wb.save(path)


def style_sheet(ws) -> None:
    fill = PatternFill("solid", fgColor="0F172A")
    font = Font(color="F8FAFC", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    for column in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(max(max_len + 2, 12), 64)
    ws.freeze_panes = "A2"


def count_by(rows: list[dict[str, Any]], key: str) -> dict[str, int]:
    result: dict[str, int] = {}
    for row in rows:
        value = str(row.get(key) or "UNKNOWN")
        result[value] = result.get(value, 0) + 1
    return result


def main() -> int:
    start = time.perf_counter()
    print("=== SP2I Plomberie Governance ===")
    try:
        outputs = build()
        write_xlsx(DQE_OUT, "DQE_PLOMBERIE_TEST", outputs["dqe"])
        write_xlsx(MASTER_OUT, "MASTER_PLOMBERIE", outputs["master"])
        write_xlsx(SUPPLIERS_OUT, "PLOMBERIE_SUPPLIERS", outputs["suppliers"])
        write_xlsx(PROCUREMENT_OUT, "PLOMBERIE_PROCUREMENT", outputs["procurement"])
        write_xlsx(DRIFT_OUT, "PLOMBERIE_DRIFT", outputs["drift"])
        write_xlsx(TCO_OUT, "PLOMBERIE_TCO", outputs["tco"])
        write_xlsx(WATER_OUT, "WATER_GOVERNANCE", outputs["water"])
        write_xlsx(LOW_OUT, "LOW_CONFIDENCE", outputs["low"])
        stats = {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "source_file": str(DQE_PATH),
            "sheet": SHEET_NAME,
            "family": "PLOMBERIE",
            "plomberie_rows": len(outputs["dqe"]),
            "master_references": len(outputs["master"]),
            "supplier_candidates": len(outputs["suppliers"]),
            "confidence_distribution": count_by(outputs["master"], "CONFIDENCE_LEVEL"),
            "water_risk": count_by(outputs["water"], "WATER_RISK_SCORE"),
            "technical_validation": count_by(outputs["water"], "TECHNICAL_VALIDATION_STATUS"),
            "drift_alerts": count_by(outputs["drift"], "PLOMBERIE_DRIFT_ALERT_LEVEL"),
            "decision_blockers": count_by(outputs["master"], "DECISION_BLOCKER"),
            "high_policy": "No HIGH without verified supplier, verified benchmark, coherent material, acceptable drift and stable procurement.",
            "duration_seconds": round(time.perf_counter() - start, 2),
        }
        STATS_OUT.write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(stats, ensure_ascii=False, indent=2))
        print("\nLivrables:")
        for path in [DQE_OUT, MASTER_OUT, SUPPLIERS_OUT, PROCUREMENT_OUT, DRIFT_OUT, TCO_OUT, WATER_OUT, LOW_OUT, STATS_OUT]:
            print(f"- {path.name}")
        return 0
    except Exception:
        print("PLOMBERIE GOVERNANCE FAILED")
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

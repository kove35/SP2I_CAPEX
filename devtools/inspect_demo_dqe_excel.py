from __future__ import annotations

import json
import sys
import time
import traceback
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "07_API_BACKEND"
SOURCE = ROOT / "03_DONNEES_REFERENCE" / "SP2I_CAPEX_DEMO_V1.xlsx"

if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from app.services.service_ai_mapping import ServiceAIMapping  # noqa: E402


def section(title: str) -> None:
    print("\n" + "=" * 88)
    print(title)
    print("=" * 88)


def non_empty_values(row: tuple[Any, ...]) -> list[Any]:
    return [value for value in row if value not in (None, "")]


def workbook_diagnostics() -> None:
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    print(f"Workbook: {SOURCE}")
    print(f"Sheets: {', '.join(workbook.sheetnames)}")

    for worksheet in workbook.worksheets:
        print(f"\n- {worksheet.title}: {worksheet.max_row} rows x {worksheet.max_column} columns")
        header_printed = False
        candidate_rows = 0
        for row_index, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
            values = non_empty_values(row)
            if len(values) < 3:
                continue
            candidate_rows += 1
            if not header_printed:
                print(f"  first non-empty row #{row_index}: {values[:20]}")
                header_printed = True
            if candidate_rows >= 5:
                break
        print(f"  non-empty preview rows: {candidate_rows}")


def certification_from_governance(governance: dict[str, Any]) -> str:
    if int(governance.get("blocking_loss_rows") or 0) > 0:
        return "BLOCKED"
    if int(governance.get("review_required_rows") or 0) > 0:
        return "REVIEW_REQUIRED"
    if int(governance.get("warning_rows") or 0) > 0:
        return "CERTIFIED_WITH_WARNINGS"
    return "CERTIFIED"


def main() -> int:
    started = time.perf_counter()
    section("DEVTOOLS | Inspect demo DQE Excel")

    if not SOURCE.exists():
        print("[FAIL] Demo DQE Excel file not found.")
        return 1

    try:
        workbook_diagnostics()

        section("Normalized parser result")
        service = ServiceAIMapping()
        lines, audit = service.extraire_lignes_normalisees(SOURCE.read_bytes(), SOURCE.name)
        confidence = audit.get("ai_confidence") or {}
        governance = confidence.get("governance_quality") or {}
        issues = audit.get("dqe_issues") or []
        issues_summary = audit.get("dqe_issues_summary") or {}
        bim_maturity = audit.get("bim_maturity") or {}
        analysis = audit.get("analyse") or {}
        sheet_selection = audit.get("sheet_selection") or {}
        certification_status = certification_from_governance(governance)

        print(f"recommended_sheet: {audit.get('feuille_recommandee')}")
        print(f"header_line: {analysis.get('ligne_entete')}")
        print(f"source_fact_metre: {sheet_selection.get('source_fact_metre')}")
        print("mapping:")
        for item in analysis.get("mapping", []):
            print(
                "  - "
                f"{item.get('champ_standard')} <- col {item.get('colonne_index')} "
                f"({item.get('colonne_excel')}) confidence={item.get('confiance')}"
            )

        print(f"normalized_lines_count: {len(lines)}")
        print(f"ignored_lines_count: {governance.get('ignored_rows', 0)}")
        print(f"data_loss_count: {governance.get('blocking_loss_rows', 0)}")
        print(f"review_required_count: {governance.get('review_required_rows', 0)}")
        print(f"trust_score: {confidence.get('trust_score', '-')}")
        print(f"certification_status: {certification_status}")
        print("issues_summary:")
        print(json.dumps(issues_summary, ensure_ascii=False, indent=2, default=str))
        print("bim_maturity:")
        print(json.dumps(bim_maturity, ensure_ascii=False, indent=2, default=str))

        section("Detected DQE issues")
        if not issues:
            print("No issue detected.")
        for issue in issues[:20]:
            print(
                f"- {issue.get('issue_type')} | severity={issue.get('severity')} "
                f"| blocking={issue.get('blocking')} | sheet={issue.get('sheet_name')} "
                f"| line={issue.get('line_number')}"
            )
            print(f"  message: {issue.get('message')}")
            print(f"  manual_fix: {issue.get('manual_fix')}")

        section("Five normalized examples")
        print(json.dumps(lines[:5], ensure_ascii=False, indent=2, default=str))

        if not lines:
            print("[FAIL] Parser returned zero normalized lines for the demo DQE.")
            return 1
        if int(governance.get("blocking_loss_rows") or 0) > 0:
            print("[FAIL] Strict data-loss rows detected.")
            return 1

        elapsed = round((time.perf_counter() - started) * 1000, 1)
        print(f"\nSummary: OK ({elapsed} ms)")
        return 0
    except Exception:
        print("[FAIL] Demo DQE inspection failed.")
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

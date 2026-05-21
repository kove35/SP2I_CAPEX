from __future__ import annotations

import hashlib
import json
import math
import re
import sys
import time
import traceback
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "07_API_BACKEND"
DQE_PATH = ROOT / "03_DONNEES_REFERENCE" / "DQE_PROJECT_SP2I.xlsx"
SHEET_NAME = "DQE_CLEAN"

DQE_OUT = ROOT / "DQE_MENUISERIE_ALU_TEST.xlsx"
MASTER_OUT = ROOT / "MASTER_REFERENCE_MENUISERIE.xlsx"
SUPPLIERS_OUT = ROOT / "MENUISERIE_SUPPLIER_REGISTRY.xlsx"
PROCUREMENT_OUT = ROOT / "MENUISERIE_PROCUREMENT_AUDIT.xlsx"
DRIFT_OUT = ROOT / "MENUISERIE_DRIFT_ANALYSIS.xlsx"
TCO_OUT = ROOT / "MENUISERIE_TCO_ANALYSIS.xlsx"
FACADE_OUT = ROOT / "MENUISERIE_FACADE_GOVERNANCE.xlsx"
LOW_OUT = ROOT / "MENUISERIE_LOW_CONFIDENCE_REFERENCES.xlsx"
STATS_OUT = ROOT / "MENUISERIE_VALIDATION_STATS.json"

if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from app.core.equipment_classification_engine import EquipmentClassificationEngine
from app.core.semantic_normalization_engine import SemanticNormalizationEngine


TODAY = date.today().isoformat()

KEYWORDS = [
    "fenetre",
    "fenêtre",
    "porte aluminium",
    "coulissant",
    "battant",
    "chassis",
    "châssis",
    "facade rideau",
    "façade rideau",
    "vitrage",
    "double vitrage",
    "quincaillerie",
    "profile aluminium",
    "profilé aluminium",
    "accessoire aluminium",
    "etancheite facade",
    "étanchéité façade",
    "garde corps",
    "alu",
    "aluminium",
    "baie",
    "mur rideau",
]

BENCHMARKS = {
    "FENETRE_ALUMINIUM": {"min": 85_000, "max": 280_000, "moq": "100 M2", "fob": 0.58, "life": 20},
    "PORTE_ALUMINIUM": {"min": 180_000, "max": 950_000, "moq": "20 U", "fob": 0.56, "life": 18},
    "COULISSANT_ALUMINIUM": {"min": 120_000, "max": 420_000, "moq": "100 M2", "fob": 0.57, "life": 18},
    "FACADE_RIDEAU": {"min": 180_000, "max": 650_000, "moq": "200 M2", "fob": 0.55, "life": 25},
    "VITRAGE": {"min": 55_000, "max": 240_000, "moq": "200 M2", "fob": 0.60, "life": 15},
    "GARDE_CORPS_VERRE_ALU": {"min": 95_000, "max": 360_000, "moq": "100 ML", "fob": 0.58, "life": 20},
    "QUINCAILLERIE_ALU": {"min": 8_000, "max": 120_000, "moq": "200 U", "fob": 0.54, "life": 10},
    "MENUISERIE_ALU_GENERAL": {"min": 65_000, "max": 450_000, "moq": "A_VERIFIER", "fob": 0.58, "life": 18},
}

SUPPLIERS = [
    {"SUPPLIER_NAME": "Foshan Aluminium Systems Export", "COUNTRY": "China", "PRODUCT_TYPES": "FENETRE_ALUMINIUM|COULISSANT_ALUMINIUM|PORTE_ALUMINIUM", "MOQ": "100 M2", "FOB_MIN": 45000, "FOB_MAX": 260000, "LEAD_TIME_DAYS": 55, "SUPPLIER_CONFIDENCE": "LOW", "VERIFIED": False},
    {"SUPPLIER_NAME": "Guangzhou Curtain Wall Materials", "COUNTRY": "China", "PRODUCT_TYPES": "FACADE_RIDEAU|VITRAGE|QUINCAILLERIE_ALU", "MOQ": "200 M2", "FOB_MIN": 35000, "FOB_MAX": 420000, "LEAD_TIME_DAYS": 65, "SUPPLIER_CONFIDENCE": "LOW", "VERIFIED": False},
    {"SUPPLIER_NAME": "Shenzhen Glass Hardware Trading", "COUNTRY": "China", "PRODUCT_TYPES": "VITRAGE|GARDE_CORPS_VERRE_ALU|QUINCAILLERIE_ALU", "MOQ": "100 M2", "FOB_MIN": 25000, "FOB_MAX": 300000, "LEAD_TIME_DAYS": 60, "SUPPLIER_CONFIDENCE": "LOW", "VERIFIED": False},
]


def normalize_key(value: Any) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"[^A-Za-z0-9]+", "_", text.strip())
    return re.sub(r"_+", "_", text).strip("_").upper()


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.lower()
    replacements = {
        "fenêtre": "fenetre",
        "façade": "facade",
        "châssis": "chassis",
        "profilé": "profile",
        "alu": "aluminium",
        "garde-corps": "garde corps",
        "mur-rideau": "mur rideau",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = re.sub(r"[^a-z0-9.,/]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def number(value: Any, default: float = 0.0) -> float:
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return default if math.isnan(value) else float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return default


def choose(row: dict[str, Any], aliases: list[str], default: Any = "") -> Any:
    for alias in aliases:
        key = normalize_key(alias)
        if key in row and row[key] not in (None, ""):
            return row[key]
    return default


def read_dqe() -> list[dict[str, Any]]:
    wb = load_workbook(DQE_PATH, data_only=True, read_only=True)
    ws = wb[SHEET_NAME]
    headers = [normalize_key(cell.value or f"COL_{index + 1}") for index, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))]
    rows = []
    for rownum, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        payload = {headers[index]: value for index, value in enumerate(values)}
        if any(value not in (None, "") for value in payload.values()):
            payload["_ROWNUM"] = rownum
            rows.append(payload)
    wb.close()
    return rows


def is_menuiserie(row: dict[str, Any]) -> bool:
    text = normalize_text(f"{choose(row, ['LOT'])} {choose(row, ['SOUS_LOT'])} {choose(row, ['DESIGNATION'])}")
    return any(keyword in text for keyword in KEYWORDS)


def menuiserie_type(text: str) -> str:
    if "facade rideau" in text or "mur rideau" in text:
        return "FACADE_RIDEAU"
    if "garde corps" in text:
        return "GARDE_CORPS_VERRE_ALU"
    if "quincaillerie" in text:
        return "QUINCAILLERIE_ALU"
    if "vitrage" in text and "fenetre" not in text and "baie" not in text:
        return "VITRAGE"
    if "porte" in text:
        return "PORTE_ALUMINIUM"
    if "coulissant" in text or "baie" in text:
        return "COULISSANT_ALUMINIUM"
    if "fenetre" in text or "chassis" in text:
        return "FENETRE_ALUMINIUM"
    return "MENUISERIE_ALU_GENERAL"


def glazing_type(text: str) -> str:
    if "44.2" in text or "securite" in text:
        return "SECURITE_44_2"
    if "double vitrage" in text or "/16/" in text:
        return "DOUBLE_VITRAGE"
    if "vitrage" in text:
        return "SIMPLE_OR_UNKNOWN"
    return "NON_APPLICABLE"


def opening_system(text: str) -> str:
    if "coulissant" in text:
        return "COULISSANT"
    if "battant" in text:
        return "BATTANT"
    if "fixe" in text:
        return "FIXE"
    return "A_VERIFIER"


def thickness(text: str) -> float:
    match = re.search(r"(\d+(?:[.,]\d+)?)\s*mm", text)
    if match:
        return number(match.group(1))
    glass = re.search(r"(\d+(?:[.,]\d+)?)/(\d+(?:[.,]\d+)?)/(\d+(?:[.,]\d+)?)", text)
    if glass:
        return sum(number(part) for part in glass.groups())
    return 0


def dimensions(text: str) -> str:
    match = re.search(r"(\d+(?:[.,]\d+)?)\s*[x×]\s*(\d+(?:[.,]\d+)?)", text)
    return f"{match.group(1)}x{match.group(2)}" if match else "A_VERIFIER"


def facade_type(mtype: str, text: str) -> str:
    if mtype == "FACADE_RIDEAU":
        return "CURTAIN_WALL"
    if "facade" in text:
        return "FACADE_ELEMENT"
    return "MENUISERIE_STANDARD"


def finish(text: str) -> str:
    if "laque" in text:
        return "LAQUE"
    if "anodise" in text:
        return "ANODISE"
    if "alucobond" in text:
        return "COMPOSITE"
    return "A_VERIFIER"


def rupture_thermique(text: str) -> str:
    if "rupture thermique" in text or "rpt" in text:
        return "YES"
    return "NO"


def technical_validation(mtype: str, glass: str, thick: float, dims: str, rupture: str) -> tuple[str, str]:
    issues = []
    if mtype in {"FENETRE_ALUMINIUM", "COULISSANT_ALUMINIUM", "FACADE_RIDEAU"} and glass == "NON_APPLICABLE":
        issues.append("vitrage non detecte")
    if glass != "NON_APPLICABLE" and thick and (thick < 4 or thick > 80):
        issues.append("epaisseur vitrage incoherente")
    if dims == "A_VERIFIER":
        issues.append("dimensions absentes")
    if mtype == "FACADE_RIDEAU" and rupture == "NO":
        issues.append("rupture thermique a confirmer pour facade")
    if not issues:
        return "VERIFIED", "coherence technique menuiserie acceptable"
    if len(issues) <= 2:
        return "PARTIAL", "; ".join(issues)
    return "REJECTED", "; ".join(issues)


def facade_risk(mtype: str, glass: str, finish_value: str, rupture: str, tech: str) -> str:
    score = 20
    if mtype in {"FACADE_RIDEAU", "GARDE_CORPS_VERRE_ALU"}:
        score += 25
    if glass in {"SIMPLE_OR_UNKNOWN", "NON_APPLICABLE"}:
        score += 15
    if finish_value == "A_VERIFIER":
        score += 12
    if rupture == "NO" and mtype in {"FACADE_RIDEAU", "COULISSANT_ALUMINIUM"}:
        score += 15
    if tech == "REJECTED":
        score += 25
    elif tech == "PARTIAL":
        score += 10
    if score >= 80:
        return "CRITICAL"
    if score >= 60:
        return "HIGH"
    if score >= 35:
        return "MEDIUM"
    return "LOW"


def supply_pose(raw: str, importability: str) -> tuple[str, str]:
    text = normalize_text(raw)
    has_pose = any(term in text for term in ["pose", "montage", "installation"])
    has_supply = any(term in text for term in ["fourniture", "achat", "fabrication"])
    if has_pose and has_supply:
        return "MIXED_MENUISERIE", "ligne mixte fourniture + pose"
    if has_pose:
        return "POSE_MENUISERIE", "pose non importable"
    if importability == "HIGH":
        return "FOURNITURE_MENUISERIE", "fourniture importable sous reserve"
    return "POSE_MENUISERIE", "non importable ou prestation probable"


def drift_alert(mtype: str, bench: dict[str, Any], risk: str) -> str:
    ratio = bench["max"] / max(bench["min"], 1)
    score = 20
    if mtype in {"FACADE_RIDEAU", "VITRAGE", "GARDE_CORPS_VERRE_ALU"}:
        score += 20
    if ratio >= 6:
        score += 30
    elif ratio >= 3:
        score += 18
    if risk in {"HIGH", "CRITICAL"}:
        score += 18
    if score >= 80:
        return "CRITICAL"
    if score >= 60:
        return "HIGH"
    if score >= 35:
        return "MEDIUM"
    return "LOW"


def confidence(tech: str, risk: str, drift: str, supplier_verified: bool) -> str:
    if supplier_verified and tech == "VERIFIED" and risk in {"LOW", "MEDIUM"} and drift in {"LOW", "MEDIUM"}:
        return "HIGH"
    if tech in {"VERIFIED", "PARTIAL"} and risk != "CRITICAL" and drift != "CRITICAL":
        return "MEDIUM"
    return "LOW"


def blocker(conf: str, split: str, risk: str, tech: str, drift: str) -> str:
    if split == "POSE_MENUISERIE":
        return "BLOCK_IMPORT"
    if tech == "REJECTED":
        return "TECHNICAL_VALIDATION_REQUIRED"
    if risk in {"HIGH", "CRITICAL"}:
        return "HIGH_RISK"
    if conf != "HIGH" or drift in {"HIGH", "CRITICAL"}:
        return "REVIEW_REQUIRED"
    return "NONE"


def supplier_registry() -> list[dict[str, Any]]:
    rows = []
    for supplier in SUPPLIERS:
        rows.append({
            "SUPPLIER_ID": "MENU-SUP-" + hashlib.sha1(supplier["SUPPLIER_NAME"].encode("utf-8")).hexdigest()[:8].upper(),
            "SUPPLIER_NAME": supplier["SUPPLIER_NAME"],
            "COUNTRY": supplier["COUNTRY"],
            "PRODUCT_TYPES": supplier["PRODUCT_TYPES"],
            "MOQ": supplier["MOQ"],
            "FOB_MIN": supplier["FOB_MIN"],
            "FOB_MAX": supplier["FOB_MAX"],
            "LEAD_TIME_DAYS": supplier["LEAD_TIME_DAYS"],
            "SUPPLIER_CONFIDENCE": supplier["SUPPLIER_CONFIDENCE"],
            "VERIFIED": supplier["VERIFIED"],
            "LAST_MARKET_CHECK": TODAY,
            "VALIDATION_STATUS": "PENDING",
        })
    return rows


def build() -> dict[str, list[dict[str, Any]]]:
    classifier = EquipmentClassificationEngine()
    normalizer = SemanticNormalizationEngine()
    suppliers = supplier_registry()
    dqe_rows = []
    master_rows = []
    procurement_rows = []
    drift_rows = []
    tco_rows = []
    facade_rows = []
    low_rows = []
    for row in read_dqe():
        if not is_menuiserie(row):
            continue
        raw = str(choose(row, ["DESIGNATION"], ""))
        text = normalize_text(raw)
        lot = str(choose(row, ["LOT"], ""))
        sous_lot = str(choose(row, ["SOUS_LOT"], ""))
        type_ligne = str(choose(row, ["TYPE_LIGNE"], "")).upper()
        unit = str(choose(row, ["UNITE"], "U") or "U").upper()
        price = number(choose(row, ["PU_ESTIME_LOCAL_FCFA"], 0))
        mtype = menuiserie_type(normalize_text(f"{lot} {sous_lot} {raw}"))
        bench = BENCHMARKS[mtype]
        glass = glazing_type(text)
        thick = thickness(text)
        dims = dimensions(text)
        finish_value = finish(text)
        rupture = rupture_thermique(text)
        opening = opening_system(text)
        ftype = facade_type(mtype, text)
        importability = "HIGH" if type_ligne in {"ARTICLE", "FOURNITURE", "EQUIPEMENT"} else "LOW"
        split, split_reason = supply_pose(raw, importability)
        tech, tech_reason = technical_validation(mtype, glass, thick, dims, rupture)
        risk = facade_risk(mtype, glass, finish_value, rupture, tech)
        drift = drift_alert(mtype, bench, risk)
        conf = confidence(tech, risk, drift, False)
        decision = blocker(conf, split, risk, tech, drift)
        normalized = normalizer.normalize(raw)["normalized_designation"].upper()
        ref_id = "SP2I-MEN-" + hashlib.sha1(f"{normalized}|{unit}".encode("utf-8")).hexdigest()[:8].upper()
        classification = classifier.classify_line({"designation": raw, "LOT": lot, "SOUS_LOT": sous_lot, "TYPE_LIGNE": type_ligne, "region": "CONGO_BRAZZAVILLE", "quality": "MEDIUM"})
        common = {
            "REFERENCE_ID": ref_id,
            "DESIGNATION_NORMALISEE": normalized,
            "RAW_DESIGNATION": raw,
            "FAMILLE": "MENUISERIE_ALUMINIUM",
            "SOUS_FAMILLE": mtype,
            "TYPE_EQUIPEMENT": classification.get("equipment_type") or mtype,
            "UNITE": unit,
            "TYPE_MENUISERIE": mtype,
            "EPAISSEUR_MM": thick,
            "TYPE_VITRAGE": glass,
            "FINITION": finish_value,
            "RUPTURE_THERMIQUE": rupture,
            "DIMENSIONS": dims,
            "SYSTEME_OUVERTURE": opening,
            "TYPE_FACADE": ftype,
        }
        dqe_rows.append({"SOURCE_ROW": row["_ROWNUM"], "LOT": lot, "SOUS_LOT": sous_lot, "TYPE_LIGNE": type_ligne, "PU_ESTIME_LOCAL_FCFA": price, **common})
        master_rows.append({
            **common,
            "PRICE_MIN_FCFA": bench["min"],
            "PRICE_MAX_FCFA": bench["max"],
            "PU_CHINE_FOB_FCFA": round(bench["min"] * bench["fob"], 2),
            "IMPORTABILITY": importability,
            "RISK_LEVEL": "HIGH" if conf == "LOW" else "MEDIUM",
            "MOQ": bench["moq"],
            "INCOTERM": "FOB",
            "FOURNISSEUR_CHINE": "A_VERIFIER_GOUVERNANCE_MENUISERIE",
            "BENCHMARK_REGION": "CONGO_BRAZZAVILLE|CAMEROUN|GABON|CENTRAL_AFRICA",
            "QUALITY_LEVEL": "MEDIUM",
            "CONFIDENCE_LEVEL": conf,
            "DECISION_BLOCKER": decision,
            "SOURCE_REFERENCE": "DQE_PROJECT_SP2I.xlsx::DQE_CLEAN::MENUISERIE_ALU_TEST",
        })
        procurement_rows.append({**common, "IMPORTABILITY": importability, "SUPPLIER_VERIFIED": False, "FOB_VALIDATION_STATUS": "PARTIAL" if split != "POSE_MENUISERIE" else "UNVERIFIED", "FOURNITURE_POSE_CLASS": split, "DECISION_BLOCKER": decision, "PROCUREMENT_REASON": "; ".join([split_reason, tech_reason])})
        drift_rows.append({**common, "MENUISERIE_DRIFT_ALERT_LEVEL": drift, "DRIFT_FACTORS": "aluminium_mondial|energie|verre|fret_maritime|usd|inflation_import", "FACADE_RISK_SCORE": risk})
        facade_rows.append({**common, "TECHNICAL_VALIDATION_STATUS": tech, "TECHNICAL_VALIDATION_REASON": tech_reason, "FACADE_RISK_SCORE": risk, "CORROSION_RISK": "MEDIUM" if finish_value == "A_VERIFIER" else "LOW", "ETANCHEITE_RISK": risk, "DURABILITY_YEARS": bench["life"]})
        maintenance = bench["min"] * (0.04 if risk in {"LOW", "MEDIUM"} else 0.08) * bench["life"]
        glass_replacement = bench["min"] * (0.20 if glass in {"DOUBLE_VITRAGE", "SECURITE_44_2"} else 0.10)
        tco_rows.append({**common, "PURCHASE_COST_FCFA": bench["min"], "MAINTENANCE_COST_FCFA": round(maintenance, 2), "CORROSION_FACTOR": risk, "ETANCHEITE_FACTOR": risk, "GLASS_REPLACEMENT_COST_FCFA": round(glass_replacement, 2), "LIFETIME_YEARS": bench["life"], "TOTAL_COST_OF_OWNERSHIP": round(bench["min"] + maintenance + glass_replacement, 2)})
        if conf != "HIGH":
            low_rows.append({**common, "CONFIDENCE_LEVEL": conf, "DECISION_BLOCKER": decision, "NEXT_ACTION": "Verifier fournisseur, vitrage, dimensions, rupture thermique, FOB et benchmark local."})
    return {"dqe": dqe_rows, "master": master_rows, "suppliers": suppliers, "procurement": procurement_rows, "drift": drift_rows, "tco": tco_rows, "facade": facade_rows, "low": low_rows}


def write_xlsx(path: Path, sheet_name: str, rows: list[dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header) for header in headers])
    else:
        ws.append(["STATUS"])
        ws.append(["NO_MENUISERIE_ALU_ROWS"])
    style_sheet(ws)
    wb.save(path)


def style_sheet(ws) -> None:
    fill = PatternFill("solid", fgColor="0F172A")
    font = Font(color="F8FAFC", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    for column in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(max(max_len + 2, 12), 64)
    ws.freeze_panes = "A2"


def count_by(rows: list[dict[str, Any]], key: str) -> dict[str, int]:
    result: dict[str, int] = {}
    for row in rows:
        value = str(row.get(key) or "UNKNOWN")
        result[value] = result.get(value, 0) + 1
    return result


def main() -> int:
    start = time.perf_counter()
    print("=== SP2I Menuiserie Aluminium Governance ===")
    try:
        outputs = build()
        write_xlsx(DQE_OUT, "DQE_MENUISERIE_ALU", outputs["dqe"])
        write_xlsx(MASTER_OUT, "MASTER_MENUISERIE", outputs["master"])
        write_xlsx(SUPPLIERS_OUT, "MENUISERIE_SUPPLIERS", outputs["suppliers"])
        write_xlsx(PROCUREMENT_OUT, "MENUISERIE_PROCUREMENT", outputs["procurement"])
        write_xlsx(DRIFT_OUT, "MENUISERIE_DRIFT", outputs["drift"])
        write_xlsx(TCO_OUT, "MENUISERIE_TCO", outputs["tco"])
        write_xlsx(FACADE_OUT, "FACADE_GOVERNANCE", outputs["facade"])
        write_xlsx(LOW_OUT, "LOW_CONFIDENCE", outputs["low"])
        stats = {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "source_file": str(DQE_PATH),
            "sheet": SHEET_NAME,
            "family": "MENUISERIE_ALUMINIUM",
            "menuiserie_rows": len(outputs["dqe"]),
            "master_references": len(outputs["master"]),
            "supplier_candidates": len(outputs["suppliers"]),
            "confidence_distribution": count_by(outputs["master"], "CONFIDENCE_LEVEL"),
            "facade_risk": count_by(outputs["facade"], "FACADE_RISK_SCORE"),
            "technical_validation": count_by(outputs["facade"], "TECHNICAL_VALIDATION_STATUS"),
            "drift_alerts": count_by(outputs["drift"], "MENUISERIE_DRIFT_ALERT_LEVEL"),
            "decision_blockers": count_by(outputs["master"], "DECISION_BLOCKER"),
            "high_policy": "No HIGH without verified supplier, verified benchmark, coherent glazing, acceptable drift and stable procurement.",
            "duration_seconds": round(time.perf_counter() - start, 2),
        }
        STATS_OUT.write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(stats, ensure_ascii=False, indent=2))
        print("\nLivrables:")
        for path in [DQE_OUT, MASTER_OUT, SUPPLIERS_OUT, PROCUREMENT_OUT, DRIFT_OUT, TCO_OUT, FACADE_OUT, LOW_OUT, STATS_OUT]:
            print(f"- {path.name}")
        return 0
    except Exception:
        print("MENUISERIE ALU GOVERNANCE FAILED")
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

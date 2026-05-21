from __future__ import annotations

import hashlib
import json
import math
import statistics
import time
import traceback
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
MASTER_V2 = ROOT / "MASTER_REFERENCE_GOVERNANCE_V2.xlsx"

SUPPLIER_REGISTRY = ROOT / "SUPPLIER_REFERENCE_REGISTRY.xlsx"
VERIFIED_PROCUREMENT = ROOT / "VERIFIED_PROCUREMENT_REFERENCES.xlsx"
FOB_AUDIT = ROOT / "FOB_VALIDATION_AUDIT.xlsx"
LOCAL_MARKET = ROOT / "LOCAL_MARKET_VALIDATION.xlsx"
CRITICAL_DRIFT = ROOT / "TOP_CRITICAL_DRIFT_REFERENCES.xlsx"
CONFIDENCE_EVOLUTION = ROOT / "PROCUREMENT_CONFIDENCE_EVOLUTION.xlsx"
STATS_PATH = ROOT / "SUPPLIER_GOVERNANCE_STATS.json"

TODAY = date.today().isoformat()
PHASE_FAMILY = "ELECTRICITE"

SUPPLIER_SEEDS = [
    {
        "SUPPLIER_NAME": "Shanghai FATO Group Co., Ltd.",
        "COUNTRY": "China",
        "CITY": "Shanghai",
        "PRODUCT_FAMILIES": "ELECTRICITE|TGBT|DISJONCTEURS|APPAREILLAGE",
        "MOQ": "100 U",
        "FOB_RANGE_MIN": 25000,
        "FOB_RANGE_MAX": 4500000,
        "INCOTERM": "FOB",
        "LEAD_TIME_DAYS": 45,
        "SUPPLIER_CONFIDENCE": "MEDIUM",
        "VERIFIED": False,
        "NOTES": "Candidat fournisseur a verifier par source marche et documentation commerciale.",
    },
    {
        "SUPPLIER_NAME": "Foshan Electrical Equipment Export Hub",
        "COUNTRY": "China",
        "CITY": "Foshan",
        "PRODUCT_FAMILIES": "ELECTRICITE|ECLAIRAGE|ACCESSOIRES",
        "MOQ": "200 U",
        "FOB_RANGE_MIN": 8000,
        "FOB_RANGE_MAX": 1200000,
        "INCOTERM": "FOB",
        "LEAD_TIME_DAYS": 50,
        "SUPPLIER_CONFIDENCE": "LOW",
        "VERIFIED": False,
        "NOTES": "Hub indicatif non valide comme fournisseur final avant controle humain.",
    },
    {
        "SUPPLIER_NAME": "Guangzhou Solar & Power Trading",
        "COUNTRY": "China",
        "CITY": "Guangzhou",
        "PRODUCT_FAMILIES": "ELECTRICITE|SOLAIRE|ONDULEURS|BATTERIES",
        "MOQ": "10 U",
        "FOB_RANGE_MIN": 150000,
        "FOB_RANGE_MAX": 8500000,
        "INCOTERM": "FOB",
        "LEAD_TIME_DAYS": 55,
        "SUPPLIER_CONFIDENCE": "LOW",
        "VERIFIED": False,
        "NOTES": "Candidat sourcing solaire a verifier, pas de HIGH automatique.",
    },
]


def number(value: Any, default: float = 0.0) -> float:
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return default if math.isnan(value) else float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return default


def supplier_id(name: str) -> str:
    digest = hashlib.sha1(name.encode("utf-8")).hexdigest()[:8].upper()
    return f"SUP-{digest}"


def read_sheet(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"Fichier introuvable: {path}")
    wb = load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Onglet {sheet_name} introuvable dans {path.name}")
    ws = wb[sheet_name]
    headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        payload = {headers[index]: value for index, value in enumerate(values)}
        if any(value not in (None, "") for value in payload.values()):
            rows.append(payload)
    wb.close()
    return rows


def supplier_registry_rows() -> list[dict[str, Any]]:
    rows = []
    for seed in SUPPLIER_SEEDS:
        rows.append({
            "SUPPLIER_ID": supplier_id(seed["SUPPLIER_NAME"]),
            "SUPPLIER_NAME": seed["SUPPLIER_NAME"],
            "COUNTRY": seed["COUNTRY"],
            "CITY": seed["CITY"],
            "PRODUCT_FAMILIES": seed["PRODUCT_FAMILIES"],
            "MOQ": seed["MOQ"],
            "FOB_RANGE_MIN": seed["FOB_RANGE_MIN"],
            "FOB_RANGE_MAX": seed["FOB_RANGE_MAX"],
            "INCOTERM": seed["INCOTERM"],
            "LEAD_TIME_DAYS": seed["LEAD_TIME_DAYS"],
            "SUPPLIER_CONFIDENCE": seed["SUPPLIER_CONFIDENCE"],
            "VERIFIED": seed["VERIFIED"],
            "LAST_MARKET_CHECK": TODAY,
            "LAST_VALIDATION_DATE": "",
            "VALIDATED_BY": "PENDING_HUMAN_VALIDATION",
            "NOTES": seed["NOTES"],
        })
    return rows


def candidate_supplier(row: dict[str, Any], suppliers: list[dict[str, Any]]) -> dict[str, Any] | None:
    subfamily = str(row.get("SOUS_FAMILLE") or "").upper()
    equipment = str(row.get("TYPE_EQUIPEMENT") or "").upper()
    target = f"{subfamily}|{equipment}|{row.get('DESIGNATION_NORMALISEE')}"
    scored = []
    for supplier in suppliers:
        product = str(supplier.get("PRODUCT_FAMILIES") or "").upper()
        score = sum(1 for token in product.split("|") if token and token in target)
        if score:
            scored.append((score, supplier))
    if not scored:
        return None
    return sorted(scored, key=lambda item: item[0], reverse=True)[0][1]


def fob_validation(row: dict[str, Any], supplier: dict[str, Any] | None) -> tuple[str, str]:
    fob = number(row.get("PU_CHINE_FOB_FCFA"))
    price_min = number(row.get("PRICE_MIN_FCFA"))
    if fob <= 0:
        return "UNVERIFIED", "FOB absent ou nul."
    if not supplier:
        return "UNVERIFIED", "Aucun fournisseur candidat associe."
    low = number(supplier.get("FOB_RANGE_MIN"))
    high = number(supplier.get("FOB_RANGE_MAX"))
    if fob < low * 0.5:
        return "REJECTED", "FOB tres inferieur au range fournisseur candidat."
    if fob > high * 1.8:
        return "REJECTED", "FOB tres superieur au range fournisseur candidat."
    if low <= fob <= high and price_min and fob <= price_min * 0.85:
        return "PARTIAL", "FOB plausible mais fournisseur non verifie humainement."
    return "PARTIAL", "FOB a confirmer par source marche."


def local_market_status(row: dict[str, Any]) -> tuple[str, str]:
    drift = number(row.get("MARKET_DRIFT_SCORE"))
    variation = str(row.get("PRICE_VARIATION_LEVEL") or "")
    if drift >= 80 or variation == "EXTREME":
        return "REJECTED", "Benchmark trop volatil ou derive critique."
    if drift >= 45 or variation == "HIGH":
        return "PARTIAL", "Benchmark utilisable uniquement avec warning."
    return "VERIFIED", "Benchmark local coherent sur la plage actuelle."


def drift_alert_level(row: dict[str, Any]) -> str:
    drift = number(row.get("MARKET_DRIFT_SCORE"))
    if drift >= 80:
        return "CRITICAL"
    if drift >= 60:
        return "HIGH"
    if drift >= 35:
        return "MEDIUM"
    return "LOW"


def confidence_after_validation(row: dict[str, Any], supplier: dict[str, Any] | None, fob_status: str, market_status: str) -> tuple[str, str]:
    supplier_verified = bool(supplier and supplier.get("VERIFIED") is True)
    taxonomy_ok = str(row.get("VALIDATION_STATUS")) in {"PARTIAL", "VERIFIED"}
    drift_ok = drift_alert_level(row) in {"LOW", "MEDIUM"}
    procurement_stable = str(row.get("PROCUREMENT_REVIEW_REQUIRED")) == "NO"
    if supplier_verified and fob_status == "VERIFIED" and market_status == "VERIFIED" and taxonomy_ok and drift_ok and procurement_stable:
        return "HIGH", "Toutes les validations reelles sont presentes."
    if fob_status in {"PARTIAL", "VERIFIED"} and market_status in {"PARTIAL", "VERIFIED"} and taxonomy_ok:
        return "MEDIUM", "Validation partielle: references utilisables avec revue procurement."
    return "LOW", "Validation fournisseur/FOB/benchmark insuffisante."


def cockpit_badge(confidence: str) -> str:
    return {"HIGH": "GREEN", "MEDIUM": "ORANGE", "LOW": "RED"}[confidence]


def build_phase3(master_rows: list[dict[str, Any]], suppliers: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    verified_rows = []
    fob_rows = []
    market_rows = []
    evolution_rows = []
    for row in master_rows:
        supplier = candidate_supplier(row, suppliers)
        fob_status, fob_reason = fob_validation(row, supplier)
        market_status, market_reason = local_market_status(row)
        confidence, confidence_reason = confidence_after_validation(row, supplier, fob_status, market_status)
        alert = drift_alert_level(row)
        supplier_verified = bool(supplier and supplier.get("VERIFIED") is True)
        supplier_name = supplier.get("SUPPLIER_NAME") if supplier else "NO_SUPPLIER_MATCH"
        common = {
            "REFERENCE_ID": row.get("REFERENCE_ID"),
            "DESIGNATION_NORMALISEE": row.get("DESIGNATION_NORMALISEE"),
            "FAMILLE": row.get("FAMILLE"),
            "SOUS_FAMILLE": row.get("SOUS_FAMILLE"),
            "TYPE_EQUIPEMENT": row.get("TYPE_EQUIPEMENT"),
            "SUPPLIER_NAME": supplier_name,
            "SUPPLIER_VERIFIED_REAL": supplier_verified,
            "FOB_VALIDATION_STATUS": fob_status,
            "LOCAL_MARKET_VALIDATION_STATUS": market_status,
            "DRIFT_ALERT_LEVEL": alert,
            "CONFIDENCE_BEFORE": row.get("CONFIDENCE_LEVEL"),
            "CONFIDENCE_AFTER": confidence,
            "COCKPIT_CONFIDENCE_BADGE": cockpit_badge(confidence),
        }
        verified_rows.append({
            **row,
            **common,
            "VALIDATION_EXPLANATION": confidence_reason,
            "PROCUREMENT_REVIEW_REQUIRED": "NO" if confidence == "HIGH" else "YES",
        })
        fob_rows.append({
            **common,
            "PU_CHINE_FOB_FCFA": row.get("PU_CHINE_FOB_FCFA"),
            "PRICE_MIN_FCFA": row.get("PRICE_MIN_FCFA"),
            "PRICE_MAX_FCFA": row.get("PRICE_MAX_FCFA"),
            "FOB_VALIDATION_REASON": fob_reason,
        })
        market_rows.append({
            **common,
            "PRICE_MIN_FCFA": row.get("PRICE_MIN_FCFA"),
            "PRICE_MAX_FCFA": row.get("PRICE_MAX_FCFA"),
            "PRICE_VARIATION_LEVEL": row.get("PRICE_VARIATION_LEVEL"),
            "MARKET_DRIFT_SCORE": row.get("MARKET_DRIFT_SCORE"),
            "LOCAL_MARKET_VALIDATION_REASON": market_reason,
        })
        evolution_rows.append({
            **common,
            "CAN_MOVE_TO_HIGH": "YES" if confidence == "HIGH" else "NO",
            "NEXT_REQUIRED_ACTION": next_action(supplier_verified, fob_status, market_status, alert),
            "CONFIDENCE_REASON": confidence_reason,
        })
    critical = sorted(
        [row for row in market_rows if row["DRIFT_ALERT_LEVEL"] in {"HIGH", "CRITICAL"}],
        key=lambda item: number(item.get("MARKET_DRIFT_SCORE")),
        reverse=True,
    )
    low_suppliers = [
        row for row in verified_rows
        if row["SUPPLIER_VERIFIED_REAL"] is False or row["CONFIDENCE_AFTER"] == "LOW"
    ]
    return {
        "verified": verified_rows,
        "fob": fob_rows,
        "market": market_rows,
        "critical": critical,
        "evolution": evolution_rows,
        "low_suppliers": low_suppliers,
    }


def next_action(supplier_verified: bool, fob_status: str, market_status: str, alert: str) -> str:
    if not supplier_verified:
        return "Verifier fournisseur reel et documentation commerciale."
    if fob_status != "VERIFIED":
        return "Confirmer FOB par source marche."
    if market_status != "VERIFIED":
        return "Confirmer benchmark local Congo/Cameroun/Gabon."
    if alert in {"HIGH", "CRITICAL"}:
        return "Revoir drift marche avant decision ROI."
    return "Reference eligible validation HIGH."


def write_xlsx(path: Path, sheet_name: str, rows: list[dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header) for header in headers])
    else:
        ws.append(["STATUS"])
        ws.append(["NO_DATA"])
    style_sheet(ws)
    wb.save(path)


def style_sheet(ws) -> None:
    fill = PatternFill("solid", fgColor="0F172A")
    font = Font(color="F8FAFC", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    for column in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(max(max_len + 2, 12), 64)
    ws.freeze_panes = "A2"


def count_by(rows: list[dict[str, Any]], key: str) -> dict[str, int]:
    result: dict[str, int] = {}
    for row in rows:
        value = str(row.get(key) or "UNKNOWN")
        result[value] = result.get(value, 0) + 1
    return result


def main() -> int:
    start = time.perf_counter()
    print("=== SP2I Supplier Governance Phase 3 ===")
    try:
        master_rows = read_sheet(MASTER_V2, "MASTER_GOVERNANCE_V2")
        suppliers = supplier_registry_rows()
        outputs = build_phase3(master_rows, suppliers)
        write_xlsx(SUPPLIER_REGISTRY, "SUPPLIER_REGISTRY", suppliers)
        write_xlsx(VERIFIED_PROCUREMENT, "VERIFIED_REFERENCES", outputs["verified"])
        write_xlsx(FOB_AUDIT, "FOB_VALIDATION", outputs["fob"])
        write_xlsx(LOCAL_MARKET, "LOCAL_MARKET", outputs["market"])
        write_xlsx(CRITICAL_DRIFT, "CRITICAL_DRIFT", outputs["critical"])
        write_xlsx(CONFIDENCE_EVOLUTION, "CONFIDENCE_EVOLUTION", outputs["evolution"])
        drift_scores = [number(row.get("MARKET_DRIFT_SCORE")) for row in outputs["market"]]
        stats = {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "source_file": str(MASTER_V2),
            "phase_family": PHASE_FAMILY,
            "references": len(master_rows),
            "suppliers_in_registry": len(suppliers),
            "verified_suppliers": len([supplier for supplier in suppliers if supplier["VERIFIED"] is True]),
            "confidence_after": count_by(outputs["verified"], "CONFIDENCE_AFTER"),
            "fob_validation": count_by(outputs["fob"], "FOB_VALIDATION_STATUS"),
            "local_market_validation": count_by(outputs["market"], "LOCAL_MARKET_VALIDATION_STATUS"),
            "drift_alerts": count_by(outputs["market"], "DRIFT_ALERT_LEVEL"),
            "critical_drift_references": len(outputs["critical"]),
            "average_market_drift_score": round(statistics.mean(drift_scores), 2) if drift_scores else 0,
            "high_policy": "No HIGH without real supplier, FOB, local benchmark, finance and procurement validation.",
            "duration_seconds": round(time.perf_counter() - start, 2),
        }
        STATS_PATH.write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(stats, ensure_ascii=False, indent=2))
        print("\nLivrables:")
        for path in [SUPPLIER_REGISTRY, VERIFIED_PROCUREMENT, FOB_AUDIT, LOCAL_MARKET, CRITICAL_DRIFT, CONFIDENCE_EVOLUTION, STATS_PATH]:
            print(f"- {path.name}")
        return 0
    except Exception:
        print("SUPPLIER GOVERNANCE PHASE 3 FAILED")
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

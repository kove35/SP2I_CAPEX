from __future__ import annotations

import hashlib
import json
import math
import re
import statistics
import time
import traceback
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
MASTER_HVAC = ROOT / "MASTER_REFERENCE_HVAC.xlsx"

HVAC_SUPPLIERS = ROOT / "HVAC_SUPPLIER_REGISTRY.xlsx"
HVAC_ENERGY = ROOT / "HVAC_ENERGY_GOVERNANCE.xlsx"
HVAC_TCO = ROOT / "HVAC_TCO_ANALYSIS.xlsx"
HVAC_FOB = ROOT / "HVAC_FOB_VALIDATION.xlsx"
HVAC_REVIEW = ROOT / "HVAC_PROCUREMENT_REVIEW.xlsx"
HVAC_DRIFT = ROOT / "HVAC_DRIFT_GOVERNANCE.xlsx"
HVAC_LOW_CONFIDENCE = ROOT / "HVAC_LOW_CONFIDENCE_REFERENCES.xlsx"
STATS_PATH = ROOT / "HVAC_GOVERNANCE_V2_STATS.json"

TODAY = date.today().isoformat()
KWH_PRICE_FCFA = 95
ANNUAL_HOURS = 2200
LIFETIME_YEARS = 10

SUPPLIERS = [
    {
        "SUPPLIER_NAME": "Midea CAC Export Division",
        "BRAND": "Midea",
        "COUNTRY": "China",
        "PRODUCT_TYPES": "SPLIT_MURAL|CASSETTE|GAINABLE",
        "MOQ": "10 U",
        "FOB_MIN": 180000,
        "FOB_MAX": 3200000,
        "LEAD_TIME_DAYS": 50,
        "SUPPLIER_CONFIDENCE": "MEDIUM",
        "VERIFIED": False,
        "VALIDATION_STATUS": "PENDING",
    },
    {
        "SUPPLIER_NAME": "Gree HVAC International",
        "BRAND": "Gree",
        "COUNTRY": "China",
        "PRODUCT_TYPES": "SPLIT_MURAL|CASSETTE|VRV_VRF",
        "MOQ": "8 U",
        "FOB_MIN": 220000,
        "FOB_MAX": 18000000,
        "LEAD_TIME_DAYS": 55,
        "SUPPLIER_CONFIDENCE": "MEDIUM",
        "VERIFIED": False,
        "VALIDATION_STATUS": "PENDING",
    },
    {
        "SUPPLIER_NAME": "Guangzhou Ventilation Equipment Co.",
        "BRAND": "Generic HVAC",
        "COUNTRY": "China",
        "PRODUCT_TYPES": "VENTILATION|EXTRACTION|CONDUITS",
        "MOQ": "50 U",
        "FOB_MIN": 50000,
        "FOB_MAX": 5500000,
        "LEAD_TIME_DAYS": 45,
        "SUPPLIER_CONFIDENCE": "LOW",
        "VERIFIED": False,
        "VALIDATION_STATUS": "PENDING",
    },
]


def number(value: Any, default: float = 0.0) -> float:
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return default if math.isnan(value) else float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return default


def supplier_id(name: str) -> str:
    return "HVAC-SUP-" + hashlib.sha1(name.encode("utf-8")).hexdigest()[:8].upper()


def read_sheet(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(f"Fichier introuvable: {path}")
    wb = load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Onglet {sheet_name} introuvable dans {path.name}")
    ws = wb[sheet_name]
    headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        payload = {headers[index]: value for index, value in enumerate(values)}
        if any(value not in (None, "") for value in payload.values()):
            rows.append(payload)
    wb.close()
    return rows


def supplier_registry() -> list[dict[str, Any]]:
    rows = []
    for supplier in SUPPLIERS:
        rows.append({
            "SUPPLIER_ID": supplier_id(supplier["SUPPLIER_NAME"]),
            "SUPPLIER_NAME": supplier["SUPPLIER_NAME"],
            "BRAND": supplier["BRAND"],
            "COUNTRY": supplier["COUNTRY"],
            "PRODUCT_TYPES": supplier["PRODUCT_TYPES"],
            "MOQ": supplier["MOQ"],
            "FOB_MIN": supplier["FOB_MIN"],
            "FOB_MAX": supplier["FOB_MAX"],
            "LEAD_TIME_DAYS": supplier["LEAD_TIME_DAYS"],
            "SUPPLIER_CONFIDENCE": supplier["SUPPLIER_CONFIDENCE"],
            "VERIFIED": supplier["VERIFIED"],
            "LAST_MARKET_CHECK": TODAY,
            "VALIDATION_STATUS": supplier["VALIDATION_STATUS"],
        })
    return rows


def match_supplier(row: dict[str, Any], suppliers: list[dict[str, Any]]) -> dict[str, Any] | None:
    type_hvac = str(row.get("TYPE_HVAC") or "").upper()
    matches = [supplier for supplier in suppliers if type_hvac in str(supplier["PRODUCT_TYPES"]).upper()]
    return matches[0] if matches else None


def supply_split(row: dict[str, Any]) -> tuple[str, str]:
    raw = str(row.get("RAW_DESIGNATION") or "").lower()
    type_line = str(row.get("SOURCE_REFERENCE") or "")
    if any(term in raw for term in ["pose", "installation", "mise en service", "raccordement"]):
        if any(term in raw for term in ["fourniture", "f & p", "f&p"]):
            return "MIXED_HVAC", "Ligne mixte fourniture + installation, importabilite partielle uniquement."
        return "INSTALLATION_HVAC", "Prestation installation non importable."
    if str(row.get("IMPORTABILITY")) == "HIGH":
        return "FOURNITURE_HVAC", "Fourniture HVAC importable sous reserve fournisseur/FOB."
    return "INSTALLATION_HVAC", f"Non importable ou type source a verifier: {type_line}."


def technical_validation(row: dict[str, Any]) -> tuple[str, str, float, str]:
    type_hvac = str(row.get("TYPE_HVAC") or "HVAC_GENERAL")
    power = number(row.get("PUISSANCE"))
    unit = str(row.get("UNITE_PUISSANCE") or "")
    price = number(row.get("PRICE_MIN_FCFA"))
    warnings: list[str] = []
    kw = 0.0
    if unit == "BTU":
        kw = power / 3412 if power else 0
    elif unit == "CV":
        kw = power * 0.7355
    elif unit == "KW":
        kw = power
    if type_hvac in {"SPLIT_MURAL", "CASSETTE", "GAINABLE", "VRV_VRF"} and not kw:
        warnings.append("puissance absente")
    if kw and price:
        ratio = price / max(kw, 0.1)
        if ratio < 90_000:
            warnings.append("ratio puissance/prix trop faible")
        if ratio > 3_500_000:
            warnings.append("ratio puissance/prix trop eleve")
    cop = expected_cop(type_hvac, row)
    if cop < 2.6:
        warnings.append("COP faible")
    if type_hvac == "VRV_VRF" and kw and kw < 8:
        warnings.append("puissance faible pour VRV/VRF")
    status = "VERIFIED" if not warnings and kw else "PARTIAL" if len(warnings) <= 1 else "REJECTED"
    return status, "; ".join(warnings) if warnings else "coherence technique HVAC acceptable", kw, cop


def expected_cop(type_hvac: str, row: dict[str, Any]) -> float:
    raw = str(row.get("RAW_DESIGNATION") or "").lower()
    inverter_bonus = 0.35 if "inverter" in raw else 0
    base = {
        "SPLIT_MURAL": 3.2,
        "CASSETTE": 3.1,
        "GAINABLE": 3.0,
        "VRV_VRF": 3.6,
        "VENTILATION": 2.8,
        "EXTRACTION": 2.6,
    }.get(type_hvac, 2.8)
    return round(base + inverter_bonus, 2)


def energy_risk(cop: float, kw: float, type_hvac: str) -> tuple[str, float]:
    annual_kwh = kw * ANNUAL_HOURS / max(cop, 1) if kw else 0
    if not kw and type_hvac in {"SPLIT_MURAL", "CASSETTE", "GAINABLE", "VRV_VRF"}:
        return "HIGH", annual_kwh
    if cop < 2.6 or annual_kwh > 28_000:
        return "CRITICAL", annual_kwh
    if cop < 3.0 or annual_kwh > 15_000:
        return "HIGH", annual_kwh
    if annual_kwh > 7_500:
        return "MEDIUM", annual_kwh
    return "LOW", annual_kwh


def fob_validation(row: dict[str, Any], supplier: dict[str, Any] | None) -> tuple[str, str]:
    fob = number(row.get("PU_CHINE_FOB_FCFA"))
    if fob <= 0:
        return "UNVERIFIED", "FOB absent."
    if not supplier:
        return "UNVERIFIED", "Aucun fournisseur HVAC candidat."
    low = number(supplier.get("FOB_MIN"))
    high = number(supplier.get("FOB_MAX"))
    if fob < low * 0.5 or fob > high * 1.8:
        return "REJECTED", "FOB hors range fournisseur candidat."
    return "PARTIAL", "FOB plausible, verification marche requise."


def drift_level(row: dict[str, Any], energy_level: str) -> str:
    variation = number(row.get("PRICE_MAX_FCFA")) / max(number(row.get("PRICE_MIN_FCFA")), 1)
    score = 20
    if variation >= 8:
        score += 45
    elif variation >= 4:
        score += 30
    if energy_level in {"HIGH", "CRITICAL"}:
        score += 18
    if str(row.get("TYPE_HVAC")) in {"CUIVRE_FRIGORIFIQUE", "CONDUITS"}:
        score += 12
    if str(row.get("TYPE_HVAC")) in {"VRV_VRF", "SPLIT_MURAL", "CASSETTE", "GAINABLE"}:
        score += 8
    score = min(score, 100)
    if score >= 80:
        return "CRITICAL"
    if score >= 60:
        return "HIGH"
    if score >= 35:
        return "MEDIUM"
    return "LOW"


def confidence(row: dict[str, Any], supplier: dict[str, Any] | None, fob_status: str, tech_status: str, energy_level: str, drift: str) -> str:
    supplier_verified = bool(supplier and supplier.get("VERIFIED") is True)
    if supplier_verified and fob_status == "VERIFIED" and tech_status == "VERIFIED" and energy_level in {"LOW", "MEDIUM"} and drift in {"LOW", "MEDIUM"}:
        return "HIGH"
    if fob_status in {"PARTIAL", "VERIFIED"} and tech_status in {"PARTIAL", "VERIFIED"} and drift != "CRITICAL":
        return "MEDIUM"
    return "LOW"


def build_outputs(master_rows: list[dict[str, Any]], suppliers: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
    energy_rows = []
    tco_rows = []
    fob_rows = []
    review_rows = []
    drift_rows = []
    low_rows = []
    for row in master_rows:
        supplier = match_supplier(row, suppliers)
        split, split_reason = supply_split(row)
        tech_status, tech_reason, kw, cop = technical_validation(row)
        energy_level, annual_kwh = energy_risk(cop, kw, str(row.get("TYPE_HVAC")))
        fob_status, fob_reason = fob_validation(row, supplier)
        drift = drift_level(row, energy_level)
        conf = confidence(row, supplier, fob_status, tech_status, energy_level, drift)
        purchase = number(row.get("PRICE_MIN_FCFA"))
        energy_cost = round(annual_kwh * KWH_PRICE_FCFA * LIFETIME_YEARS, 2)
        maintenance = round(purchase * 0.08 * LIFETIME_YEARS, 2)
        installation = round(purchase * (0.18 if split == "INSTALLATION_HVAC" else 0.12), 2)
        tco = round(purchase + energy_cost + maintenance + installation, 2)
        base = {
            "REFERENCE_ID": row.get("REFERENCE_ID"),
            "DESIGNATION_NORMALISEE": row.get("DESIGNATION_NORMALISEE"),
            "RAW_DESIGNATION": row.get("RAW_DESIGNATION"),
            "TYPE_HVAC": row.get("TYPE_HVAC"),
            "PUISSANCE_KW_ESTIMEE": round(kw, 3),
            "COP_ESTIME": cop,
            "FOURNITURE_INSTALLATION_CLASS": split,
            "CONFIDENCE_LEVEL": conf,
            "COCKPIT_CONFIDENCE_BADGE": {"HIGH": "GREEN", "MEDIUM": "ORANGE", "LOW": "RED"}[conf],
        }
        energy_rows.append({
            **base,
            "TECHNICAL_VALIDATION_STATUS": tech_status,
            "TECHNICAL_VALIDATION_REASON": tech_reason,
            "ENERGY_RISK_SCORE": energy_level,
            "ANNUAL_KWH_ESTIMATE": round(annual_kwh, 2),
            "INVERTER_DETECTED": "YES" if "inverter" in str(row.get("RAW_DESIGNATION") or "").lower() else "NO",
        })
        tco_rows.append({
            **base,
            "PURCHASE_COST_FCFA": purchase,
            "ENERGY_COST_10Y_FCFA": energy_cost,
            "MAINTENANCE_10Y_FCFA": maintenance,
            "INSTALLATION_COST_FCFA": installation,
            "TOTAL_COST_OF_OWNERSHIP": tco,
            "LIFETIME_YEARS": LIFETIME_YEARS,
        })
        fob_rows.append({
            **base,
            "SUPPLIER_NAME": supplier.get("SUPPLIER_NAME") if supplier else "NO_MATCH",
            "SUPPLIER_VERIFIED": bool(supplier and supplier.get("VERIFIED") is True),
            "PU_CHINE_FOB_FCFA": row.get("PU_CHINE_FOB_FCFA"),
            "FOB_VALIDATION_STATUS": fob_status,
            "FOB_VALIDATION_REASON": fob_reason,
        })
        review_rows.append({
            **base,
            "IMPORTABILITY": row.get("IMPORTABILITY"),
            "PROCUREMENT_REVIEW_REQUIRED": "YES",
            "PROCUREMENT_REVIEW_REASON": "; ".join([split_reason, fob_reason, tech_reason]),
            "SUPPLIER_NAME": supplier.get("SUPPLIER_NAME") if supplier else "NO_MATCH",
        })
        drift_rows.append({
            **base,
            "HVAC_DRIFT_ALERT_LEVEL": drift,
            "DRIFT_FACTORS": "cuivre|fret_maritime|energie|gaz_refrigerant|saison|inflation_import",
            "ENERGY_RISK_SCORE": energy_level,
            "PRICE_VARIATION_RATIO": round(number(row.get("PRICE_MAX_FCFA")) / max(number(row.get("PRICE_MIN_FCFA")), 1), 3),
        })
        if conf != "HIGH":
            low_rows.append({
                **base,
                "NEXT_ACTION": next_action(fob_status, tech_status, energy_level, drift, supplier),
            })
    return {
        "energy": energy_rows,
        "tco": tco_rows,
        "fob": fob_rows,
        "review": review_rows,
        "drift": drift_rows,
        "low": low_rows,
    }


def next_action(fob_status: str, tech_status: str, energy_level: str, drift: str, supplier: dict[str, Any] | None) -> str:
    if not supplier or supplier.get("VERIFIED") is not True:
        return "Verifier fournisseur HVAC reel."
    if fob_status != "VERIFIED":
        return "Confirmer FOB HVAC."
    if tech_status != "VERIFIED":
        return "Valider puissance BTU/CV/HP/COP."
    if energy_level in {"HIGH", "CRITICAL"}:
        return "Revoir efficacite energetique et TCO."
    if drift in {"HIGH", "CRITICAL"}:
        return "Revoir drift cuivre/fret/energie."
    return "Eligible validation HIGH."


def write_xlsx(path: Path, sheet_name: str, rows: list[dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header) for header in headers])
    else:
        ws.append(["STATUS"])
        ws.append(["NO_DATA"])
    style_sheet(ws)
    wb.save(path)


def style_sheet(ws) -> None:
    fill = PatternFill("solid", fgColor="0F172A")
    font = Font(color="F8FAFC", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    for column in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(max(max_len + 2, 12), 64)
    ws.freeze_panes = "A2"


def count_by(rows: list[dict[str, Any]], key: str) -> dict[str, int]:
    result: dict[str, int] = {}
    for row in rows:
        value = str(row.get(key) or "UNKNOWN")
        result[value] = result.get(value, 0) + 1
    return result


def main() -> int:
    start = time.perf_counter()
    print("=== SP2I HVAC Governance V2 ===")
    try:
        master_rows = read_sheet(MASTER_HVAC, "MASTER_REFERENCE_HVAC")
        suppliers = supplier_registry()
        outputs = build_outputs(master_rows, suppliers)
        write_xlsx(HVAC_SUPPLIERS, "HVAC_SUPPLIERS", suppliers)
        write_xlsx(HVAC_ENERGY, "HVAC_ENERGY", outputs["energy"])
        write_xlsx(HVAC_TCO, "HVAC_TCO", outputs["tco"])
        write_xlsx(HVAC_FOB, "HVAC_FOB", outputs["fob"])
        write_xlsx(HVAC_REVIEW, "HVAC_REVIEW", outputs["review"])
        write_xlsx(HVAC_DRIFT, "HVAC_DRIFT", outputs["drift"])
        write_xlsx(HVAC_LOW_CONFIDENCE, "HVAC_LOW_CONFIDENCE", outputs["low"])
        stats = {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "source_file": str(MASTER_HVAC),
            "family": "HVAC",
            "references": len(master_rows),
            "supplier_candidates": len(suppliers),
            "verified_suppliers": len([supplier for supplier in suppliers if supplier["VERIFIED"] is True]),
            "confidence_distribution": count_by(outputs["energy"], "CONFIDENCE_LEVEL"),
            "energy_risk": count_by(outputs["energy"], "ENERGY_RISK_SCORE"),
            "technical_validation": count_by(outputs["energy"], "TECHNICAL_VALIDATION_STATUS"),
            "fob_validation": count_by(outputs["fob"], "FOB_VALIDATION_STATUS"),
            "drift_alerts": count_by(outputs["drift"], "HVAC_DRIFT_ALERT_LEVEL"),
            "low_confidence_references": len(outputs["low"]),
            "high_policy": "No HIGH without verified supplier, verified FOB, verified benchmark, coherent power/COP and acceptable drift.",
            "duration_seconds": round(time.perf_counter() - start, 2),
        }
        (ROOT / "HVAC_GOVERNANCE_V2_STATS.json").write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(stats, ensure_ascii=False, indent=2))
        print("\nLivrables:")
        for path in [HVAC_SUPPLIERS, HVAC_ENERGY, HVAC_TCO, HVAC_FOB, HVAC_REVIEW, HVAC_DRIFT, HVAC_LOW_CONFIDENCE, ROOT / "HVAC_GOVERNANCE_V2_STATS.json"]:
            print(f"- {path.name}")
        return 0
    except Exception:
        print("HVAC GOVERNANCE V2 FAILED")
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

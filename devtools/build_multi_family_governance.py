from __future__ import annotations

import json
import math
import statistics
import time
import traceback
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
TODAY = date.today().isoformat()

FAMILY_INDEX_OUT = ROOT / "FAMILY_GOVERNANCE_INDEX.xlsx"
CANDIDATES_OUT = ROOT / "MASTER_REFERENCE_ENTERPRISE_CANDIDATES.xlsx"
CROSS_AUDIT_OUT = ROOT / "CROSS_FAMILY_GOVERNANCE_AUDIT.xlsx"
MONITORING_OUT = ROOT / "GOVERNANCE_MONITORING_REPORT.xlsx"
STATS_OUT = ROOT / "MULTI_FAMILY_GOVERNANCE_STATS.json"

FAMILIES = [
    {
        "name": "ELECTRICITE",
        "master": ROOT / "MASTER_REFERENCE_GOVERNANCE_V2.xlsx",
        "fallback_master": ROOT / "MASTER_REFERENCE_ENTERPRISE.xlsx",
        "stats": ROOT / "SUPPLIER_GOVERNANCE_STATS.json",
        "fallback_stats": ROOT / "PROCUREMENT_GOVERNANCE_STATS.json",
        "supplier": ROOT / "SUPPLIER_REFERENCE_REGISTRY.xlsx",
        "tco": None,
    },
    {
        "name": "HVAC",
        "master": ROOT / "MASTER_REFERENCE_HVAC.xlsx",
        "stats": ROOT / "HVAC_GOVERNANCE_V2_STATS.json",
        "fallback_stats": ROOT / "HVAC_VALIDATION_STATS.json",
        "supplier": ROOT / "HVAC_SUPPLIER_REGISTRY.xlsx",
        "tco": ROOT / "HVAC_TCO_ANALYSIS.xlsx",
    },
    {
        "name": "PLOMBERIE",
        "master": ROOT / "MASTER_REFERENCE_PLOMBERIE.xlsx",
        "stats": ROOT / "PLOMBERIE_VALIDATION_STATS.json",
        "supplier": ROOT / "PLOMBERIE_SUPPLIER_REGISTRY.xlsx",
        "tco": ROOT / "PLOMBERIE_TCO_ANALYSIS.xlsx",
    },
    {
        "name": "MENUISERIE_ALU",
        "master": ROOT / "MASTER_REFERENCE_MENUISERIE.xlsx",
        "stats": ROOT / "MENUISERIE_VALIDATION_STATS.json",
        "supplier": ROOT / "MENUISERIE_SUPPLIER_REGISTRY.xlsx",
        "tco": ROOT / "MENUISERIE_TCO_ANALYSIS.xlsx",
    },
]

STATUS_RULES = {
    "VERIFIED": "Gouvernance stable, integration possible.",
    "CONDITIONALLY_READY": "Integration possible avec supervision.",
    "REVIEW_REQUIRED": "Validation metier requise.",
    "HIGH_RISK": "Risques importants, integration non recommandee sans arbitrage.",
    "BLOCKED": "Non integrable en l'etat.",
}


def load_json(path: Path | None) -> dict[str, Any]:
    if not path or not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def resolve_existing(*paths: Path | None) -> Path | None:
    for path in paths:
        if path and path.exists():
            return path
    return None


def normalize_header(value: Any) -> str:
    return str(value or "").strip().upper().replace(" ", "_")


def read_xlsx_rows(path: Path | None) -> list[dict[str, Any]]:
    if not path or not path.exists():
        return []
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [normalize_header(cell) for cell in rows[0]]
    data: list[dict[str, Any]] = []
    for values in rows[1:]:
        row = {headers[index]: value for index, value in enumerate(values) if index < len(headers)}
        if any(value not in (None, "") for value in row.values()):
            data.append(row)
    return data


def count_total(distribution: dict[str, Any]) -> int:
    return sum(int(value or 0) for value in distribution.values())


def weighted_score(distribution: dict[str, Any], weights: dict[str, float], default: float = 60.0) -> float:
    total = count_total(distribution)
    if not total:
        return default
    score = 0.0
    for key, value in distribution.items():
        score += int(value or 0) * weights.get(str(key).upper(), default)
    return round(score / total, 2)


def ratio(numerator: float, denominator: float) -> float:
    if not denominator:
        return 0.0
    return round(numerator / denominator, 4)


def clamp(value: float, low: float = 0.0, high: float = 100.0) -> float:
    return max(low, min(high, value))


def confidence_distribution(stats: dict[str, Any]) -> dict[str, int]:
    for key in ("confidence_after", "confidence_distribution"):
        value = stats.get(key)
        if isinstance(value, dict):
            return {str(k).upper(): int(v or 0) for k, v in value.items()}
    return {}


def drift_distribution(stats: dict[str, Any]) -> dict[str, int]:
    value = stats.get("drift_alerts") or {}
    if isinstance(value, dict):
        return {str(k).upper(): int(v or 0) for k, v in value.items()}
    return {}


def technical_distribution(stats: dict[str, Any]) -> dict[str, int]:
    value = stats.get("technical_validation") or {}
    if isinstance(value, dict):
        return {str(k).upper(): int(v or 0) for k, v in value.items()}
    references = int(stats.get("references") or stats.get("master_references") or 0)
    review = int(stats.get("review_required") or stats.get("procurement_review_required") or 0)
    if references:
        return {"PARTIAL": min(references, review) or references}
    return {}


def blocker_distribution(stats: dict[str, Any]) -> dict[str, int]:
    value = stats.get("decision_blockers") or {}
    if isinstance(value, dict):
        return {str(k).upper(): int(v or 0) for k, v in value.items()}
    review = int(stats.get("review_required") or stats.get("procurement_review_required") or 0)
    return {"REVIEW_REQUIRED": review} if review else {}


def reference_count(stats: dict[str, Any], rows: list[dict[str, Any]]) -> int:
    return int(
        stats.get("references")
        or stats.get("master_references")
        or stats.get("menuiserie_rows")
        or stats.get("plomberie_rows")
        or stats.get("hvac_rows")
        or len(rows)
    )


def tco_maturity(family: dict[str, Any], refs: int) -> float:
    tco_path = family.get("tco")
    if not tco_path:
        return 55.0
    rows = read_xlsx_rows(tco_path)
    if not rows:
        return 55.0
    coverage = min(1.0, len(rows) / refs) if refs else 0.0
    return round(60.0 + coverage * 30.0, 2)


def procurement_score(conf_score: float, blocker_rate: float, review_rate: float) -> float:
    score = conf_score - blocker_rate * 35.0 - review_rate * 10.0
    return round(clamp(score), 2)


def integration_score(
    conf_score: float,
    procurement: float,
    drift_score: float,
    tco_score: float,
    technical_score: float,
    blocker_rate: float,
) -> float:
    score = (
        conf_score * 0.22
        + procurement * 0.22
        + drift_score * 0.18
        + tco_score * 0.13
        + technical_score * 0.17
        + (100.0 - blocker_rate * 100.0) * 0.08
    )
    return round(clamp(score), 2)


def readiness_label(score: float) -> str:
    if score > 85:
        return "READY"
    if score >= 70:
        return "REVIEW"
    return "BLOCKED"


def family_status(score: float, blocker_rate: float, critical_drift_rate: float) -> str:
    if blocker_rate >= 0.65:
        return "BLOCKED"
    if critical_drift_rate >= 0.35:
        return "HIGH_RISK"
    if score > 85 and blocker_rate < 0.10:
        return "VERIFIED"
    if score >= 70:
        return "CONDITIONALLY_READY"
    return "REVIEW_REQUIRED"


def badge_from_confidence(value: str) -> str:
    level = str(value or "").upper()
    if level == "HIGH":
        return "GREEN"
    if level == "MEDIUM":
        return "ORANGE"
    return "RED"


def badge_from_risk(value: str) -> str:
    level = str(value or "").upper()
    if level in {"CRITICAL", "HIGH", "BLOCKED", "HIGH_RISK"}:
        return "RED"
    if level in {"MEDIUM", "REVIEW", "REVIEW_REQUIRED"}:
        return "ORANGE"
    return "GREEN"


def choose(row: dict[str, Any], *keys: str, default: Any = "") -> Any:
    for key in keys:
        value = row.get(normalize_header(key))
        if value not in (None, ""):
            return value
    return default


def as_text(value: Any) -> str:
    return str(value or "").strip()


def build_family_records() -> tuple[list[dict[str, Any]], dict[str, list[dict[str, Any]]]]:
    family_index: list[dict[str, Any]] = []
    family_rows: dict[str, list[dict[str, Any]]] = {}

    for family in FAMILIES:
        master_path = resolve_existing(family.get("master"), family.get("fallback_master"))
        stats_path = resolve_existing(family.get("stats"), family.get("fallback_stats"))
        stats = load_json(stats_path)
        rows = read_xlsx_rows(master_path)
        refs = reference_count(stats, rows)
        conf = confidence_distribution(stats)
        drift = drift_distribution(stats)
        tech = technical_distribution(stats)
        blockers = blocker_distribution(stats)

        family_rows[family["name"]] = rows

        conf_score = weighted_score(conf, {"HIGH": 100, "MEDIUM": 70, "LOW": 35}, 55)
        drift_score = weighted_score(drift, {"LOW": 100, "MEDIUM": 78, "HIGH": 45, "CRITICAL": 18}, 65)
        tech_score = weighted_score(tech, {"VERIFIED": 100, "PARTIAL": 62, "UNVERIFIED": 45, "REJECTED": 20}, 60)
        blocker_count = sum(value for key, value in blockers.items() if key in {"BLOCK_IMPORT", "HIGH_RISK", "TECHNICAL_VALIDATION_REQUIRED"})
        review_count = sum(value for key, value in blockers.items() if key in {"REVIEW_REQUIRED"})
        blocker_rate = ratio(blocker_count, refs)
        review_rate = ratio(review_count, refs)
        tco_score = tco_maturity(family, refs)
        procurement = procurement_score(conf_score, blocker_rate, review_rate)
        integration = integration_score(conf_score, procurement, drift_score, tco_score, tech_score, blocker_rate)
        critical_drift_rate = ratio(drift.get("CRITICAL", 0), refs)
        status = family_status(integration, blocker_rate, critical_drift_rate)

        family_index.append(
            {
                "FAMILY_NAME": family["name"],
                "FAMILY_STATUS": status,
                "REFERENCE_COUNT": refs,
                "CONFIDENCE_MEDIUM": conf.get("MEDIUM", 0),
                "CONFIDENCE_LOW": conf.get("LOW", 0),
                "CONFIDENCE_HIGH": conf.get("HIGH", 0),
                "PROCUREMENT_SCORE": procurement,
                "DRIFT_SCORE": drift_score,
                "TCO_MATURITY": tco_score,
                "TECHNICAL_VALIDATION_SCORE": tech_score,
                "INTEGRATION_READINESS_SCORE": integration,
                "INTEGRATION_GATE": readiness_label(integration),
                "DECISION_BLOCKER_RATE": blocker_rate,
                "REVIEW_REQUIRED_RATE": review_rate,
                "CRITICAL_DRIFT_RATE": critical_drift_rate,
                "LAST_REVIEW_DATE": TODAY,
                "SOURCE_MASTER": master_path.name if master_path else "MISSING",
                "SOURCE_STATS": stats_path.name if stats_path else "MISSING",
                "STATUS_RULE": STATUS_RULES.get(status, ""),
            }
        )

    return family_index, family_rows


def build_candidates(family_index: list[dict[str, Any]], family_rows: dict[str, list[dict[str, Any]]]) -> list[dict[str, Any]]:
    scores = {row["FAMILY_NAME"]: float(row["INTEGRATION_READINESS_SCORE"]) for row in family_index}
    gates = {row["FAMILY_NAME"]: row["INTEGRATION_GATE"] for row in family_index}
    candidates: list[dict[str, Any]] = []

    for family_name, rows in family_rows.items():
        for index, row in enumerate(rows, start=1):
            confidence = as_text(choose(row, "CONFIDENCE_LEVEL", "CONFIDENCE", default="LOW")).upper() or "LOW"
            drift = as_text(choose(row, "DRIFT_ALERT_LEVEL", "HVAC_DRIFT_ALERT_LEVEL", "PLOMBERIE_DRIFT_ALERT_LEVEL", "MENUISERIE_DRIFT_ALERT_LEVEL", "MARKET_DRIFT_ALERT_LEVEL", default="MEDIUM")).upper()
            technical = as_text(choose(row, "TECHNICAL_VALIDATION", "VALIDATION_STATUS", "FOB_VALIDATION_STATUS", default="PARTIAL")).upper()
            blocker = as_text(choose(row, "DECISION_BLOCKER", "PROCUREMENT_REVIEW_REQUIRED", default="REVIEW_REQUIRED")).upper()
            designation = as_text(
                choose(
                    row,
                    "DESIGNATION_NORMALISEE",
                    "NORMALIZED_DESIGNATION",
                    "RAW_DESIGNATION",
                    "DESIGNATION",
                    "DESCRIPTION",
                    default=f"{family_name}_REFERENCE_{index:03d}",
                )
            )
            ref_id = as_text(choose(row, "REFERENCE_ID", default=f"{family_name[:3]}-CAND-{index:04d}"))
            family_score = scores.get(family_name, 0.0)
            row_penalty = 0.0
            if confidence == "LOW":
                row_penalty += 12.0
            if drift in {"HIGH", "CRITICAL"}:
                row_penalty += 12.0 if drift == "HIGH" else 22.0
            if blocker in {"BLOCK_IMPORT", "HIGH_RISK", "TECHNICAL_VALIDATION_REQUIRED"}:
                row_penalty += 18.0
            if technical in {"REJECTED", "UNVERIFIED"}:
                row_penalty += 15.0
            readiness = round(clamp(family_score - row_penalty), 2)
            status = readiness_label(readiness)
            manual_review = status != "READY" or confidence != "HIGH" or blocker not in {"NONE", "FALSE", "NO"}

            candidates.append(
                {
                    "REFERENCE_ID": ref_id,
                    "FAMILY_NAME": family_name,
                    "DESIGNATION_NORMALISEE": designation,
                    "CONFIDENCE_LEVEL": confidence,
                    "PROCUREMENT_STATUS": as_text(choose(row, "PROCUREMENT_STATUS", "IMPORTABILITY", "VALIDATION_STATUS", default=gates.get(family_name, "REVIEW"))),
                    "DRIFT_LEVEL": drift or "MEDIUM",
                    "TECHNICAL_VALIDATION": technical or "PARTIAL",
                    "INTEGRATION_STATUS": status,
                    "INTEGRATION_READINESS_SCORE": readiness,
                    "DECISION_BLOCKER": blocker or "REVIEW_REQUIRED",
                    "MANUAL_REVIEW_REQUIRED": manual_review,
                    "COCKPIT_GOVERNANCE_STATUS": status,
                    "COCKPIT_RISK_BADGE": badge_from_risk(blocker or drift),
                    "COCKPIT_CONFIDENCE_BADGE": badge_from_confidence(confidence),
                    "COCKPIT_DRIFT_ALERT": drift or "MEDIUM",
                    "COCKPIT_REVIEW_REQUIRED": manual_review,
                }
            )

    return candidates


def read_suppliers() -> list[dict[str, Any]]:
    suppliers: list[dict[str, Any]] = []
    for family in FAMILIES:
        for row in read_xlsx_rows(family.get("supplier")):
            name = as_text(choose(row, "SUPPLIER_NAME", "FOURNISSEUR_CHINE", default="UNKNOWN_SUPPLIER"))
            suppliers.append(
                {
                    "FAMILY_NAME": family["name"],
                    "SUPPLIER_NAME": name,
                    "COUNTRY": choose(row, "COUNTRY", default=""),
                    "PRODUCT_FAMILIES": choose(row, "PRODUCT_FAMILIES", "PRODUCT_TYPES", default=""),
                    "VERIFIED": choose(row, "VERIFIED", default=False),
                }
            )
    return suppliers


def build_cross_audit(family_index: list[dict[str, Any]], candidates: list[dict[str, Any]]) -> list[dict[str, Any]]:
    audit: list[dict[str, Any]] = []
    suppliers = read_suppliers()
    by_supplier: dict[str, set[str]] = {}
    for supplier in suppliers:
        key = as_text(supplier["SUPPLIER_NAME"]).upper()
        by_supplier.setdefault(key, set()).add(supplier["FAMILY_NAME"])

    for supplier_name, families in sorted(by_supplier.items()):
        if supplier_name and len(families) > 1:
            audit.append(
                {
                    "AUDIT_TYPE": "DUPLICATED_SUPPLIER",
                    "SEVERITY": "WARNING",
                    "FAMILIES": "|".join(sorted(families)),
                    "REFERENCE": supplier_name,
                    "FINDING": "Fournisseur present sur plusieurs familles, verifier absence de conflit procurement.",
                    "RECOMMENDATION": "Centraliser la validation fournisseur avant integration master.",
                }
            )

    family_by_name = {row["FAMILY_NAME"]: row for row in family_index}
    for row in family_index:
        if row["CRITICAL_DRIFT_RATE"] >= 0.25:
            audit.append(
                {
                    "AUDIT_TYPE": "DRIFT_INCOHERENCE",
                    "SEVERITY": "CRITICAL" if row["CRITICAL_DRIFT_RATE"] >= 0.5 else "HIGH",
                    "FAMILIES": row["FAMILY_NAME"],
                    "REFERENCE": "DRIFT",
                    "FINDING": f"Drift critique eleve ({row['CRITICAL_DRIFT_RATE']:.2%}).",
                    "RECOMMENDATION": "Geler integration automatique et lancer revue benchmark / devises / fret.",
                }
            )

    designation_map: dict[str, set[str]] = {}
    for candidate in candidates:
        key = as_text(candidate["DESIGNATION_NORMALISEE"]).upper()
        if len(key) >= 8:
            designation_map.setdefault(key, set()).add(candidate["FAMILY_NAME"])
    for designation, families in designation_map.items():
        if len(families) > 1:
            audit.append(
                {
                    "AUDIT_TYPE": "REFERENCE_OVERLAP",
                    "SEVERITY": "WARNING",
                    "FAMILIES": "|".join(sorted(families)),
                    "REFERENCE": designation[:120],
                    "FINDING": "Designation proche detectee dans plusieurs familles.",
                    "RECOMMENDATION": "Verifier taxonomy avant toute fusion master.",
                }
            )

    for candidate in candidates:
        family = family_by_name.get(candidate["FAMILY_NAME"], {})
        if candidate["CONFIDENCE_LEVEL"] == "LOW" and family.get("INTEGRATION_GATE") == "READY":
            audit.append(
                {
                    "AUDIT_TYPE": "CONFIDENCE_CONFLICT",
                    "SEVERITY": "WARNING",
                    "FAMILIES": candidate["FAMILY_NAME"],
                    "REFERENCE": candidate["REFERENCE_ID"],
                    "FINDING": "Famille prete mais reference faible confiance.",
                    "RECOMMENDATION": "Exclure la reference du lot candidat tant que non validee.",
                }
            )

    if not audit:
        audit.append(
            {
                "AUDIT_TYPE": "NO_MAJOR_CROSS_FAMILY_CONFLICT",
                "SEVERITY": "INFO",
                "FAMILIES": "ALL",
                "REFERENCE": "N/A",
                "FINDING": "Aucun conflit transverse majeur detecte par les regles actuelles.",
                "RECOMMENDATION": "Maintenir revue humaine avant fusion master.",
            }
        )
    return audit


def global_kpis(family_index: list[dict[str, Any]]) -> dict[str, Any]:
    if not family_index:
        return {}
    return {
        "GLOBAL_GOVERNANCE_SCORE": round(statistics.mean(float(row["INTEGRATION_READINESS_SCORE"]) for row in family_index), 2),
        "GLOBAL_PROCUREMENT_SCORE": round(statistics.mean(float(row["PROCUREMENT_SCORE"]) for row in family_index), 2),
        "GLOBAL_CONFIDENCE_SCORE": round(
            statistics.mean(
                weighted_score(
                    {"HIGH": row["CONFIDENCE_HIGH"], "MEDIUM": row["CONFIDENCE_MEDIUM"], "LOW": row["CONFIDENCE_LOW"]},
                    {"HIGH": 100, "MEDIUM": 70, "LOW": 35},
                    55,
                )
                for row in family_index
            ),
            2,
        ),
        "GLOBAL_DRIFT_SCORE": round(statistics.mean(float(row["DRIFT_SCORE"]) for row in family_index), 2),
        "GLOBAL_TCO_SCORE": round(statistics.mean(float(row["TCO_MATURITY"]) for row in family_index), 2),
        "GLOBAL_TECHNICAL_VALIDATION_SCORE": round(statistics.mean(float(row["TECHNICAL_VALIDATION_SCORE"]) for row in family_index), 2),
        "FAMILY_COUNT": len(family_index),
        "REFERENCE_COUNT": sum(int(row["REFERENCE_COUNT"]) for row in family_index),
        "READY_FAMILIES": sum(1 for row in family_index if row["INTEGRATION_GATE"] == "READY"),
        "REVIEW_FAMILIES": sum(1 for row in family_index if row["INTEGRATION_GATE"] == "REVIEW"),
        "BLOCKED_FAMILIES": sum(1 for row in family_index if row["INTEGRATION_GATE"] == "BLOCKED"),
    }


def write_xlsx(path: Path, sheets: dict[str, list[dict[str, Any]]]) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(sheet_name[:31])
        headers = sorted({key for row in rows for key in row.keys()}) if rows else ["STATUS"]
        if rows:
            priority = list(rows[0].keys())
            headers = priority + [key for key in headers if key not in priority]
        ws.append(headers)
        for row in rows or [{"STATUS": "NO_DATA"}]:
            ws.append([row.get(header, "") for header in headers])
        style_sheet(ws)
    wb.save(path)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="103344")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            value = str(cell.value or "").upper()
            if value in {"BLOCKED", "CRITICAL", "RED", "HIGH_RISK", "BLOCK_IMPORT"}:
                cell.fill = PatternFill("solid", fgColor="F4CCCC")
            elif value in {"REVIEW", "REVIEW_REQUIRED", "HIGH", "ORANGE", "PARTIAL", "MEDIUM"}:
                cell.fill = PatternFill("solid", fgColor="FCE5CD")
            elif value in {"READY", "GREEN", "VERIFIED", "LOW", "NONE"}:
                cell.fill = PatternFill("solid", fgColor="D9EAD3")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        column_letter = get_column_letter(column[0].column)
        width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        ws.column_dimensions[column_letter].width = width


def main() -> int:
    start = time.perf_counter()
    print("=== SP2I Multi-Family Governance Layer ===")
    try:
        family_index, family_rows = build_family_records()
        candidates = build_candidates(family_index, family_rows)
        cross_audit = build_cross_audit(family_index, candidates)
        kpis = global_kpis(family_index)
        monitoring_rows = [
            {
                "MONITORING_DATE": TODAY,
                "FAMILY_NAME": row["FAMILY_NAME"],
                "FAMILY_STATUS": row["FAMILY_STATUS"],
                "INTEGRATION_READINESS_SCORE": row["INTEGRATION_READINESS_SCORE"],
                "PROCUREMENT_SCORE": row["PROCUREMENT_SCORE"],
                "DRIFT_SCORE": row["DRIFT_SCORE"],
                "TCO_MATURITY": row["TCO_MATURITY"],
                "TECHNICAL_VALIDATION_SCORE": row["TECHNICAL_VALIDATION_SCORE"],
                "DECISION_BLOCKER_RATE": row["DECISION_BLOCKER_RATE"],
                "NEXT_ACTION": "Manual review before master integration" if row["INTEGRATION_GATE"] != "READY" else "Supervised integration candidate",
            }
            for row in family_index
        ]
        kpi_rows = [{"KPI": key, "VALUE": value} for key, value in kpis.items()]

        write_xlsx(FAMILY_INDEX_OUT, {"FAMILY_GOVERNANCE_INDEX": family_index})
        write_xlsx(CANDIDATES_OUT, {"MASTER_CANDIDATES": candidates})
        write_xlsx(CROSS_AUDIT_OUT, {"CROSS_FAMILY_AUDIT": cross_audit})
        write_xlsx(
            MONITORING_OUT,
            {
                "GLOBAL_KPIS": kpi_rows,
                "FAMILY_MONITORING": monitoring_rows,
                "STATUS_RULES": [{"STATUS": key, "RULE": value} for key, value in STATUS_RULES.items()],
            },
        )

        stats = {
            "generated_at": datetime.now().replace(microsecond=0).isoformat(),
            "families": [family["name"] for family in FAMILIES],
            "global_kpis": kpis,
            "family_status": {row["FAMILY_NAME"]: row["FAMILY_STATUS"] for row in family_index},
            "integration_gates": {row["FAMILY_NAME"]: row["INTEGRATION_GATE"] for row in family_index},
            "candidate_references": len(candidates),
            "manual_review_required": sum(1 for row in candidates if row["MANUAL_REVIEW_REQUIRED"]),
            "cross_family_findings": len(cross_audit),
            "duration_seconds": round(time.perf_counter() - start, 2),
        }
        STATS_OUT.write_text(json.dumps(stats, indent=2, ensure_ascii=False), encoding="utf-8")

        print(json.dumps(stats, indent=2, ensure_ascii=False))
        print("\nLivrables:")
        for path in [FAMILY_INDEX_OUT, CANDIDATES_OUT, CROSS_AUDIT_OUT, MONITORING_OUT, STATS_OUT]:
            print(f"- {path.name}")
        return 0
    except Exception:
        print("\nERREUR MULTI-FAMILY GOVERNANCE")
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

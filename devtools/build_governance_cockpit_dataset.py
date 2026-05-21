from __future__ import annotations

import json
import statistics
import time
import traceback
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
TODAY = date.today().isoformat()

QUEUE = ROOT / "PROCUREMENT_REVIEW_QUEUE.xlsx"
FAMILY_INDEX = ROOT / "FAMILY_GOVERNANCE_INDEX.xlsx"
OVERRIDE_LOG = ROOT / "GOVERNANCE_OVERRIDE_LOG.xlsx"
COMMENTS = ROOT / "PROCUREMENT_GOVERNANCE_COMMENTS.xlsx"
REVIEW_AUDIT = ROOT / "PROCUREMENT_REVIEW_AUDIT.xlsx"
MULTI_STATS = ROOT / "MULTI_FAMILY_GOVERNANCE_STATS.json"
WORKBENCH_STATS = ROOT / "PROCUREMENT_WORKBENCH_STATS.json"

DATASET_OUT = ROOT / "GOVERNANCE_COCKPIT_DATASET.xlsx"
WORKFLOW_OUT = ROOT / "GOVERNANCE_WORKFLOW_TIMELINE.xlsx"
ESCALATION_OUT = ROOT / "GOVERNANCE_ESCALATION_MATRIX.xlsx"
EXPLAINABILITY_OUT = ROOT / "GOVERNANCE_EXPLAINABILITY_DATA.xlsx"
AUDIT_TIMELINE_OUT = ROOT / "GOVERNANCE_AUDIT_TIMELINE.xlsx"
STATS_OUT = ROOT / "COCKPIT_GOVERNANCE_STATS.json"

WORKFLOW_STATES = ["PENDING", "IN_REVIEW", "ESCALATED", "APPROVED", "REJECTED", "NEEDS_MORE_DATA"]
BADGE_COLORS = {
    "GREEN": "Stable / validated",
    "ORANGE": "Review required",
    "RED": "High risk",
    "DARK_RED": "Critical / blocked",
}


def normalize_header(value: Any) -> str:
    return str(value or "").strip().upper().replace(" ", "_")


def as_text(value: Any) -> str:
    return str(value or "").strip()


def as_float(value: Any, default: float = 0.0) -> float:
    try:
        if value in (None, ""):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def load_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def read_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [normalize_header(cell) for cell in rows[0]]
    data: list[dict[str, Any]] = []
    for values in rows[1:]:
        row = {headers[index]: value for index, value in enumerate(values) if index < len(headers)}
        if any(value not in (None, "") for value in row.values()):
            data.append(row)
    return data


def choose(row: dict[str, Any], *keys: str, default: Any = "") -> Any:
    for key in keys:
        value = row.get(normalize_header(key))
        if value not in (None, ""):
            return value
    return default


def count_by(rows: list[dict[str, Any]], key: str) -> dict[str, int]:
    values: dict[str, int] = {}
    for row in rows:
        value = as_text(choose(row, key, default="EMPTY")).upper() or "EMPTY"
        values[value] = values.get(value, 0) + 1
    return dict(sorted(values.items()))


def severity_level(value: str) -> str:
    level = value.upper()
    if level in {"CRITICAL", "BLOCK_IMPORT", "SENIOR_PROCUREMENT_APPROVAL", "EXECUTIVE_ESCALATION"}:
        return "CRITICAL"
    if level in {"HIGH", "HIGH_RISK", "DOUBLE_VALIDATION", "SENIOR_REVIEW"}:
        return "HIGH"
    if level in {"MEDIUM", "PROCUREMENT_REVIEW", "PENDING", "REVIEW_REQUIRED"}:
        return "MEDIUM"
    return "LOW"


def badge_for_confidence(value: Any) -> str:
    level = as_text(value).upper()
    if level == "HIGH":
        return "GREEN"
    if level == "MEDIUM":
        return "ORANGE"
    if level == "LOW":
        return "RED"
    return "ORANGE"


def badge_for_drift(value: Any) -> str:
    level = as_text(value).upper()
    if level == "CRITICAL":
        return "DARK_RED"
    if level == "HIGH":
        return "RED"
    if level == "MEDIUM":
        return "ORANGE"
    return "GREEN"


def badge_for_blocker(value: Any) -> str:
    blocker = as_text(value).upper()
    if blocker == "BLOCK_IMPORT":
        return "DARK_RED"
    if blocker in {"HIGH_RISK", "TECHNICAL_VALIDATION_REQUIRED"}:
        return "RED"
    if blocker == "REVIEW_REQUIRED":
        return "ORANGE"
    return "GREEN"


def alert_level(row: dict[str, Any], alert_type: str) -> str:
    drift = as_text(choose(row, "DRIFT_LEVEL", default="MEDIUM")).upper()
    blocker = as_text(choose(row, "DECISION_BLOCKER", default="REVIEW_REQUIRED")).upper()
    confidence = as_text(choose(row, "CONFIDENCE_LEVEL", default="LOW")).upper()
    technical = as_text(choose(row, "TECHNICAL_VALIDATION", default="PARTIAL")).upper()
    priority = as_text(choose(row, "REVIEW_PRIORITY", default="MEDIUM")).upper()

    if alert_type == "DRIFT":
        return "CRITICAL" if drift == "CRITICAL" else "HIGH" if drift == "HIGH" else "MEDIUM" if drift == "MEDIUM" else "LOW"
    if alert_type == "PROCUREMENT":
        if blocker == "BLOCK_IMPORT":
            return "CRITICAL"
        if blocker == "HIGH_RISK":
            return "HIGH"
        return "MEDIUM" if blocker == "REVIEW_REQUIRED" else "LOW"
    if alert_type == "TECHNICAL":
        if technical == "REJECTED":
            return "CRITICAL"
        if technical in {"UNVERIFIED", "PARTIAL"}:
            return "HIGH"
        return "LOW"
    if alert_type == "BLOCK_IMPORT":
        return "CRITICAL" if blocker == "BLOCK_IMPORT" else "LOW"
    if priority == "CRITICAL" or confidence == "LOW":
        return "CRITICAL" if priority == "CRITICAL" else "HIGH"
    return "HIGH" if priority == "HIGH" else "MEDIUM"


def explainability(row: dict[str, Any]) -> dict[str, Any]:
    ref_id = as_text(choose(row, "REFERENCE_ID", default="UNKNOWN_REF"))
    family = as_text(choose(row, "FAMILY_NAME", default="UNKNOWN"))
    blocker = as_text(choose(row, "DECISION_BLOCKER", default="REVIEW_REQUIRED")).upper()
    drift = as_text(choose(row, "DRIFT_LEVEL", default="MEDIUM")).upper()
    confidence = as_text(choose(row, "CONFIDENCE_LEVEL", default="LOW")).upper()
    escalation = as_text(choose(row, "REQUIRED_ESCALATION_LEVEL", default="STANDARD_REVIEW")).upper()
    technical = as_text(choose(row, "TECHNICAL_VALIDATION", default="PARTIAL")).upper()
    explanation = as_text(choose(row, "PROCUREMENT_EXPLANATION", default="Validation humaine requise."))
    reasons: list[str] = []

    if blocker == "BLOCK_IMPORT":
        reasons.append("Import bloque par decision governance.")
    elif blocker == "HIGH_RISK":
        reasons.append("Risque decisionnel eleve detecte.")
    elif blocker == "REVIEW_REQUIRED":
        reasons.append("Revue procurement obligatoire.")
    if drift in {"CRITICAL", "HIGH"}:
        reasons.append(f"Drift marche {drift.lower()} a verifier.")
    if confidence == "LOW":
        reasons.append("Confidence faible, validation humaine requise.")
    if technical in {"PARTIAL", "REJECTED", "UNVERIFIED"}:
        reasons.append(f"Validation technique {technical.lower()}.")

    return {
        "REFERENCE_ID": ref_id,
        "FAMILY_NAME": family,
        "WHY_BLOCKED": " ".join(reasons) if reasons else "Non bloque, revue de prudence.",
        "WHY_ESCALATED": escalation.replace("_", " ").title() if escalation != "STANDARD_REVIEW" else "Pas d'escalation senior.",
        "WHY_REVIEW_REQUIRED": "Decision critique non automatisee; revue humaine obligatoire.",
        "WHY_CONFIDENCE_LOW": "Fournisseur, benchmark, FOB ou drift insuffisamment verifies." if confidence == "LOW" else "Confidence non faible.",
        "WHY_DRIFT_CRITICAL": "Derive marche critique detectee." if drift == "CRITICAL" else "Drift non critique.",
        "PROCUREMENT_EXPLANATION": explanation,
        "NEXT_BEST_ACTION": choose(row, "RECOMMENDED_ACTIONS", default="ADD_GOVERNANCE_COMMENT"),
    }


def family_map() -> dict[str, dict[str, Any]]:
    return {as_text(row.get("FAMILY_NAME")).upper(): row for row in read_xlsx_rows(FAMILY_INDEX)}


def build_cockpit_dataset() -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    queue = read_xlsx_rows(QUEUE)
    families = family_map()
    rows: list[dict[str, Any]] = []
    explain_rows: list[dict[str, Any]] = []

    for row in queue:
        family = as_text(choose(row, "FAMILY_NAME", default="UNKNOWN")).upper()
        family_row = families.get(family, {})
        item = {
            "REFERENCE_ID": choose(row, "REFERENCE_ID", default=""),
            "FAMILY_NAME": family,
            "DESIGNATION_NORMALISEE": choose(row, "DESIGNATION_NORMALISEE", default=""),
            "FAMILY_STATUS": choose(family_row, "FAMILY_STATUS", default="UNKNOWN"),
            "READINESS_SCORE": choose(family_row, "INTEGRATION_READINESS_SCORE", default=""),
            "CONFIDENCE_LEVEL": choose(row, "CONFIDENCE_LEVEL", default="LOW"),
            "PROCUREMENT_SCORE": choose(row, "PROCUREMENT_SCORE", default=choose(family_row, "PROCUREMENT_SCORE", default="")),
            "DRIFT_LEVEL": choose(row, "DRIFT_LEVEL", default="MEDIUM"),
            "TECHNICAL_VALIDATION": choose(row, "TECHNICAL_VALIDATION", default="PARTIAL"),
            "DECISION_BLOCKER": choose(row, "DECISION_BLOCKER", default="REVIEW_REQUIRED"),
            "REVIEW_PRIORITY": choose(row, "REVIEW_PRIORITY", default="MEDIUM"),
            "REVIEW_STATUS": choose(row, "REVIEW_STATUS", default="PENDING"),
            "ESCALATION_LEVEL": choose(row, "REQUIRED_ESCALATION_LEVEL", "COCKPIT_ESCALATION_LEVEL", default="STANDARD_REVIEW"),
            "COCKPIT_GOVERNANCE_ALERT": alert_level(row, "GOVERNANCE"),
            "COCKPIT_DRIFT_ALERT": alert_level(row, "DRIFT"),
            "COCKPIT_PROCUREMENT_ALERT": alert_level(row, "PROCUREMENT"),
            "COCKPIT_TECHNICAL_ALERT": alert_level(row, "TECHNICAL"),
            "COCKPIT_BLOCK_IMPORT_ALERT": alert_level(row, "BLOCK_IMPORT"),
            "CONFIDENCE_BADGE": badge_for_confidence(choose(row, "CONFIDENCE_LEVEL", default="LOW")),
            "PROCUREMENT_BADGE": badge_for_blocker(choose(row, "DECISION_BLOCKER", default="REVIEW_REQUIRED")),
            "DRIFT_BADGE": badge_for_drift(choose(row, "DRIFT_LEVEL", default="MEDIUM")),
            "RISK_BADGE": badge_for_blocker(choose(row, "DECISION_BLOCKER", default="REVIEW_REQUIRED")),
            "REVIEW_BADGE": badge_for_blocker(choose(row, "REVIEW_PRIORITY", default="MEDIUM")),
            "MANUAL_VALIDATION_REQUIRED": choose(row, "COCKPIT_MANUAL_VALIDATION", default=True),
            "PROCUREMENT_EXPLANATION": choose(row, "PROCUREMENT_EXPLANATION", default=""),
        }
        rows.append(item)
        explain_rows.append(explainability(row))

    family_kpis = build_family_kpis(rows, families)
    return rows, family_kpis, explain_rows


def build_family_kpis(rows: list[dict[str, Any]], families: dict[str, dict[str, Any]]) -> list[dict[str, Any]]:
    by_family: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        by_family.setdefault(as_text(row["FAMILY_NAME"]), []).append(row)
    kpis: list[dict[str, Any]] = []
    for family, items in sorted(by_family.items()):
        family_row = families.get(family, {})
        kpis.append(
            {
                "FAMILY_NAME": family,
                "FAMILY_STATUS": choose(family_row, "FAMILY_STATUS", default="UNKNOWN"),
                "READINESS_SCORE": choose(family_row, "INTEGRATION_READINESS_SCORE", default=""),
                "REVIEW_REQUIRED_COUNT": len([item for item in items if item["REVIEW_STATUS"] in {"PENDING", "IN_REVIEW"}]),
                "BLOCK_IMPORT_COUNT": len([item for item in items if item["DECISION_BLOCKER"] == "BLOCK_IMPORT"]),
                "DRIFT_CRITICAL_COUNT": len([item for item in items if item["DRIFT_LEVEL"] == "CRITICAL"]),
                "VERIFIED_REFERENCE_COUNT": len([item for item in items if item["CONFIDENCE_LEVEL"] == "HIGH"]),
                "HIGH_RISK_COUNT": len([item for item in items if item["DECISION_BLOCKER"] == "HIGH_RISK"]),
                "CRITICAL_ALERT_COUNT": len([item for item in items if item["COCKPIT_GOVERNANCE_ALERT"] == "CRITICAL"]),
            }
        )
    return kpis


def global_kpis(cockpit_rows: list[dict[str, Any]], family_kpis: list[dict[str, Any]]) -> list[dict[str, Any]]:
    multi = load_json(MULTI_STATS).get("global_kpis", {})
    workbench = load_json(WORKBENCH_STATS)
    block_import = len([row for row in cockpit_rows if row["DECISION_BLOCKER"] == "BLOCK_IMPORT"])
    high_risk = len([row for row in cockpit_rows if row["DECISION_BLOCKER"] == "HIGH_RISK"])
    escalation_count = len([row for row in cockpit_rows if row["ESCALATION_LEVEL"] != "STANDARD_REVIEW"])
    values = {
        "GLOBAL_GOVERNANCE_SCORE": multi.get("GLOBAL_GOVERNANCE_SCORE", 0),
        "GLOBAL_CONFIDENCE_SCORE": multi.get("GLOBAL_CONFIDENCE_SCORE", 0),
        "GLOBAL_PROCUREMENT_SCORE": multi.get("GLOBAL_PROCUREMENT_SCORE", 0),
        "GLOBAL_DRIFT_SCORE": multi.get("GLOBAL_DRIFT_SCORE", 0),
        "GLOBAL_TCO_SCORE": multi.get("GLOBAL_TCO_SCORE", 0),
        "GLOBAL_REVIEW_BACKLOG": workbench.get("review_items", len(cockpit_rows)),
        "GLOBAL_ESCALATION_COUNT": escalation_count,
        "GLOBAL_BLOCK_IMPORT_COUNT": block_import,
        "GLOBAL_HIGH_RISK_COUNT": high_risk,
        "FAMILY_COUNT": len(family_kpis),
    }
    return [{"KPI": key, "VALUE": value} for key, value in values.items()]


def workflow_timeline(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    timeline: list[dict[str, Any]] = []
    for state in WORKFLOW_STATES:
        state_rows = [row for row in rows if row["REVIEW_STATUS"] == state]
        timeline.append(
            {
                "TIMELINE_DATE": TODAY,
                "WORKFLOW_STATE": state,
                "REFERENCE_COUNT": len(state_rows),
                "CRITICAL_COUNT": len([row for row in state_rows if row["REVIEW_PRIORITY"] == "CRITICAL"]),
                "HIGH_COUNT": len([row for row in state_rows if row["REVIEW_PRIORITY"] == "HIGH"]),
                "MANUAL_VALIDATION_REQUIRED": len([row for row in state_rows if row["MANUAL_VALIDATION_REQUIRED"]]),
            }
        )
    return timeline


def escalation_matrix(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    matrix: list[dict[str, Any]] = []
    families = sorted({row["FAMILY_NAME"] for row in rows})
    escalations = sorted({row["ESCALATION_LEVEL"] for row in rows})
    for family in families:
        for escalation in escalations:
            items = [row for row in rows if row["FAMILY_NAME"] == family and row["ESCALATION_LEVEL"] == escalation]
            if not items:
                continue
            matrix.append(
                {
                    "FAMILY_NAME": family,
                    "ESCALATION_LEVEL": escalation,
                    "REFERENCE_COUNT": len(items),
                    "CRITICAL": len([row for row in items if row["REVIEW_PRIORITY"] == "CRITICAL"]),
                    "HIGH": len([row for row in items if row["REVIEW_PRIORITY"] == "HIGH"]),
                    "AVERAGE_PROCUREMENT_SCORE": round(statistics.mean(as_float(row["PROCUREMENT_SCORE"], 0) for row in items), 2),
                    "ALERT_LEVEL": severity_level(escalation),
                }
            )
    return matrix


def audit_timeline() -> list[dict[str, Any]]:
    audit_rows = read_xlsx_rows(REVIEW_AUDIT)
    overrides = [row for row in read_xlsx_rows(OVERRIDE_LOG) if as_text(choose(row, "STATUS", default="")) != "TEMPLATE_NO_ACTIVE_OVERRIDE"]
    comments = [row for row in read_xlsx_rows(COMMENTS) if as_text(choose(row, "STATUS", default="")) != "TEMPLATE_NO_ACTIVE_COMMENT"]
    timeline: list[dict[str, Any]] = []
    for row in audit_rows:
        timeline.append(
            {
                "EVENT_DATE": choose(row, "AUDIT_DATE", default=TODAY),
                "EVENT_TYPE": choose(row, "AUDIT_EVENT", default="AUDIT_EVENT"),
                "REFERENCE_ID": choose(row, "REFERENCE_ID", default=""),
                "FAMILY_NAME": choose(row, "FAMILY_NAME", default=""),
                "ACTOR": choose(row, "AUDIT_BY", default="SYSTEM"),
                "DECISION": choose(row, "REVIEW_STATUS", default=""),
                "JUSTIFICATION": choose(row, "JUSTIFICATION", default=""),
                "RISK_ACCEPTED": choose(row, "RISK_ACCEPTED", default=False),
            }
        )
    for row in overrides:
        timeline.append(
            {
                "EVENT_DATE": choose(row, "OVERRIDE_DATE", default=TODAY),
                "EVENT_TYPE": "GOVERNANCE_OVERRIDE",
                "REFERENCE_ID": choose(row, "REFERENCE_ID", default=""),
                "FAMILY_NAME": "",
                "ACTOR": choose(row, "OVERRIDE_BY", default=""),
                "DECISION": f"{choose(row, 'OLD_DECISION', default='')} -> {choose(row, 'NEW_DECISION', default='')}",
                "JUSTIFICATION": choose(row, "OVERRIDE_REASON", default=""),
                "RISK_ACCEPTED": choose(row, "RISK_ACCEPTED", default=False),
            }
        )
    for row in comments:
        timeline.append(
            {
                "EVENT_DATE": choose(row, "COMMENT_DATE", default=TODAY),
                "EVENT_TYPE": "GOVERNANCE_COMMENT",
                "REFERENCE_ID": choose(row, "REFERENCE_ID", default=""),
                "FAMILY_NAME": choose(row, "FAMILY_NAME", default=""),
                "ACTOR": choose(row, "COMMENT_BY", default=""),
                "DECISION": choose(row, "COMMENT_TYPE", default=""),
                "JUSTIFICATION": choose(row, "COMMENT_TEXT", default=""),
                "RISK_ACCEPTED": False,
            }
        )
    return timeline


def component_specs() -> list[dict[str, Any]]:
    return [
        {"COMPONENT": "GovernanceKpiStrip", "PURPOSE": "Afficher scores globaux, backlog, escalations et blockers.", "DATA_SOURCE": "GLOBAL_KPIS", "UX_RULE": "Ne jamais masquer les faibles confiances."},
        {"COMPONENT": "GovernanceReviewQueue", "PURPOSE": "Piloter les references PENDING/IN_REVIEW avec priorites.", "DATA_SOURCE": "COCKPIT_REFERENCES", "UX_RULE": "Tri par CRITICAL puis HIGH."},
        {"COMPONENT": "GovernanceEscalationPanel", "PURPOSE": "Visualiser senior approvals, double validations et revues procurement.", "DATA_SOURCE": "GOVERNANCE_ESCALATION_MATRIX.xlsx", "UX_RULE": "Escalations visibles en premier viewport."},
        {"COMPONENT": "GovernanceDriftHeatmap", "PURPOSE": "Montrer derives marche par famille et reference.", "DATA_SOURCE": "COCKPIT_REFERENCES", "UX_RULE": "DARK_RED seulement pour CRITICAL."},
        {"COMPONENT": "GovernanceConfidenceMatrix", "PURPOSE": "Comparer confidence, procurement et readiness.", "DATA_SOURCE": "FAMILY_KPIS", "UX_RULE": "Aucun signal vert si validation humaine absente."},
        {"COMPONENT": "GovernanceAuditTimeline", "PURPOSE": "Tracer validations, overrides, commentaires et reviewers.", "DATA_SOURCE": "GOVERNANCE_AUDIT_TIMELINE.xlsx", "UX_RULE": "Afficher qui, quand, pourquoi."},
        {"COMPONENT": "GovernanceExplainabilityPanel", "PURPOSE": "Expliquer blocage, escalation, confidence faible et drift.", "DATA_SOURCE": "GOVERNANCE_EXPLAINABILITY_DATA.xlsx", "UX_RULE": "Texte court, orienté decision."},
        {"COMPONENT": "GovernanceFamilyStatusBoard", "PURPOSE": "Board familles avec status, readiness et blockers.", "DATA_SOURCE": "FAMILY_KPIS", "UX_RULE": "BLOCKED/HIGH_RISK predominants visuellement."},
    ]


def badge_specs() -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for badge_type in ["CONFIDENCE_BADGE", "PROCUREMENT_BADGE", "DRIFT_BADGE", "RISK_BADGE", "REVIEW_BADGE"]:
        for color, meaning in BADGE_COLORS.items():
            rows.append({"BADGE_TYPE": badge_type, "COLOR": color, "MEANING": meaning})
    return rows


def write_xlsx(path: Path, sheets: dict[str, list[dict[str, Any]]]) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(sheet_name[:31])
        if rows:
            headers = list(rows[0].keys())
            for row in rows[1:]:
                for key in row:
                    if key not in headers:
                        headers.append(key)
        else:
            headers = ["STATUS"]
            rows = [{"STATUS": "NO_DATA"}]
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header, "") for header in headers])
        style_sheet(ws)
    wb.save(path)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="0F2A3A")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            value = str(cell.value or "").upper()
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if value in {"CRITICAL", "DARK_RED", "BLOCK_IMPORT", "BLOCKED"}:
                cell.fill = PatternFill("solid", fgColor="990000")
                cell.font = Font(color="FFFFFF")
            elif value in {"HIGH", "RED", "HIGH_RISK"}:
                cell.fill = PatternFill("solid", fgColor="F4CCCC")
            elif value in {"MEDIUM", "ORANGE", "PENDING", "REVIEW_REQUIRED"}:
                cell.fill = PatternFill("solid", fgColor="FCE5CD")
            elif value in {"LOW", "GREEN", "APPROVED", "VERIFIED"}:
                cell.fill = PatternFill("solid", fgColor="D9EAD3")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        width = min(65, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        ws.column_dimensions[get_column_letter(column[0].column)].width = width


def main() -> int:
    start = time.perf_counter()
    print("=== SP2I Governance Cockpit Dataset ===")
    try:
        cockpit_rows, family_kpis, explain_rows = build_cockpit_dataset()
        kpi_rows = global_kpis(cockpit_rows, family_kpis)
        workflow_rows = workflow_timeline(cockpit_rows)
        escalation_rows = escalation_matrix(cockpit_rows)
        audit_rows = audit_timeline()

        write_xlsx(
            DATASET_OUT,
            {
                "GLOBAL_KPIS": kpi_rows,
                "FAMILY_KPIS": family_kpis,
                "COCKPIT_REFERENCES": cockpit_rows,
                "BADGE_RULES": badge_specs(),
            },
        )
        write_xlsx(WORKFLOW_OUT, {"WORKFLOW_TIMELINE": workflow_rows})
        write_xlsx(ESCALATION_OUT, {"ESCALATION_MATRIX": escalation_rows})
        write_xlsx(EXPLAINABILITY_OUT, {"EXPLAINABILITY": explain_rows})
        write_xlsx(AUDIT_TIMELINE_OUT, {"AUDIT_TIMELINE": audit_rows})

        stats = {
            "generated_at": datetime.now().replace(microsecond=0).isoformat(),
            "source_queue": QUEUE.name,
            "cockpit_references": len(cockpit_rows),
            "family_kpis": len(family_kpis),
            "workflow_states": count_by(cockpit_rows, "REVIEW_STATUS"),
            "governance_alerts": count_by(cockpit_rows, "COCKPIT_GOVERNANCE_ALERT"),
            "drift_alerts": count_by(cockpit_rows, "COCKPIT_DRIFT_ALERT"),
            "procurement_alerts": count_by(cockpit_rows, "COCKPIT_PROCUREMENT_ALERT"),
            "technical_alerts": count_by(cockpit_rows, "COCKPIT_TECHNICAL_ALERT"),
            "escalations": count_by(cockpit_rows, "ESCALATION_LEVEL"),
            "audit_events": len(audit_rows),
            "components_prepared": len(component_specs()),
            "duration_seconds": round(time.perf_counter() - start, 2),
        }
        STATS_OUT.write_text(json.dumps(stats, indent=2, ensure_ascii=False), encoding="utf-8")
        print(json.dumps(stats, indent=2, ensure_ascii=False))
        print("\nLivrables:")
        for path in [DATASET_OUT, WORKFLOW_OUT, ESCALATION_OUT, EXPLAINABILITY_OUT, AUDIT_TIMELINE_OUT, STATS_OUT]:
            print(f"- {path.name}")
        return 0
    except Exception:
        print("\nERREUR GOVERNANCE COCKPIT DATASET")
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

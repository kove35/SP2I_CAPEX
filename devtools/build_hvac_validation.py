from __future__ import annotations

import hashlib
import json
import math
import re
import sys
import time
import traceback
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "07_API_BACKEND"
DQE_PATH = ROOT / "03_DONNEES_REFERENCE" / "DQE_PROJECT_SP2I.xlsx"
SHEET_NAME = "DQE_CLEAN"

DQE_HVAC = ROOT / "DQE_HVAC_TEST.xlsx"
MASTER_HVAC = ROOT / "MASTER_REFERENCE_HVAC.xlsx"
BENCHMARKS_HVAC = ROOT / "HVAC_BENCHMARKS.xlsx"
PROCUREMENT_HVAC = ROOT / "HVAC_PROCUREMENT_AUDIT.xlsx"
DRIFT_HVAC = ROOT / "HVAC_DRIFT_ANALYSIS.xlsx"
CONFIDENCE_HVAC = ROOT / "HVAC_CONFIDENCE_AUDIT.xlsx"
STATS_PATH = ROOT / "HVAC_VALIDATION_STATS.json"

if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from app.core.equipment_classification_engine import EquipmentClassificationEngine
from app.core.semantic_normalization_engine import SemanticNormalizationEngine


TODAY = date.today().isoformat()
REGION_FACTORS = {
    "CONGO_BRAZZAVILLE": 1.10,
    "CAMEROUN": 1.04,
    "GABON": 1.14,
    "CENTRAL_AFRICA": 1.0,
}

HVAC_KEYWORDS = [
    "split",
    "cassette",
    "gainable",
    "vrv",
    "vrf",
    "ventilation",
    "extraction",
    "climatisation",
    "climatiseur",
    "cuivre frigorifique",
    "conduit",
    "gaine",
    "isolation thermique",
    "thermostat",
    "cta",
    "btu",
]

HVAC_BENCHMARKS = {
    "SPLIT_MURAL": {"min": 350_000, "max": 1_450_000, "moq": "10 U", "fob_factor": 0.58, "lead_time": 45},
    "CASSETTE": {"min": 850_000, "max": 2_600_000, "moq": "5 U", "fob_factor": 0.56, "lead_time": 50},
    "GAINABLE": {"min": 1_500_000, "max": 5_500_000, "moq": "2 U", "fob_factor": 0.55, "lead_time": 55},
    "VRV_VRF": {"min": 8_500_000, "max": 65_000_000, "moq": "1 LOT", "fob_factor": 0.52, "lead_time": 75},
    "VENTILATION": {"min": 250_000, "max": 9_000_000, "moq": "10 U", "fob_factor": 0.60, "lead_time": 50},
    "EXTRACTION": {"min": 300_000, "max": 8_500_000, "moq": "10 U", "fob_factor": 0.60, "lead_time": 50},
    "CUIVRE_FRIGORIFIQUE": {"min": 8_000, "max": 55_000, "moq": "500 ML", "fob_factor": 0.62, "lead_time": 45},
    "CONDUITS": {"min": 15_000, "max": 120_000, "moq": "300 ML", "fob_factor": 0.61, "lead_time": 45},
    "ISOLATION_THERMIQUE": {"min": 6_000, "max": 45_000, "moq": "500 M2", "fob_factor": 0.63, "lead_time": 45},
    "THERMOSTATS": {"min": 25_000, "max": 250_000, "moq": "50 U", "fob_factor": 0.54, "lead_time": 40},
    "HVAC_GENERAL": {"min": 250_000, "max": 20_000_000, "moq": "A_VERIFIER", "fob_factor": 0.58, "lead_time": 55},
}


def normalize_key(value: Any) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"[^A-Za-z0-9]+", "_", text.strip())
    return re.sub(r"_+", "_", text).strip("_").upper()


def normalize_text(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.lower()
    replacements = {
        "clim": "climatisation",
        "v.r.v": "vrv",
        "v.r.f": "vrf",
        "split ac": "split",
        "cuivre frigo": "cuivre frigorifique",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    text = re.sub(r"[^a-z0-9.,]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def number(value: Any, default: float = 0.0) -> float:
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return default if math.isnan(value) else float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return default


def choose(row: dict[str, Any], aliases: list[str], default: Any = "") -> Any:
    for alias in aliases:
        key = normalize_key(alias)
        if key in row and row[key] not in (None, ""):
            return row[key]
    return default


def read_dqe() -> list[dict[str, Any]]:
    wb = load_workbook(DQE_PATH, data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Onglet {SHEET_NAME} introuvable dans {DQE_PATH.name}")
    ws = wb[SHEET_NAME]
    headers = [normalize_key(cell.value or f"COL_{index + 1}") for index, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))]
    rows = []
    for rownum, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        payload = {headers[index]: value for index, value in enumerate(values)}
        if any(value not in (None, "") for value in payload.values()):
            payload["_ROWNUM"] = rownum
            rows.append(payload)
    wb.close()
    return rows


def is_hvac_row(row: dict[str, Any]) -> bool:
    text = normalize_text(" ".join(str(choose(row, field, "")) for field in [["LOT"], ["SOUS_LOT"], ["DESIGNATION"]]))
    return any(keyword in text for keyword in HVAC_KEYWORDS)


def hvac_type(text: str) -> str:
    if "vrv" in text or "vrf" in text:
        return "VRV_VRF"
    if "cassette" in text:
        return "CASSETTE"
    if "gainable" in text:
        return "GAINABLE"
    if "split" in text or "climatiseur" in text or "climatisation" in text:
        return "SPLIT_MURAL"
    if "extraction" in text:
        return "EXTRACTION"
    if "ventilation" in text or "cta" in text:
        return "VENTILATION"
    if "cuivre frigorifique" in text:
        return "CUIVRE_FRIGORIFIQUE"
    if "conduit" in text or "gaine" in text:
        return "CONDUITS"
    if "isolation thermique" in text:
        return "ISOLATION_THERMIQUE"
    if "thermostat" in text:
        return "THERMOSTATS"
    return "HVAC_GENERAL"


def extract_power(text: str) -> tuple[float, str]:
    btu = re.search(r"(\d+(?:[.,]\d+)?)\s*btu", text)
    if btu:
        return number(btu.group(1)), "BTU"
    kw = re.search(r"(\d+(?:[.,]\d+)?)\s*kw", text)
    if kw:
        return number(kw.group(1)), "KW"
    cv = re.search(r"(\d+(?:[.,]\d+)?)\s*cv", text)
    if cv:
        return number(cv.group(1)), "CV"
    return 0, ""


def reference_id(normalized: str, unit: str) -> str:
    digest = hashlib.sha1(f"HVAC|{normalized}|{unit}".encode("utf-8")).hexdigest()[:8].upper()
    return f"SP2I-HVA-{digest}"


def status(price: float, price_min: float, price_max: float) -> str:
    if price <= 0:
        return "CRITICAL_ZERO_PRICE"
    if price < price_min * 0.2:
        return "CRITICAL_LOW_MAGNITUDE"
    if price < price_min:
        return "WARNING_BELOW_BENCHMARK"
    if price > price_max * 2:
        return "CRITICAL_HIGH_MAGNITUDE"
    if price > price_max:
        return "WARNING_ABOVE_BENCHMARK"
    return "OK"


def confidence(price_status: str, supplier_verified: bool, fob_status: str, drift: float) -> str:
    if supplier_verified and fob_status == "VERIFIED" and price_status == "OK" and drift < 35:
        return "HIGH"
    if price_status == "OK" and fob_status in {"PARTIAL", "VERIFIED"} and drift < 65:
        return "MEDIUM"
    return "LOW"


def drift_score(price_min: float, price_max: float, price_status: str) -> float:
    ratio = price_max / max(price_min, 1)
    score = 15
    if ratio >= 8:
        score += 55
    elif ratio >= 4:
        score += 35
    elif ratio >= 2.5:
        score += 18
    if price_status.startswith("CRITICAL"):
        score += 25
    elif price_status.startswith("WARNING"):
        score += 12
    return round(min(score, 100), 2)


def drift_level(score: float) -> str:
    if score >= 80:
        return "CRITICAL"
    if score >= 60:
        return "HIGH"
    if score >= 35:
        return "MEDIUM"
    return "LOW"


def build_hvac() -> dict[str, list[dict[str, Any]]]:
    classifier = EquipmentClassificationEngine()
    normalizer = SemanticNormalizationEngine()
    dqe_rows = []
    master_rows = []
    benchmark_rows = []
    procurement_rows = []
    drift_rows = []
    confidence_rows = []

    for row in read_dqe():
        if not is_hvac_row(row):
            continue
        designation = str(choose(row, ["DESIGNATION"], ""))
        lot = str(choose(row, ["LOT"], ""))
        sous_lot = str(choose(row, ["SOUS_LOT"], ""))
        type_ligne = str(choose(row, ["TYPE_LIGNE"], "")).upper()
        unit = str(choose(row, ["UNITE"], "U") or "U").upper()
        price = number(choose(row, ["PU_ESTIME_LOCAL_FCFA"], 0))
        normalized = normalizer.normalize(designation)["normalized_designation"].upper()
        text = normalize_text(f"{lot} {sous_lot} {designation}")
        type_hvac = hvac_type(text)
        power, power_unit = extract_power(text)
        benchmark = HVAC_BENCHMARKS[type_hvac]
        price_min = benchmark["min"]
        price_max = benchmark["max"]
        price_status = status(price, price_min, price_max)
        drift = drift_score(price_min, price_max, price_status)
        supplier_verified = False
        fob_status = "PARTIAL" if price > 0 and price_status != "CRITICAL_LOW_MAGNITUDE" else "UNVERIFIED"
        conf = confidence(price_status, supplier_verified, fob_status, drift)
        classification = classifier.classify_line({
            "designation": designation,
            "LOT": lot,
            "SOUS_LOT": sous_lot,
            "TYPE_LIGNE": type_ligne,
            "region": "CONGO_BRAZZAVILLE",
            "quality": "MEDIUM",
        })
        ref_id = reference_id(normalized, unit)
        dqe_item = {
            "SOURCE_ROW": row["_ROWNUM"],
            "LOT": lot,
            "SOUS_LOT": sous_lot,
            "TYPE_LIGNE": type_ligne,
            "DESIGNATION": designation,
            "UNITE": unit,
            "QTE": choose(row, ["QTE"], ""),
            "PU_ESTIME_LOCAL_FCFA": price,
            "PUISSANCE": power,
            "UNITE_PUISSANCE": power_unit,
            "TYPE_HVAC": type_hvac,
            "TECHNOLOGIE": "VRV/VRF" if type_hvac == "VRV_VRF" else "DX_STANDARD" if type_hvac in {"SPLIT_MURAL", "CASSETTE", "GAINABLE"} else "AIRFLOW",
            "MARQUE_REFERENCE": "A_VERIFIER",
        }
        dqe_rows.append(dqe_item)
        master = {
            "REFERENCE_ID": ref_id,
            "DESIGNATION_NORMALISEE": normalized,
            "RAW_DESIGNATION": designation,
            "FAMILLE": "HVAC",
            "SOUS_FAMILLE": type_hvac,
            "TYPE_EQUIPEMENT": classification.get("equipment_type") or type_hvac,
            "UNITE": unit,
            "PUISSANCE": power,
            "UNITE_PUISSANCE": power_unit,
            "TYPE_HVAC": type_hvac,
            "TECHNOLOGIE": dqe_item["TECHNOLOGIE"],
            "MARQUE_REFERENCE": "A_VERIFIER",
            "PRICE_MIN_FCFA": price_min,
            "PRICE_MAX_FCFA": price_max,
            "PU_CHINE_FOB_FCFA": round(price_min * benchmark["fob_factor"], 2),
            "IMPORTABILITY": "HIGH" if type_ligne in {"ARTICLE", "FOURNITURE", "EQUIPEMENT"} else "LOW",
            "RISK_LEVEL": "HIGH" if conf == "LOW" else "MEDIUM",
            "MOQ": benchmark["moq"],
            "INCOTERM": "FOB",
            "FOURNISSEUR_CHINE": "A_VERIFIER_GOUVERNANCE_HVAC",
            "BENCHMARK_REGION": "CONGO_BRAZZAVILLE|CAMEROUN|GABON|CENTRAL_AFRICA",
            "QUALITY_LEVEL": "MEDIUM",
            "CONFIDENCE_LEVEL": conf,
            "LAST_VALIDATION_DATE": TODAY,
            "SOURCE_REFERENCE": "DQE_PROJECT_SP2I.xlsx::DQE_CLEAN::HVAC_TEST",
        }
        master_rows.append(master)
        for region, factor in REGION_FACTORS.items():
            benchmark_rows.append({
                "REFERENCE_ID": ref_id,
                "DESIGNATION_NORMALISEE": normalized,
                "TYPE_HVAC": type_hvac,
                "REGION": region,
                "QUALITY_LEVEL": "MEDIUM",
                "PRICE_MIN_FCFA": round(price_min * factor, 2),
                "PRICE_MAX_FCFA": round(price_max * factor, 2),
                "BENCHMARK_CONFIDENCE": "MEDIUM" if price_status == "OK" else "LOW",
                "PRICE_STATUS": price_status,
            })
        procurement_rows.append({
            "REFERENCE_ID": ref_id,
            "DESIGNATION_NORMALISEE": normalized,
            "TYPE_HVAC": type_hvac,
            "IMPORTABILITY": master["IMPORTABILITY"],
            "FOURNISSEUR_CHINE": master["FOURNISSEUR_CHINE"],
            "SUPPLIER_VERIFIED": False,
            "FOB_VALIDATION_STATUS": fob_status,
            "MOQ": benchmark["moq"],
            "LEAD_TIME_DAYS": benchmark["lead_time"],
            "PROCUREMENT_REVIEW_REQUIRED": "YES",
            "PROCUREMENT_WARNING": "Fournisseur HVAC et FOB a verifier avant decision import.",
        })
        drift_rows.append({
            "REFERENCE_ID": ref_id,
            "DESIGNATION_NORMALISEE": normalized,
            "TYPE_HVAC": type_hvac,
            "PRICE_STATUS": price_status,
            "MARKET_DRIFT_SCORE": drift,
            "DRIFT_ALERT_LEVEL": drift_level(drift),
            "DRIFT_FACTORS": "inflation|usd|fret_maritime|energie|cuivre|sourcing_chine",
        })
        confidence_rows.append({
            "REFERENCE_ID": ref_id,
            "DESIGNATION_NORMALISEE": normalized,
            "TYPE_HVAC": type_hvac,
            "CLASSIFICATION_CONFIDENCE": classification.get("classification_confidence"),
            "BENCHMARK_STATUS": price_status,
            "FOB_VALIDATION_STATUS": fob_status,
            "SUPPLIER_VERIFIED": False,
            "DRIFT_ALERT_LEVEL": drift_level(drift),
            "CONFIDENCE_LEVEL": conf,
            "COCKPIT_CONFIDENCE_BADGE": {"HIGH": "GREEN", "MEDIUM": "ORANGE", "LOW": "RED"}[conf],
            "CONFIDENCE_REASON": "HIGH bloque tant que fournisseur, benchmark, FOB et drift ne sont pas verifies.",
        })
    return {
        "dqe": dqe_rows,
        "master": master_rows,
        "benchmarks": benchmark_rows,
        "procurement": procurement_rows,
        "drift": drift_rows,
        "confidence": confidence_rows,
    }


def style_sheet(ws) -> None:
    fill = PatternFill("solid", fgColor="0F172A")
    font = Font(color="F8FAFC", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    for column in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(max(max_len + 2, 12), 64)
    ws.freeze_panes = "A2"


def write_xlsx(path: Path, sheet_name: str, rows: list[dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    if rows:
        headers = list(rows[0].keys())
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header) for header in headers])
    else:
        ws.append(["STATUS"])
        ws.append(["NO_HVAC_ROWS"])
    style_sheet(ws)
    wb.save(path)


def count_by(rows: list[dict[str, Any]], key: str) -> dict[str, int]:
    result: dict[str, int] = {}
    for row in rows:
        value = str(row.get(key) or "UNKNOWN")
        result[value] = result.get(value, 0) + 1
    return result


def main() -> int:
    start = time.perf_counter()
    print("=== SP2I HVAC Validation Pipeline ===")
    try:
        outputs = build_hvac()
        write_xlsx(DQE_HVAC, "DQE_HVAC_TEST", outputs["dqe"])
        write_xlsx(MASTER_HVAC, "MASTER_REFERENCE_HVAC", outputs["master"])
        write_xlsx(BENCHMARKS_HVAC, "HVAC_BENCHMARKS", outputs["benchmarks"])
        write_xlsx(PROCUREMENT_HVAC, "HVAC_PROCUREMENT", outputs["procurement"])
        write_xlsx(DRIFT_HVAC, "HVAC_DRIFT", outputs["drift"])
        write_xlsx(CONFIDENCE_HVAC, "HVAC_CONFIDENCE", outputs["confidence"])
        stats = {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "source_file": str(DQE_PATH),
            "sheet": SHEET_NAME,
            "family": "HVAC",
            "hvac_rows": len(outputs["dqe"]),
            "master_references": len(outputs["master"]),
            "benchmark_rows": len(outputs["benchmarks"]),
            "confidence_distribution": count_by(outputs["confidence"], "CONFIDENCE_LEVEL"),
            "drift_alerts": count_by(outputs["drift"], "DRIFT_ALERT_LEVEL"),
            "procurement_review_required": len([row for row in outputs["procurement"] if row["PROCUREMENT_REVIEW_REQUIRED"] == "YES"]),
            "high_policy": "No HIGH without supplier, benchmark, FOB and acceptable drift validation.",
            "duration_seconds": round(time.perf_counter() - start, 2),
        }
        STATS_PATH.write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(stats, ensure_ascii=False, indent=2))
        print("\nLivrables:")
        for path in [DQE_HVAC, MASTER_HVAC, BENCHMARKS_HVAC, PROCUREMENT_HVAC, DRIFT_HVAC, CONFIDENCE_HVAC, STATS_PATH]:
            print(f"- {path.name}")
        return 0
    except Exception:
        print("HVAC VALIDATION FAILED")
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

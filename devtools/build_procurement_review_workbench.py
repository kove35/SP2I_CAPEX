from __future__ import annotations

import json
import statistics
import time
import traceback
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
TODAY = date.today().isoformat()

CANDIDATES = ROOT / "MASTER_REFERENCE_ENTERPRISE_CANDIDATES.xlsx"
FAMILY_INDEX = ROOT / "FAMILY_GOVERNANCE_INDEX.xlsx"

QUEUE_OUT = ROOT / "PROCUREMENT_REVIEW_QUEUE.xlsx"
OVERRIDE_OUT = ROOT / "GOVERNANCE_OVERRIDE_LOG.xlsx"
COMMENTS_OUT = ROOT / "PROCUREMENT_GOVERNANCE_COMMENTS.xlsx"
DASHBOARD_OUT = ROOT / "PROCUREMENT_REVIEW_DASHBOARD_DATA.xlsx"
AUDIT_OUT = ROOT / "PROCUREMENT_REVIEW_AUDIT.xlsx"
STATS_OUT = ROOT / "PROCUREMENT_WORKBENCH_STATS.json"

REVIEW_ACTIONS = ["PENDING", "IN_REVIEW", "APPROVED", "REJECTED", "ESCALATED", "NEEDS_MORE_DATA"]
HUMAN_ACTIONS = [
    "VALIDATE_SUPPLIER",
    "VALIDATE_BENCHMARK",
    "VALIDATE_FOB",
    "APPROVE_IMPORT",
    "REJECT_IMPORT",
    "OVERRIDE_AI_RECOMMENDATION",
    "REQUEST_TECHNICAL_REVIEW",
    "REQUEST_PROCUREMENT_REVIEW",
    "ADD_GOVERNANCE_COMMENT",
]


def normalize_header(value: Any) -> str:
    return str(value or "").strip().upper().replace(" ", "_")


def as_text(value: Any) -> str:
    return str(value or "").strip()


def as_float(value: Any, default: float = 0.0) -> float:
    try:
        if value in (None, ""):
            return default
        return float(value)
    except (TypeError, ValueError):
        return default


def read_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [normalize_header(cell) for cell in rows[0]]
    data: list[dict[str, Any]] = []
    for values in rows[1:]:
        row = {headers[index]: value for index, value in enumerate(values) if index < len(headers)}
        if any(value not in (None, "") for value in row.values()):
            data.append(row)
    return data


def choose(row: dict[str, Any], *keys: str, default: Any = "") -> Any:
    for key in keys:
        value = row.get(normalize_header(key))
        if value not in (None, ""):
            return value
    return default


def family_index_map() -> dict[str, dict[str, Any]]:
    return {as_text(row.get("FAMILY_NAME")).upper(): row for row in read_xlsx_rows(FAMILY_INDEX)}


def priority_score(row: dict[str, Any], family: dict[str, Any]) -> float:
    confidence = as_text(choose(row, "CONFIDENCE_LEVEL", default="LOW")).upper()
    drift = as_text(choose(row, "DRIFT_LEVEL", "COCKPIT_DRIFT_ALERT", default="MEDIUM")).upper()
    blocker = as_text(choose(row, "DECISION_BLOCKER", default="REVIEW_REQUIRED")).upper()
    technical = as_text(choose(row, "TECHNICAL_VALIDATION", default="PARTIAL")).upper()
    readiness = as_float(choose(row, "INTEGRATION_READINESS_SCORE", default=0))
    family_procurement = as_float(family.get("PROCUREMENT_SCORE"), 40.0)
    blocker_rate = as_float(family.get("DECISION_BLOCKER_RATE"), 0.0)

    score = 25.0
    if confidence == "LOW":
        score += 18.0
    elif confidence == "MEDIUM":
        score += 8.0
    if drift == "CRITICAL":
        score += 28.0
    elif drift == "HIGH":
        score += 18.0
    elif drift == "MEDIUM":
        score += 8.0
    if blocker == "BLOCK_IMPORT":
        score += 28.0
    elif blocker == "HIGH_RISK":
        score += 24.0
    elif blocker == "TECHNICAL_VALIDATION_REQUIRED":
        score += 20.0
    elif blocker == "REVIEW_REQUIRED":
        score += 12.0
    if technical in {"REJECTED", "UNVERIFIED"}:
        score += 18.0
    elif technical == "PARTIAL":
        score += 8.0
    if readiness < 50:
        score += 12.0
    if family_procurement < 45:
        score += 10.0
    score += min(12.0, blocker_rate * 12.0)
    return round(min(100.0, score), 2)


def priority_label(score: float) -> str:
    if score >= 85:
        return "CRITICAL"
    if score >= 65:
        return "HIGH"
    if score >= 45:
        return "MEDIUM"
    return "LOW"


def escalation_level(row: dict[str, Any], priority: str) -> str:
    blocker = as_text(choose(row, "DECISION_BLOCKER", default="")).upper()
    drift = as_text(choose(row, "DRIFT_LEVEL", default="")).upper()
    if blocker == "BLOCK_IMPORT":
        return "SENIOR_PROCUREMENT_APPROVAL"
    if blocker == "HIGH_RISK":
        return "DOUBLE_VALIDATION"
    if drift == "CRITICAL":
        return "EXECUTIVE_ESCALATION"
    if priority == "CRITICAL":
        return "SENIOR_REVIEW"
    if priority == "HIGH":
        return "PROCUREMENT_REVIEW"
    return "STANDARD_REVIEW"


def explanation(row: dict[str, Any], priority: str) -> str:
    reasons: list[str] = []
    confidence = as_text(choose(row, "CONFIDENCE_LEVEL", default="LOW")).upper()
    drift = as_text(choose(row, "DRIFT_LEVEL", default="MEDIUM")).upper()
    blocker = as_text(choose(row, "DECISION_BLOCKER", default="REVIEW_REQUIRED")).upper()
    technical = as_text(choose(row, "TECHNICAL_VALIDATION", default="PARTIAL")).upper()

    if confidence == "LOW":
        reasons.append("confiance faible")
    if drift == "CRITICAL":
        reasons.append("drift critique")
    elif drift == "HIGH":
        reasons.append("drift eleve")
    if blocker == "BLOCK_IMPORT":
        reasons.append("import bloque")
    elif blocker == "HIGH_RISK":
        reasons.append("risque decisionnel eleve")
    elif blocker == "TECHNICAL_VALIDATION_REQUIRED":
        reasons.append("validation technique requise")
    elif blocker == "REVIEW_REQUIRED":
        reasons.append("revue procurement requise")
    if technical in {"PARTIAL", "REJECTED", "UNVERIFIED"}:
        reasons.append(f"validation technique {technical.lower()}")
    if not reasons:
        reasons.append("validation humaine requise avant integration master")
    return f"Priorite {priority}: " + " + ".join(reasons) + ". Decision automatique interdite."


def recommended_actions(row: dict[str, Any]) -> str:
    actions = ["ADD_GOVERNANCE_COMMENT"]
    confidence = as_text(choose(row, "CONFIDENCE_LEVEL", default="LOW")).upper()
    drift = as_text(choose(row, "DRIFT_LEVEL", default="MEDIUM")).upper()
    blocker = as_text(choose(row, "DECISION_BLOCKER", default="REVIEW_REQUIRED")).upper()
    technical = as_text(choose(row, "TECHNICAL_VALIDATION", default="PARTIAL")).upper()

    if confidence == "LOW":
        actions.extend(["VALIDATE_SUPPLIER", "VALIDATE_BENCHMARK", "VALIDATE_FOB"])
    if drift in {"HIGH", "CRITICAL"}:
        actions.append("REQUEST_PROCUREMENT_REVIEW")
    if technical in {"PARTIAL", "REJECTED", "UNVERIFIED"}:
        actions.append("REQUEST_TECHNICAL_REVIEW")
    if blocker == "BLOCK_IMPORT":
        actions.append("REJECT_IMPORT")
    elif blocker == "HIGH_RISK":
        actions.append("ESCALATE_DOUBLE_VALIDATION")
    else:
        actions.append("NEEDS_MORE_DATA")
    return "|".join(dict.fromkeys(actions))


def cockpit_alert(priority: str, escalation: str) -> str:
    if priority == "CRITICAL" or escalation in {"EXECUTIVE_ESCALATION", "SENIOR_PROCUREMENT_APPROVAL"}:
        return "CRITICAL_MANUAL_REVIEW"
    if priority == "HIGH":
        return "HIGH_PRIORITY_REVIEW"
    if priority == "MEDIUM":
        return "STANDARD_REVIEW_REQUIRED"
    return "LOW_PRIORITY_REVIEW"


def build_queue() -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    families = family_index_map()
    candidates = read_xlsx_rows(CANDIDATES)
    queue: list[dict[str, Any]] = []
    audit: list[dict[str, Any]] = []
    dashboard: list[dict[str, Any]] = []

    for row in candidates:
        family_name = as_text(choose(row, "FAMILY_NAME", default="UNKNOWN")).upper()
        family = families.get(family_name, {})
        score = priority_score(row, family)
        priority = priority_label(score)
        escalation = escalation_level(row, priority)
        ref_id = as_text(choose(row, "REFERENCE_ID", default="UNKNOWN_REF"))
        review_status = "PENDING"
        manual = True
        item = {
            "REFERENCE_ID": ref_id,
            "FAMILY_NAME": family_name,
            "DESIGNATION_NORMALISEE": choose(row, "DESIGNATION_NORMALISEE", default=""),
            "CONFIDENCE_LEVEL": choose(row, "CONFIDENCE_LEVEL", default="LOW"),
            "PROCUREMENT_SCORE": family.get("PROCUREMENT_SCORE", ""),
            "DRIFT_LEVEL": choose(row, "DRIFT_LEVEL", "COCKPIT_DRIFT_ALERT", default="MEDIUM"),
            "TECHNICAL_VALIDATION": choose(row, "TECHNICAL_VALIDATION", default="PARTIAL"),
            "DECISION_BLOCKER": choose(row, "DECISION_BLOCKER", default="REVIEW_REQUIRED"),
            "REVIEW_PRIORITY": priority,
            "REVIEW_PRIORITY_SCORE": score,
            "REVIEW_STATUS": review_status,
            "ASSIGNED_REVIEWER": "UNASSIGNED",
            "LAST_REVIEW_DATE": "",
            "REQUIRED_ESCALATION_LEVEL": escalation,
            "HUMAN_ACTIONS_AVAILABLE": "|".join(HUMAN_ACTIONS),
            "RECOMMENDED_ACTIONS": recommended_actions(row),
            "PROCUREMENT_EXPLANATION": explanation(row, priority),
            "COCKPIT_REVIEW_STATUS": review_status,
            "COCKPIT_ESCALATION_LEVEL": escalation,
            "COCKPIT_GOVERNANCE_ALERT": cockpit_alert(priority, escalation),
            "COCKPIT_MANUAL_VALIDATION": manual,
        }
        queue.append(item)
        audit.append(
            {
                "AUDIT_ID": f"AUD-{len(audit) + 1:05d}",
                "REFERENCE_ID": ref_id,
                "FAMILY_NAME": family_name,
                "AUDIT_EVENT": "REVIEW_QUEUE_CREATED",
                "REVIEW_STATUS": review_status,
                "REVIEW_PRIORITY": priority,
                "ESCALATION_LEVEL": escalation,
                "AUDIT_DATE": TODAY,
                "AUDIT_BY": "SYSTEM_GOVERNANCE_WORKBENCH",
                "RISK_ACCEPTED": False,
                "JUSTIFICATION": item["PROCUREMENT_EXPLANATION"],
            }
        )
        dashboard.append(
            {
                "FAMILY_NAME": family_name,
                "REFERENCE_ID": ref_id,
                "REVIEW_STATUS": review_status,
                "REVIEW_PRIORITY": priority,
                "ESCALATION_LEVEL": escalation,
                "CONFIDENCE_LEVEL": item["CONFIDENCE_LEVEL"],
                "DRIFT_LEVEL": item["DRIFT_LEVEL"],
                "DECISION_BLOCKER": item["DECISION_BLOCKER"],
                "COCKPIT_REVIEW_STATUS": item["COCKPIT_REVIEW_STATUS"],
                "COCKPIT_ESCALATION_LEVEL": item["COCKPIT_ESCALATION_LEVEL"],
                "COCKPIT_GOVERNANCE_ALERT": item["COCKPIT_GOVERNANCE_ALERT"],
                "COCKPIT_MANUAL_VALIDATION": item["COCKPIT_MANUAL_VALIDATION"],
            }
        )

    queue.sort(key=lambda item: (-as_float(item["REVIEW_PRIORITY_SCORE"]), item["FAMILY_NAME"], item["REFERENCE_ID"]))
    return queue, audit, dashboard


def build_override_log() -> list[dict[str, Any]]:
    return [
        {
            "OVERRIDE_ID": "TEMPLATE-0001",
            "REFERENCE_ID": "",
            "OLD_DECISION": "",
            "NEW_DECISION": "",
            "OVERRIDE_REASON": "",
            "OVERRIDE_BY": "",
            "OVERRIDE_DATE": "",
            "RISK_ACCEPTED": False,
            "APPROVAL_LEVEL": "SENIOR_PROCUREMENT_REQUIRED",
            "STATUS": "TEMPLATE_NO_ACTIVE_OVERRIDE",
        }
    ]


def build_comments_template() -> list[dict[str, Any]]:
    return [
        {
            "COMMENT_ID": "TEMPLATE-0001",
            "REFERENCE_ID": "",
            "FAMILY_NAME": "",
            "COMMENT_TYPE": "PROCUREMENT|FINANCE|TECHNICAL|REVIEWER",
            "COMMENT_TEXT": "",
            "COMMENT_BY": "",
            "COMMENT_DATE": "",
            "FOLLOW_UP_REQUIRED": True,
            "STATUS": "TEMPLATE_NO_ACTIVE_COMMENT",
        }
    ]


def dashboard_summary(queue: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_family: dict[str, list[dict[str, Any]]] = {}
    for item in queue:
        by_family.setdefault(item["FAMILY_NAME"], []).append(item)

    rows: list[dict[str, Any]] = []
    for family, items in sorted(by_family.items()):
        priorities = {key: sum(1 for item in items if item["REVIEW_PRIORITY"] == key) for key in ["CRITICAL", "HIGH", "MEDIUM", "LOW"]}
        escalations = sum(1 for item in items if item["REQUIRED_ESCALATION_LEVEL"] != "STANDARD_REVIEW")
        rows.append(
            {
                "FAMILY_NAME": family,
                "PENDING_REVIEWS": len(items),
                "CRITICAL": priorities["CRITICAL"],
                "HIGH": priorities["HIGH"],
                "MEDIUM": priorities["MEDIUM"],
                "LOW": priorities["LOW"],
                "ESCALATIONS_REQUIRED": escalations,
                "AVERAGE_PRIORITY_SCORE": round(statistics.mean(as_float(item["REVIEW_PRIORITY_SCORE"]) for item in items), 2),
                "MANUAL_VALIDATION_REQUIRED": len(items),
            }
        )
    return rows


def write_xlsx(path: Path, sheets: dict[str, list[dict[str, Any]]]) -> None:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(sheet_name[:31])
        if rows:
            headers = list(rows[0].keys())
            for row in rows[1:]:
                for key in row:
                    if key not in headers:
                        headers.append(key)
        else:
            headers = ["STATUS"]
            rows = [{"STATUS": "NO_DATA"}]
        ws.append(headers)
        for row in rows:
            ws.append([row.get(header, "") for header in headers])
        style_sheet(ws)
    wb.save(path)


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="102A43")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            value = str(cell.value or "").upper()
            if value in {"CRITICAL", "BLOCK_IMPORT", "SENIOR_PROCUREMENT_APPROVAL", "EXECUTIVE_ESCALATION"}:
                cell.fill = PatternFill("solid", fgColor="F4CCCC")
            elif value in {"HIGH", "HIGH_RISK", "DOUBLE_VALIDATION", "SENIOR_REVIEW", "NEEDS_MORE_DATA", "PENDING"}:
                cell.fill = PatternFill("solid", fgColor="FCE5CD")
            elif value in {"MEDIUM", "PROCUREMENT_REVIEW", "IN_REVIEW"}:
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
            elif value in {"LOW", "APPROVED", "STANDARD_REVIEW"}:
                cell.fill = PatternFill("solid", fgColor="D9EAD3")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        ws.column_dimensions[get_column_letter(column[0].column)].width = width


def main() -> int:
    start = time.perf_counter()
    print("=== SP2I Procurement Review Workbench ===")
    try:
        queue, audit, dashboard_rows = build_queue()
        summary = dashboard_summary(queue)
        override_log = build_override_log()
        comments = build_comments_template()

        write_xlsx(QUEUE_OUT, {"REVIEW_QUEUE": queue})
        write_xlsx(OVERRIDE_OUT, {"OVERRIDE_LOG": override_log, "APPROVAL_RULES": approval_rules()})
        write_xlsx(COMMENTS_OUT, {"GOVERNANCE_COMMENTS": comments})
        write_xlsx(DASHBOARD_OUT, {"DASHBOARD_SUMMARY": summary, "DASHBOARD_ROWS": dashboard_rows})
        write_xlsx(AUDIT_OUT, {"AUDIT_TRAIL": audit, "REVIEW_ACTIONS": [{"ACTION": action} for action in REVIEW_ACTIONS]})

        stats = {
            "generated_at": datetime.now().replace(microsecond=0).isoformat(),
            "source_candidates": CANDIDATES.name,
            "review_items": len(queue),
            "manual_validation_required": sum(1 for item in queue if item["COCKPIT_MANUAL_VALIDATION"]),
            "priority_distribution": count_by(queue, "REVIEW_PRIORITY"),
            "status_distribution": count_by(queue, "REVIEW_STATUS"),
            "escalation_distribution": count_by(queue, "REQUIRED_ESCALATION_LEVEL"),
            "families": count_by(queue, "FAMILY_NAME"),
            "active_overrides": 0,
            "active_comments": 0,
            "duration_seconds": round(time.perf_counter() - start, 2),
        }
        STATS_OUT.write_text(json.dumps(stats, indent=2, ensure_ascii=False), encoding="utf-8")
        print(json.dumps(stats, indent=2, ensure_ascii=False))
        print("\nLivrables:")
        for path in [QUEUE_OUT, OVERRIDE_OUT, COMMENTS_OUT, DASHBOARD_OUT, AUDIT_OUT, STATS_OUT]:
            print(f"- {path.name}")
        return 0
    except Exception:
        print("\nERREUR PROCUREMENT REVIEW WORKBENCH")
        traceback.print_exc()
        return 1


def count_by(rows: list[dict[str, Any]], key: str) -> dict[str, int]:
    values: dict[str, int] = {}
    for row in rows:
        value = as_text(row.get(key)) or "EMPTY"
        values[value] = values.get(value, 0) + 1
    return dict(sorted(values.items()))


def approval_rules() -> list[dict[str, Any]]:
    return [
        {
            "RULE_ID": "CONFIDENCE_LOW",
            "CONDITION": "CONFIDENCE_LEVEL = LOW",
            "REQUIRED_ACTION": "Review obligatoire",
            "APPROVAL_LEVEL": "PROCUREMENT_REVIEWER",
        },
        {
            "RULE_ID": "DRIFT_CRITICAL",
            "CONDITION": "DRIFT_LEVEL = CRITICAL",
            "REQUIRED_ACTION": "Escalation obligatoire",
            "APPROVAL_LEVEL": "EXECUTIVE_ESCALATION",
        },
        {
            "RULE_ID": "BLOCK_IMPORT",
            "CONDITION": "DECISION_BLOCKER = BLOCK_IMPORT",
            "REQUIRED_ACTION": "Validation senior obligatoire",
            "APPROVAL_LEVEL": "SENIOR_PROCUREMENT_APPROVAL",
        },
        {
            "RULE_ID": "HIGH_RISK",
            "CONDITION": "DECISION_BLOCKER = HIGH_RISK",
            "REQUIRED_ACTION": "Double validation obligatoire",
            "APPROVAL_LEVEL": "DOUBLE_VALIDATION",
        },
    ]


if __name__ == "__main__":
    raise SystemExit(main())

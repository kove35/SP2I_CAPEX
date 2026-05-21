from __future__ import annotations

import hashlib
import json
import math
import re
import sys
import time
import traceback
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "07_API_BACKEND"
DQE_PATH = ROOT / "03_DONNEES_REFERENCE" / "DQE_PROJECT_SP2I.xlsx"
SHEET_NAME = "DQE_CLEAN"

MASTER_PATH = ROOT / "MASTER_REFERENCE_ENTERPRISE.xlsx"
QUALITY_AUDIT_PATH = ROOT / "MASTER_REFERENCE_QUALITY_AUDIT.xlsx"
BENCHMARK_PATH = ROOT / "BENCHMARK_AFRIQUE_CENTRALE.xlsx"
LOW_CONFIDENCE_PATH = ROOT / "TOP_LOW_CONFIDENCE_REFERENCES.xlsx"
STATS_PATH = ROOT / "MASTER_REFERENCE_ENTERPRISE_STATS.json"

PHASE_FAMILY = "ELECTRICITE"
REGION_FACTORS = {
    "CONGO_BRAZZAVILLE": 1.08,
    "CAMEROUN": 1.03,
    "GABON": 1.12,
    "CENTRAL_AFRICA": 1.0,
}

if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from app.core.equipment_classification_engine import EquipmentClassificationEngine
from app.core.financial_reference_engine import FinancialReferenceEngine
from app.core.semantic_normalization_engine import SemanticNormalizationEngine


def normalize_key(value: Any) -> str:
    text = "" if value is None else str(value)
    text = re.sub(r"[^A-Za-z0-9]+", "_", text.strip())
    return re.sub(r"_+", "_", text).strip("_").upper()


def number(value: Any, default: float = 0.0) -> float:
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return default if math.isnan(value) else float(value)
    text = str(value).strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return default


def choose(row: dict[str, Any], aliases: list[str], default: Any = "") -> Any:
    for alias in aliases:
        key = normalize_key(alias)
        if key in row and row[key] not in (None, ""):
            return row[key]
    return default


def read_dqe() -> list[dict[str, Any]]:
    if not DQE_PATH.exists():
        raise FileNotFoundError(f"DQE officiel introuvable: {DQE_PATH}")
    wb = load_workbook(DQE_PATH, data_only=True, read_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Onglet {SHEET_NAME} introuvable dans {DQE_PATH.name}")
    ws = wb[SHEET_NAME]
    raw_headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    headers = [normalize_key(header or f"COL_{index + 1}") for index, header in enumerate(raw_headers)]
    rows: list[dict[str, Any]] = []
    for rownum, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        payload = {headers[index]: value for index, value in enumerate(values)}
        if any(value not in (None, "") for value in payload.values()):
            payload["_ROWNUM"] = rownum
            rows.append(payload)
    wb.close()
    return rows


def reference_id(family: str, normalized: str, unit: str) -> str:
    digest = hashlib.sha1(f"{family}|{normalized}|{unit}".encode("utf-8")).hexdigest()[:8].upper()
    return f"SP2I-{family[:3]}-{digest}"


def confidence_level(score: float, price_status: str, source_count: int) -> str:
    if price_status.startswith("CRITICAL") or score < 58:
        return "LOW"
    if score >= 78 and source_count >= 2 and price_status == "OK":
        return "HIGH"
    return "MEDIUM"


def risk_level(confidence: str, price_status: str, importability: str) -> str:
    if confidence == "LOW" or price_status.startswith("CRITICAL"):
        return "HIGH"
    if importability == "LOW" or price_status == "WARNING":
        return "MEDIUM"
    return "LOW"


def price_status(price: float, price_min: float, price_max: float) -> str:
    if price <= 0:
        return "CRITICAL_ZERO_PRICE"
    if price_min and price < price_min * 0.2:
        return "CRITICAL_LOW_MAGNITUDE"
    if price_min and price < price_min:
        return "WARNING_BELOW_BENCHMARK"
    if price_max and price > price_max * 2:
        return "CRITICAL_HIGH_MAGNITUDE"
    if price_max and price > price_max:
        return "WARNING_ABOVE_BENCHMARK"
    return "OK"


def importability_for(classification: dict[str, Any], type_ligne: str) -> str:
    if type_ligne in {"TOTAL", "DOCUMENTAIRE", "VALIDATION", "TITRE_SECTION", "TRAVAUX", "PRESTATION"}:
        return "LOW"
    equipment_type = str(classification.get("equipment_type") or "")
    family = str(classification.get("family") or "")
    if equipment_type in {"UNKNOWN", "NON_IMPORTABLE", ""}:
        return "MEDIUM" if family in {"ELECTRICITE", "SOLAIRE"} else "LOW"
    return "HIGH"


def build_rows() -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    classifier = EquipmentClassificationEngine()
    references = FinancialReferenceEngine()
    normalizer = SemanticNormalizationEngine()
    grouped: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    quality_rows: list[dict[str, Any]] = []

    for row in read_dqe():
        designation = str(choose(row, ["DESIGNATION"], ""))
        lot = str(choose(row, ["LOT"], ""))
        sous_lot = str(choose(row, ["SOUS_LOT"], ""))
        type_ligne = str(choose(row, ["TYPE_LIGNE"], "")).upper()
        unit = str(choose(row, ["UNITE"], "U") or "U").upper()
        price = number(choose(row, ["PU_ESTIME_LOCAL_FCFA"], 0))
        context = {
            "designation": designation,
            "LOT": lot,
            "SOUS_LOT": sous_lot,
            "TYPE_LIGNE": type_ligne,
            "region": "CONGO_BRAZZAVILLE",
            "quality": "MEDIUM",
        }
        classification = classifier.classify_line(context)
        family = str(classification.get("family") or "GENERAL")
        if family not in {PHASE_FAMILY, "SOLAIRE"}:
            continue
        normalized = normalizer.normalize(designation)["normalized_designation"].upper()
        reference = references.get_reference(classification.get("equipment_type"), region="CONGO_BRAZZAVILLE", quality="MEDIUM")
        price_min = number(reference.get("price_min"), 0)
        price_max = number(reference.get("price_max"), 0)
        if not price_min or not price_max:
            prices = [p for p in [price] if p > 0]
            price_min = min(prices) if prices else 0
            price_max = max(prices) if prices else 0
        status = price_status(price, price_min, price_max)
        importability = importability_for(classification, type_ligne)
        key = (normalized, family, unit)
        item = {
            "row": row["_ROWNUM"],
            "raw_designation": designation,
            "normalized": normalized,
            "family": family,
            "subcategory": classification.get("subcategory") or sous_lot,
            "equipment_type": classification.get("equipment_type") or "UNKNOWN",
            "unit": unit,
            "price": price,
            "price_min": price_min,
            "price_max": price_max,
            "type_ligne": type_ligne,
            "importability": importability,
            "classification_score": number(classification.get("semantic_confidence_score"), 0),
            "classification_confidence": classification.get("classification_confidence") or "LOW",
            "classification_reason": classification.get("classification_reason") or "",
            "price_status": status,
        }
        grouped[key].append(item)
        quality_rows.append({
            "ROW": item["row"],
            "RAW_DESIGNATION": designation,
            "DESIGNATION_NORMALISEE": normalized,
            "FAMILLE": family,
            "SOUS_FAMILLE": item["subcategory"],
            "TYPE_EQUIPEMENT": item["equipment_type"],
            "TYPE_LIGNE": type_ligne,
            "PU_LOCAL_FCFA": price,
            "PRICE_MIN_FCFA": price_min,
            "PRICE_MAX_FCFA": price_max,
            "PRICE_STATUS": status,
            "IMPORTABILITY": importability,
            "CLASSIFICATION_CONFIDENCE": item["classification_confidence"],
            "CLASSIFICATION_REASON": item["classification_reason"],
        })

    master_rows: list[dict[str, Any]] = []
    benchmark_rows: list[dict[str, Any]] = []
    today = date.today().isoformat()
    for (normalized, family, unit), items in sorted(grouped.items()):
        prices = [item["price"] for item in items if item["price"] > 0]
        source_count = len(items)
        representative = max(items, key=lambda item: item["classification_score"])
        ref_min = max(min([item["price_min"] for item in items if item["price_min"] > 0] or prices or [0]), 0)
        ref_max = max([item["price_max"] for item in items if item["price_max"] > 0] or prices or [0])
        observed_min = min(prices) if prices else 0
        observed_max = max(prices) if prices else 0
        price_min = min([value for value in [ref_min, observed_min] if value > 0] or [0])
        price_max = max(ref_max, observed_max)
        status_counts = defaultdict(int)
        for item in items:
            status_counts[item["price_status"]] += 1
        dominant_status = "OK" if status_counts["OK"] else max(status_counts.items(), key=lambda entry: entry[1])[0]
        avg_classification = sum(item["classification_score"] for item in items) / max(source_count, 1)
        confidence = confidence_level(avg_classification, dominant_status, source_count)
        importability = representative["importability"]
        risk = risk_level(confidence, dominant_status, importability)
        fob = round(price_min * 0.62, 2) if price_min else 0
        moq = "1 U" if unit in {"U", "ENS", "FORFAIT"} else f"100 {unit}"
        master = {
            "REFERENCE_ID": reference_id(family, normalized, unit),
            "DESIGNATION_NORMALISEE": normalized,
            "RAW_DESIGNATION": representative["raw_designation"],
            "FAMILLE": PHASE_FAMILY if family == "SOLAIRE" else family,
            "SOUS_FAMILLE": representative["subcategory"],
            "TYPE_EQUIPEMENT": representative["equipment_type"],
            "UNITE": unit,
            "PRICE_MIN_FCFA": round(price_min, 2),
            "PRICE_MAX_FCFA": round(price_max, 2),
            "PU_CHINE_FOB_FCFA": fob,
            "IMPORTABILITY": importability,
            "RISK_LEVEL": risk,
            "MOQ": moq,
            "INCOTERM": "FOB",
            "FOURNISSEUR_CHINE": "A_VERIFIER_GOUVERNANCE",
            "BENCHMARK_REGION": "CONGO_BRAZZAVILLE|CAMEROUN|GABON|CENTRAL_AFRICA",
            "QUALITY_LEVEL": "MEDIUM",
            "CONFIDENCE_LEVEL": confidence,
            "LAST_VALIDATION_DATE": today,
            "SOURCE_REFERENCE": "DQE_PROJECT_SP2I.xlsx::DQE_CLEAN",
        }
        master_rows.append(master)
        for region, factor in REGION_FACTORS.items():
            benchmark_rows.append({
                "REFERENCE_ID": master["REFERENCE_ID"],
                "DESIGNATION_NORMALISEE": normalized,
                "FAMILLE": master["FAMILLE"],
                "SOUS_FAMILLE": master["SOUS_FAMILLE"],
                "REGION": region,
                "QUALITY_LEVEL": "MEDIUM",
                "PRICE_MIN_FCFA": round(price_min * factor, 2),
                "PRICE_MAX_FCFA": round(price_max * factor, 2),
                "PU_CHINE_FOB_FCFA": fob,
                "CONFIDENCE_LEVEL": confidence,
                "PRICE_STATUS": dominant_status,
                "SOURCE_COUNT": source_count,
            })
    return master_rows, quality_rows, benchmark_rows


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="0F172A")
    header_font = Font(color="F8FAFC", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 12), 58)
    ws.freeze_panes = "A2"


def write_xlsx(path: Path, sheets: dict[str, list[dict[str, Any]]]) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title[:31])
        if rows:
            headers = list(rows[0].keys())
            ws.append(headers)
            for row in rows:
                ws.append([row.get(header) for header in headers])
        else:
            ws.append(["STATUS"])
            ws.append(["NO_DATA"])
        style_sheet(ws)
    wb.save(path)


def main() -> int:
    start = time.perf_counter()
    print("=== SP2I_CAPEX Master Reference Enterprise Builder ===")
    print(f"Source: {DQE_PATH} | Sheet: {SHEET_NAME}")
    try:
        master_rows, quality_rows, benchmark_rows = build_rows()
        low_confidence = [
            row for row in master_rows
            if row["CONFIDENCE_LEVEL"] != "HIGH" or row["RISK_LEVEL"] != "LOW"
        ]
        low_confidence.sort(key=lambda row: (row["CONFIDENCE_LEVEL"], row["RISK_LEVEL"], row["DESIGNATION_NORMALISEE"]))
        write_xlsx(MASTER_PATH, {"MASTER_REFERENCE": master_rows})
        write_xlsx(QUALITY_AUDIT_PATH, {"QUALITY_AUDIT": quality_rows})
        write_xlsx(BENCHMARK_PATH, {"BENCHMARKS_MEDIUM": benchmark_rows})
        write_xlsx(LOW_CONFIDENCE_PATH, {"LOW_CONFIDENCE": low_confidence})
        stats = {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "source_file": str(DQE_PATH),
            "sheet": SHEET_NAME,
            "phase_family": PHASE_FAMILY,
            "master_references": len(master_rows),
            "quality_rows": len(quality_rows),
            "benchmark_rows": len(benchmark_rows),
            "low_confidence_references": len(low_confidence),
            "confidence_distribution": {
                level: len([row for row in master_rows if row["CONFIDENCE_LEVEL"] == level])
                for level in ["HIGH", "MEDIUM", "LOW"]
            },
            "risk_distribution": {
                level: len([row for row in master_rows if row["RISK_LEVEL"] == level])
                for level in ["LOW", "MEDIUM", "HIGH"]
            },
            "duration_seconds": round(time.perf_counter() - start, 2),
        }
        STATS_PATH.write_text(json.dumps(stats, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps(stats, ensure_ascii=False, indent=2))
        print("\nLivrables:")
        for path in [MASTER_PATH, QUALITY_AUDIT_PATH, BENCHMARK_PATH, LOW_CONFIDENCE_PATH, STATS_PATH]:
            print(f"- {path.name}")
        return 0
    except Exception:
        print("MASTER REFERENCE BUILD FAILED")
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
